"""Nemo v92 "Wire": the libraries installed through Forge become chat commands, each run safely.

Offline. The helper programs are real Python snippets, but in most tests the runner is replaced by a recorder; a second group runs the REAL runner on tiny programs (guard, one at a time,
clean environment, time limit, failure words). A third group runs every helper program against the REAL libraries, only when NEMO_WIRE_PY points at a Python that has them
(pdfplumber, pytesseract + tesseract, trafilatura, pandas-ta-classic, quantstats); it is skipped otherwise.
No Telegram, no network, no paid calls, no trades, nothing is ever sent to the broker.

    NEMO_FILE=/path/to/nemotron_bot.py python -m unittest tests.test_wire92 -v
    NEMO_WIRE_PY=/path/to/python-with-the-libraries ... (adds the real-library group)
"""
import ast
import io
import json
import os
import re
import shutil
import stat
import subprocess
import sys
import time
import unittest
import zipfile
from unittest import mock

from tests import test_cortex83 as base
from tests.test_forge91 import ForgeCase, HttpResp, GB, OWNER_ID

PNG = None


def setUpModule():
    if base.m is None:
        base.setUpModule()


def wire_source():
    src = open(base.NEMO_FILE, encoding='utf-8').read()
    i = src.index('# NEMO 92 - WIRE')
    j = src.find('# NEMO 93 - SCOUT', i)
    return src[i:(j if j > 0 else src.rindex("if __name__"))]


class WireCase(ForgeCase):
    def setUp(self):
        super().setUp()
        m = self.m
        self.docs = []
        self.child_calls = []
        self.child_result = {}

        def send_document(cid, path, name, mime):
            with open(path, 'rb') as fh:
                self.docs.append({'cid': cid, 'name': name, 'mime': mime, 'bytes': fh.read()})
            return True

        def fake_child(code, args=(), need_mb=200, timeout=120, job=None):
            self.child_calls.append({'code': code, 'args': [str(a) for a in args], 'need': need_mb, 'job': job, 'timeout': timeout})
            key = next((k for k in ('PDF', 'OCR', 'ARTICLE', 'IND', 'QS', 'NSE', 'STT') if code is getattr(m, '_N92_CODE_' + k)), None)
            r = self.child_result.get(key)
            if isinstance(r, Exception):
                raise r
            if callable(r):
                return r(args, job)
            return r if r is not None else {}
        self.start(m, 'send_document', send_document)
        self.start(m, '_n92_child', fake_child)
        self.start(m, '_n92_guard', lambda need: (True, ''))
        for k in m._N92_STATS:
            m._N92_STATS[k] = 0
        self.job_dirs_before = set(os.listdir(m._n92_dir()))

    def texts_to_owner(self, since=0):
        return '\n'.join(t for c, t in self.sent[since:] if c == OWNER_ID)

    def say(self, text, **extra):
        n = len(self.sent)
        self.m.handle(self.msg(text, **extra))
        return self.texts_to_owner(n)

    def pdf_msg(self, caption='tables from this pdf', reply=False):
        doc = {'file_id': 'f1', 'file_name': 'statement.pdf', 'mime_type': 'application/pdf', 'file_size': 1000}
        if reply:
            return self.msg(caption, reply_to_message={'document': doc})
        return self.msg('', caption=caption, document=doc)

    def no_job_folders_left(self):
        self.assertEqual(set(os.listdir(self.m._n92_dir())), self.job_dirs_before, 'every job folder is removed afterwards')


# ===================================================================================================================
# 1. THE RUNNER (real): one at a time, only when memory is free, a clean environment, plain failure words
# ===================================================================================================================
class TestRunner(ForgeCase):
    def setUp(self):
        super().setUp()
        for k in self.m._N92_STATS:
            self.m._N92_STATS[k] = 0

    def test_a_helper_prints_one_json_line_and_that_is_the_result(self):
        out = self.m._n92_child('import json; print("noise"); print(json.dumps({"a": 1, "b": [2, 3]}))', need_mb=1)
        self.assertEqual(out, {'a': 1, 'b': [2, 3]})
        self.assertEqual(self.m._N92_STATS['jobs'], 1)

    def test_nothing_readable_is_reported_plainly(self):
        for code in ('print("not json")', 'print("[1, 2]")', 'pass'):
            with self.assertRaises(self.m._N92Err) as cm:
                self.m._n92_child(code, need_mb=1)
            self.assertEqual(cm.exception.code, 'bad_output')

    def test_a_helper_never_sees_a_key_or_token(self):
        for k in ('NVIDIA_API_KEY', 'GITHUB_TOKEN', 'TELEGRAM_BOT_TOKEN', 'MY_SECRET_THING'):
            os.environ[k] = 'NEMO-WIRE-SECRET-VALUE'
        try:
            out = self.m._n92_child('import json, os; print(json.dumps({"seen": sorted(k for k in os.environ if any(w in k for w in ("KEY", "TOKEN", "SECRET")))}))', need_mb=1)
        finally:
            for k in ('NVIDIA_API_KEY', 'GITHUB_TOKEN', 'TELEGRAM_BOT_TOKEN', 'MY_SECRET_THING'):
                os.environ.pop(k, None)
        self.assertEqual(out['seen'], [])

    def test_it_runs_at_low_priority_without_a_shell(self):
        seen = []
        real = self.m._n91_run
        self.start(self.m, '_n91_run', lambda argv, timeout=120, env=None, cwd=None: (seen.append(list(argv)) or real(argv, timeout, env, cwd)))
        self.m._n92_child('import json; print(json.dumps({}))', ['x.txt', 5], need_mb=1)
        argv = seen[0]
        if shutil.which('nice'):
            self.assertEqual(argv[:3], ['nice', '-n', '10'])
        self.assertIn('-c', argv)
        self.assertEqual(argv[-2:], ['x.txt', '5'])
        self.assertNotIn('sh', argv[:1])

    def test_a_job_that_takes_too_long_is_stopped(self):
        with self.assertRaises(self.m._N92Err) as cm:
            self.m._n92_child('import time; time.sleep(30)', need_mb=1, timeout=1)
        self.assertEqual(cm.exception.code, 'timeout')
        self.assertIn('took too long', cm.exception.msg)
        self.assertFalse(self.m._N92_LOCK.locked(), 'the lock is released after a failure')

    def test_a_missing_library_says_what_to_install(self):
        e = self.m._n92_failure(1, '', "Traceback...\nModuleNotFoundError: No module named 'pdfplumber'")
        self.assertEqual(e.code, 'not_installed')
        self.assertIn('Say “install pdfplumber into yourself”', e.msg)
        e = self.m._n92_failure(1, '', "ModuleNotFoundError: No module named 'pandas_ta_classic'")
        self.assertIn('install pandas-ta-classic into yourself', e.msg)
        e = self.m._n92_failure(1, '', "pytesseract.pytesseract.TesseractNotFoundError: tesseract is not installed or it's not in your PATH")
        self.assertEqual(e.code, 'no_tesseract')
        self.assertIn('forge apt tesseract-ocr tesseract-ocr-eng tesseract-ocr-hin', e.msg)
        e = self.m._n92_failure(1, '', 'httpx.ConnectError: huggingface.co unreachable')
        self.assertEqual(e.code, 'model_download')
        self.assertIn('huggingface.co', e.msg)
        e = self.m._n92_failure(1, '', 'ValueError: something odd with sk-' + 'a1' * 20)
        self.assertEqual(e.code, 'failed')
        self.assertNotIn('a1a1a1', e.msg, 'hidden values are masked')

    def test_a_real_import_failure_goes_through_the_whole_path(self):
        with self.assertRaises(self.m._N92Err) as cm:
            self.m._n92_child('import definitely_not_a_module_xyz', need_mb=1)
        self.assertEqual(cm.exception.code, 'not_installed')

    def test_low_memory_refuses_before_anything_starts(self):
        started = []
        self.start(self.m, '_n91_run', lambda *a, **k: started.append(a) or (0, '{}', ''))
        self.start(self.m, '_n91_meminfo', lambda: {'MemAvailable': 400 * 1048576, 'MemTotal': 4 * GB})
        with self.assertRaises(self.m._N92Err) as cm:
            self.m._n92_child('print(1)', need_mb=250)
        self.assertEqual(cm.exception.code, 'low_memory')
        self.assertIn('0.4 GB available', cm.exception.msg)
        self.assertIn('trading bot stays out of swap', cm.exception.msg)
        self.assertEqual(started, [])
        self.assertEqual(self.m._N92_STATS['refused_memory'], 1)

    def test_the_guard_leaves_a_margin_and_notices_a_stalling_server(self):
        self.start(self.m, '_n91_meminfo', lambda: {'MemAvailable': 600 * 1048576})
        self.start(self.m, '_n91_psi', lambda: None)
        self.assertTrue(self.m._n92_guard(250)[0])
        self.assertFalse(self.m._n92_guard(350)[0], '600 MB free is not enough for a 350 MB job plus the margin')
        self.start(self.m, '_n91_meminfo', lambda: {'MemAvailable': 3 * GB})
        self.start(self.m, '_n91_psi', lambda: {'some': 35.0})
        ok, why = self.m._n92_guard(100)
        self.assertFalse(ok)
        self.assertIn('waiting for memory', why)
        self.start(self.m, '_n91_meminfo', lambda: {})
        self.assertTrue(self.m._n92_guard(5000)[0], 'a system that does not say is not blocked')

    def test_only_one_heavy_job_at_a_time(self):
        self.assertTrue(self.m._N92_LOCK.acquire(blocking=False))
        try:
            with self.assertRaises(self.m._N92Err) as cm:
                self.m._n92_child('print(1)', need_mb=1)
        finally:
            self.m._N92_LOCK.release()
        self.assertEqual(cm.exception.code, 'busy')
        self.assertEqual(self.m._N92_STATS['refused_busy'], 1)

    def test_job_folders_are_private_and_only_such_folders_are_ever_deleted(self):
        job = self.m._n92_new_job()
        self.assertEqual(os.stat(job).st_mode & 0o777, 0o700)
        self.assertTrue(os.path.basename(job).startswith('J92-'))
        outside = os.path.join(self.tmp.name, 'precious')
        os.makedirs(outside)
        self.m._n92_rm_job(outside)
        self.m._n92_rm_job('/etc')
        self.assertTrue(os.path.isdir(outside) and os.path.isdir('/etc'))
        self.m._n92_rm_job(job)
        self.assertFalse(os.path.exists(job))

    def test_old_job_folders_are_swept_at_start_up(self):
        job = self.m._n92_new_job()
        os.utime(job, (time.time() - 2 * 86400,) * 2)
        keep = self.m._n92_new_job()
        self.assertEqual(self.m._n92_prune_jobs(), 1)
        self.assertFalse(os.path.exists(job))
        self.assertTrue(os.path.exists(keep))


# ===================================================================================================================
# 2. PDF TABLES
# ===================================================================================================================
class TestPdfTables(WireCase):
    TABLE = {'pages': 2, 'text_chars': 900, 'tables': [{'page': 1, 'rows': [['Contract', 'Qty', 'Price'], ['NIFTY OCT FUT', '75', '24,590.50'], ['BANKNIFTY OCT FUT', '30', '52110.00']]},
                                                         {'page': 2, 'rows': [['Total', '105', '-1,200']]}]}

    def setUp(self):
        super().setUp()
        self.downloaded = []
        self.start(self.m, '_n85_tg_download', lambda fid, size=0: (self.downloaded.append((fid, size)) or b'%PDF-1.4 fake'))

    def test_a_pdf_with_the_words_becomes_a_spreadsheet_with_real_numbers(self):
        from openpyxl import load_workbook
        self.child_result['PDF'] = self.TABLE
        out = self.say('', caption='tables from this pdf', document={'file_id': 'f1', 'file_name': 'Contract Note.pdf', 'mime_type': 'application/pdf', 'file_size': 1000})
        self.assertIn('2 tables found in Contract Note.pdf (2 pages read)', out)
        self.assertIn('NIFTY OCT FUT | 75 | 24,590.50', out)
        self.assertIn('check them against the PDF', out)
        self.assertEqual(self.downloaded, [('f1', 1000)])
        self.assertEqual(len(self.docs), 1)
        self.assertEqual((self.docs[0]['name'], self.docs[0]['mime']), ('tables_Contract_Note.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
        wb = load_workbook(io.BytesIO(self.docs[0]['bytes']))
        self.assertEqual(wb.sheetnames, ['p1_t1', 'p2_t2'])
        ws = wb['p1_t1']
        self.assertEqual([c.value for c in ws[2]], ['NIFTY OCT FUT', 75.0, 24590.5])
        self.assertEqual([c.value for c in ws[3]], ['BANKNIFTY OCT FUT', 30.0, 52110.0])
        self.assertEqual([c.value for c in wb['p2_t2'][1]], ['Total', 105.0, -1200.0])
        self.assertEqual(self.passed, [], 'the older layers never see it')
        self.no_job_folders_left()

    def test_replying_to_a_pdf_works_too_and_the_helper_gets_the_file_and_a_page_limit(self):
        self.child_result['PDF'] = self.TABLE
        self.say('tables from this pdf', reply_to_message={'document': {'file_id': 'f9', 'file_name': 'a.pdf', 'mime_type': 'application/pdf', 'file_size': 5}})
        self.assertEqual(self.downloaded[0][0], 'f9')
        call = self.child_calls[0]
        self.assertTrue(call['args'][0].endswith('in.pdf'))
        self.assertEqual(call['args'][1], '40')
        self.assertEqual((call['need'], call['timeout']), (250, 150))

    def test_the_phrases_that_count(self):
        for text in ('tables from this pdf', 'Tables', 'extract tables from this pdf', 'get the tables', 'pdf tables', 'please extract the tables from this statement', 'show tables in the file'):
            self.child_calls.clear()
            self.child_result['PDF'] = self.TABLE
            self.m.handle(self.pdf_msg(text))
            self.assertEqual(len(self.child_calls), 1, text)

    def test_other_talk_about_tables_is_left_alone(self):
        for text in ('set the tables for dinner', 'tables are nice', 'my tables in the report look wrong', 'how do I make tables in word'):
            n = len(self.passed)
            self.m.handle(self.pdf_msg(text))
            self.assertEqual(len(self.passed), n + 1, text)
        self.assertEqual(self.child_calls, [])

    def test_without_a_pdf_it_says_what_to_send(self):
        out = self.say('tables from this pdf')
        self.assertIn('Send me the PDF with the words', out)
        self.assertEqual(self.child_calls, [])

    def test_a_scan_without_text_is_called_a_scan(self):
        self.child_result['PDF'] = {'pages': 3, 'text_chars': 10, 'tables': []}
        out = self.say('', caption='tables from this pdf', document=self.pdf_msg()['document'])
        self.assertIn('I found no tables in statement.pdf (3 pages)', out)
        self.assertIn('probably a scan', out)
        self.assertEqual(self.docs, [])

    def test_a_text_pdf_without_ruled_tables_says_so_and_points_to_the_other_method(self):
        self.child_result['PDF'] = {'pages': 2, 'text_chars': 4000, 'tables': []}
        out = self.say('', caption='tables from this pdf', document=self.pdf_msg()['document'])
        self.assertIn('Only tables with ruled lines or clear columns are found; try /doc85 tables', out)

    def test_a_file_that_cannot_be_downloaded_or_a_helper_that_fails_gets_one_plain_sentence(self):
        self.start(self.m, '_n85_tg_download', lambda fid, size=0: (_ for _ in ()).throw(ValueError('file is larger than 20 MB')))
        self.assertIn('I could not get that file: file is larger than 20 MB.', self.say('', caption='tables from this pdf', document=self.pdf_msg()['document']))
        self.start(self.m, '_n85_tg_download', lambda fid, size=0: b'x')
        self.child_result['PDF'] = self.m._N92Err('not_installed', 'pdfplumber is not installed in me yet. Say “install pdfplumber into yourself” (you get a card to approve).')
        self.assertIn('Say “install pdfplumber into yourself”', self.say('', caption='tables from this pdf', document=self.pdf_msg()['document']))
        self.no_job_folders_left()

    def test_numbers_become_numbers_only_when_they_look_like_numbers(self):
        c = self.m._n92_cell
        self.assertEqual((c('1,234.50'), c('-7'), c('0.5'), c('100')), (1234.5, -7.0, 0.5, 100.0))
        for keep in ('NIFTY', '12-05-2026', '1,23', '1.2.3', '+91 98765', '', None, '24590.50 CE'):
            self.assertEqual(c(keep), str(keep if keep is not None else '').strip(), keep)


# ===================================================================================================================
# 3. THE TEXT IN A PICTURE
# ===================================================================================================================
class TestOcr(WireCase):
    def setUp(self):
        super().setUp()
        self.start(self.m, '_n90_source_image', lambda cid, msg: (b'\x89PNG fake', 'attached') if (msg.get('photo') or msg.get('reply_to_message')) else (None, ''))

    def photo(self, caption, **extra):
        return self.msg('', caption=caption, photo=[{'file_id': 'p1', 'file_size': 100}], **extra)

    def test_a_picture_with_the_words_gives_its_text(self):
        self.child_result['OCR'] = {'text': 'NIFTY OCT FUT Qty 75 Price 24590.50\nभारतीय रिज़र्व बैंक', 'lang': 'eng+hin', 'size': [900, 160]}
        n = len(self.sent)
        self.m.handle(self.photo('read this picture'))
        out = self.texts_to_owner(n)
        self.assertIn('Text found (eng+hin, picture 900x160)', out)
        self.assertIn('NIFTY OCT FUT Qty 75 Price 24590.50', out)
        self.assertIn('भारतीय रिज़र्व बैंक', out)
        self.assertIn('Check numbers against the picture', out)
        self.assertEqual(self.child_calls[0]['args'][1], 'eng+hin')
        self.assertEqual((self.child_calls[0]['need'], self.child_calls[0]['timeout']), (350, 120))
        self.assertEqual(self.passed, [])
        self.no_job_folders_left()

    def test_a_language_can_be_asked_for(self):
        self.child_result['OCR'] = {'text': 'x', 'lang': 'hin', 'size': [1, 1]}
        for words, want in (('read this picture in hindi', 'hin'), ('read this picture in english', 'eng'), ('ocr', 'eng+hin'), ('read this picture in gujarati', 'guj+eng')):
            self.child_calls.clear()
            self.m.handle(self.photo(words))
            self.assertEqual(self.child_calls[0]['args'][1], want, words)

    def test_the_phrases_that_count_and_the_ones_that_do_not(self):
        for text in ('read this picture', 'ocr this', 'read the text in this screenshot', 'extract text from the image', "what's written on this photo", 'transcribe this scan'):
            self.child_calls.clear()
            self.child_result['OCR'] = {'text': 'x', 'lang': 'eng', 'size': [1, 1]}
            self.m.handle(self.photo(text))
            self.assertEqual(len(self.child_calls), 1, text)
        for text in ('make this picture black and white', 'read my mail', 'what a nice photo', 'draw a cat', 'read this'):
            n = len(self.passed)
            self.m.handle(self.photo(text))
            self.assertEqual(len(self.passed), n + 1, text)

    def test_with_no_picture_it_says_what_to_send(self):
        self.assertIn('Send me the picture with the words', self.say('read this picture'))
        self.assertEqual(self.child_calls, [])

    def test_long_text_also_comes_as_a_file_and_empty_text_is_explained(self):
        self.child_result['OCR'] = {'text': 'line of text\n' * 400, 'lang': 'eng', 'size': [2000, 3000]}
        out = self.say('', caption='read this picture', photo=[{'file_id': 'p1'}])
        self.assertIn('(cut; the full text is in the file)', out)
        self.assertEqual(self.docs[0]['name'], 'picture_text.txt')
        self.assertGreater(len(self.docs[0]['bytes']), 3500)
        self.docs.clear()
        self.child_result['OCR'] = {'text': '  ', 'lang': 'eng+hin', 'size': [10, 10]}
        out = self.say('', caption='read this picture', photo=[{'file_id': 'p1'}])
        self.assertIn('I could not find readable text in that picture (language used: eng+hin)', out)
        self.assertEqual(self.docs, [])

    def test_a_missing_system_program_is_explained(self):
        self.child_result['OCR'] = self.m._n92_failure(1, '', 'TesseractNotFoundError: tesseract is not installed')
        out = self.say('', caption='read this picture', photo=[{'file_id': 'p1'}])
        self.assertIn('forge apt tesseract-ocr tesseract-ocr-eng tesseract-ocr-hin', out)


# ===================================================================================================================
# 4. THE CLEAN TEXT OF A WEB PAGE (through the existing safe fetcher)
# ===================================================================================================================
class TestArticle(WireCase):
    def setUp(self):
        super().setUp()
        self.fetched = []
        self.page = {'url': 'https://example.com/oil', 'final_url': 'https://example.com/oil?x=1', 'status': 200, 'ctype': 'text/html', 'text': 'Oil jumps\nplain page text ' * 20,
                     'html': '<html><body><article><p>Brent rose.</p></article></body></html>', 'title': 'Oil jumps', 'meta': {}, 'jsonld': []}
        self.start(self.m, '_n85_fetch', lambda url, **k: (self.fetched.append(url) or self.page))

    def test_a_link_with_the_words_gives_the_clean_article(self):
        self.child_result['ARTICLE'] = {'text': 'Oil jumps 4% on supply fears\n\nBrent rose 4% on Friday.\n\nRefiners may see pressure. ' + 'More detail follows in the full report. ' * 6, 'meta': {'title': 'Oil jumps 4%', 'author': 'A Reporter', 'date': '2026-10-02', 'sitename': 'Test Wire'}}
        out = self.say('read this article https://example.com/oil')
        self.assertEqual(self.fetched, ['https://example.com/oil'])
        self.assertIn('📰 Oil jumps 4%', out)
        self.assertIn('Test Wire · A Reporter · 2026-10-02', out)
        self.assertIn('Source: https://example.com/oil?x=1', out)
        self.assertIn('[page text, written by a third party: untrusted]', out)
        self.assertIn('Brent rose 4% on Friday.\n\nRefiners may see pressure.', out, 'paragraphs are kept')
        self.assertNotIn('found little here', out)
        self.assertEqual(self.http.calls, [], 'the page is opened only by the existing safe fetcher')
        self.assertEqual(self.passed, [])
        self.no_job_folders_left()

    def test_the_phrases_that_count(self):
        for text in ('read article https://example.com/a', 'clean text of https://example.com/a', 'article https://example.com/a', 'extract the page https://example.com/a', 'Read the story from https://example.com/a'):
            self.child_calls.clear()
            self.child_result['ARTICLE'] = {'text': 'x' * 300, 'meta': {}}
            self.say(text)
            self.assertEqual(len(self.child_calls), 1, text)

    def test_other_messages_with_links_are_left_to_the_older_layers(self):
        for text in ('read https://example.com/a', 'summarise https://example.com/a', 'https://example.com/a', 'what do you think of https://example.com/a', 'read the article'):
            n = len(self.passed)
            self.m.handle(self.msg(text))
            self.assertEqual(len(self.passed), n + 1, text)
        self.assertEqual(self.fetched, [])

    def test_hidden_values_in_a_page_are_masked(self):
        self.child_result['ARTICLE'] = {'text': 'Leaked: sk-' + 'a1' * 20 + ' and more words here. ' * 12, 'meta': {'title': 'T'}}
        out = self.say('read article https://example.com/a')
        self.assertNotIn('sk-' + 'a1' * 20, out)
        self.assertIn('and more words here.', out)

    def test_when_the_extractor_finds_little_the_plain_page_text_is_used_and_says_so(self):
        self.child_result['ARTICLE'] = {'text': 'short', 'meta': {}}
        out = self.say('read article https://example.com/a')
        self.assertIn('the article extractor found little here, so this is the plain page text', out)
        self.assertIn('plain page text', out)

    def test_a_long_article_also_comes_as_a_file(self):
        self.child_result['ARTICLE'] = {'text': 'Paragraph of the long article. ' * 300, 'meta': {'title': 'Long'}}
        self.say('read article https://example.com/a')
        self.assertEqual(self.docs[0]['name'], 'article.txt')

    def test_a_page_the_safe_fetcher_refuses_gets_its_plain_reason(self):
        self.start(self.m, '_n85_fetch', lambda url, **k: (_ for _ in ()).throw(ValueError('blocked: private address')))
        out = self.say('read article http://192.168.0.1/admin')
        self.assertIn('I could not open that page: blocked: private address.', out)
        self.assertEqual(self.child_calls, [])


# ===================================================================================================================
# 5. INDICATORS
# ===================================================================================================================
class TestIndicators(WireCase):
    IND = {'n': 60, 'close': 24270.95, 'prev_close': 24288.79, 'rsi': 39.1, 'rsi_prev': 41.1, 'atr': 65.92, 'adx': 15.6, 'plus_di': 15.2, 'minus_di': 23.7, 'st_line': 24468.49, 'st_dir': -1,
           'bb_low': 24253.76, 'bb_mid': 24320.0, 'bb_up': 24391.65, 'ema20': 24330.21, 'ema50': 24390.34, 'macd': 1.0, 'macd_signal': 0.4, 'macd_hist': 0.61}

    def setUp(self):
        super().setUp()
        self.asked = []
        closes = [24000.0 + i for i in range(60)]
        self.start(self.m, 'get_history', lambda t: (self.asked.append(t) or {'name': t, 'closes': closes, 'highs': [c + 10 for c in closes], 'lows': [c - 10 for c in closes], 'cur': 'INR'}))

    def test_indicators_for_an_index_use_the_usual_reader_and_read_in_plain_words(self):
        self.child_result['IND'] = self.IND
        out = self.say('indicators NIFTY')
        self.assertEqual(self.asked, ['^NSEI'], 'the index name is mapped the way the rest of Nemo does it')
        for part in ('📊 NIFTY: daily candles, last 60 days (descriptive, not a signal; nothing is sent to the broker)', 'Close 24270.95 (previous 24288.79)', 'RSI(14) 39.1, leaning weak (was 41.1)',
                     'ADX(14) 15.6: weak or no trend, sellers ahead (-DI 23.7 over +DI 15.2)', 'Supertrend(10,3): price is below the line at 24468.49 (downtrend)', 'Bollinger(20,2): 24253.76 to 24391.65, price at 12% of the band',
                     'EMA20 24330.21 below EMA50 24390.34', 'MACD histogram 0.61 (momentum rising)', 'ATR(14) 65.92, about 0.3% of the price', 'computed with pandas-ta-classic'):
            self.assertIn(part, out)
        self.assertEqual(self.passed, [])
        self.no_job_folders_left()

    def test_the_candles_go_to_the_helper_as_a_file(self):
        self.child_result['IND'] = lambda args, job: (self.assertEqual(sorted(json.load(open(args[0]))), ['closes', 'highs', 'lows']) or self.IND)
        self.say('indicators RELIANCE')
        self.assertEqual(self.asked, ['RELIANCE'])
        self.assertEqual((self.child_calls[0]['need'], self.child_calls[0]['timeout']), (200, 90))

    def test_the_phrases_that_count_and_the_ones_that_do_not(self):
        self.child_result['IND'] = self.IND
        for text, sym in (('indicators for BANKNIFTY', '^NSEBANK'), ('technical read RELIANCE', 'RELIANCE'), ('show the indicators of NIFTY', '^NSEI'), ('indicator TCS', 'TCS')):
            self.asked.clear()
            self.say(text)
            self.assertEqual(self.asked, [sym], text)
        for text in ('technical analysis of NIFTY', 'what are indicators', 'indicators are great', 'rsi'):
            n = len(self.passed)
            self.m.handle(self.msg(text))
            self.assertEqual(len(self.passed), n + 1, text)

    def test_not_enough_candles_gets_an_honest_answer(self):
        self.start(self.m, 'get_history', lambda t: {'name': t, 'closes': [1.0] * 10, 'highs': [1.0] * 10, 'lows': [1.0] * 10, 'cur': ''})
        self.assertIn('I could not get enough daily candles for XYZ (I need at least 35)', self.say('indicators XYZ'))
        self.start(self.m, 'get_history', lambda t: None)
        self.assertIn('I could not get enough daily candles', self.say('indicators XYZ'))
        self.assertEqual(self.child_calls, [])

    def test_the_words_follow_the_numbers_at_the_edges(self):
        t = lambda **kw: self.m._n92_ind_text('X', dict({'n': 60, 'close': 100.0}, **kw))
        self.assertIn('RSI(14) 72.0, in the overbought zone', t(rsi=72.0))
        self.assertIn('RSI(14) 28.0, in the oversold zone', t(rsi=28.0))
        self.assertIn('RSI(14) 60.0, leaning strong', t(rsi=60.0))
        self.assertIn('RSI(14) 50.0, neutral', t(rsi=50.0))
        self.assertIn('very strong trend', t(adx=45.0))
        self.assertIn('strong trend', t(adx=30.0))
        self.assertIn('trend developing', t(adx=22.0))
        self.assertIn('buyers ahead', t(adx=30.0, plus_di=30.0, minus_di=10.0))
        self.assertIn('price is above the line at 95.00 (uptrend)', t(st_line=95.0, st_dir=1))
        self.assertIn('(above the top)', t(bb_low=90.0, bb_up=99.0))
        self.assertIn('(below the bottom)', t(bb_low=101.0, bb_up=110.0))
        self.assertIn('n/a', self.m._n92_ind_text('X', {'n': 5, 'close': None}))
        self.assertNotIn('buy', t(rsi=20.0, adx=50.0).lower().replace('buyers', ''), 'it describes, it never says buy or sell')
        self.assertNotIn('sell', t(rsi=80.0).lower().replace('sellers', ''))


# ===================================================================================================================
# 6. THE JOURNAL REPORT
# ===================================================================================================================
class TestTearsheet(WireCase):
    def trades(self, n=14, days=None):
        return [{'trade_date': '2026-09-%02d' % (1 + i), 'pnl': float(((-1) ** i) * (500 + 40 * i)), 'sym': 'NIFTY'} for i in range(n)]

    def setUp(self):
        super().setUp()
        self.asked_days = []
        self.rows = self.trades()
        self.start(self.m, '_n84_trades', lambda cid, days=None, **k: (self.asked_days.append(days) or list(self.rows)))

        def qs(args, job):
            self.qs_input = json.load(open(args[0]))
            with open(args[1], 'w') as f:
                f.write('<html>report</html>')
            return {'days': 14, 'sharpe': -0.39, 'max_drawdown': -0.0021, 'win_rate': 0.5, 'best': 0.0019, 'worst': -0.0020}
        self.child_result['QS'] = qs

    def test_the_journal_becomes_a_report_file_and_a_short_summary(self):
        out = self.say('tearsheet capital=5L')
        self.assertIn('JOURNAL REPORT: 14 trades on 14 days', out)
        self.assertIn('winning days 50.0%', out)
        self.assertIn('best day 0.2%', out)
        self.assertIn('max drawdown -0.2%', out)
        self.assertIn('Sharpe -0.39', out)
        self.assertIn('divided by capital 500000 (from you)', out)
        self.assertIn('nothing here is a forecast', out)
        self.assertEqual((self.docs[0]['name'], self.docs[0]['mime']), ('journal_report.html', 'text/html'))
        self.assertEqual(self.docs[0]['bytes'], b'<html>report</html>')
        self.assertEqual(len(self.qs_input['dates']), 14)
        self.assertAlmostEqual(self.qs_input['values'][0], 500.0 / 500000)
        self.assertEqual(self.passed, [])
        self.no_job_folders_left()

    def test_the_default_capital_comes_from_the_existing_settings_and_is_said_so(self):
        out = self.say('tearsheet')
        self.assertIn('capital %.0f (from your settings)' % self.m._P75_DEFAULTS['capital'], out)

    def test_days_and_capital_are_read_from_the_words(self):
        self.say('tearsheet days=30 capital=2.5L')
        self.assertEqual(self.asked_days[-1], 30)
        self.assertAlmostEqual(self.qs_input['values'][0], 500.0 / 250000)
        self.say('performance report last 90 days')
        self.assertEqual(self.asked_days[-1], 90)
        self.assertEqual(self.m._n92_money('5L'), 500000.0)
        self.assertEqual(self.m._n92_money('1.5 cr'), 15000000.0)
        self.assertEqual(self.m._n92_money('2,50,000'), 250000.0)
        self.assertEqual(self.m._n92_money('abc'), 0.0)

    def test_the_phrases_that_count_and_the_ones_that_do_not(self):
        for text in ('tearsheet', 'tear sheet of my journal', 'journal report', 'risk report from my trades', 'performance report capital=3L'):
            self.child_calls.clear()
            self.say(text)
            self.assertEqual(len(self.child_calls), 1, text)
        for text in ('my report is ready', 'write a report on oil', 'sheet music'):
            n = len(self.passed)
            self.m.handle(self.msg(text))
            self.assertEqual(len(self.passed), n + 1, text)

    def test_a_short_journal_is_not_dressed_up(self):
        self.rows = self.trades(6)
        out = self.say('tearsheet')
        self.assertIn('I need at least 10 journal trades on 8 different days', out)
        self.assertIn('you have 6 trades on 6 days', out)
        self.assertEqual(self.child_calls, [])
        self.rows = [t for t in self.trades(14)] + [{'trade_date': '2026-09-01', 'pnl': None}]
        self.assertIn('JOURNAL REPORT', self.say('tearsheet'))

    def test_several_trades_on_one_day_are_added_up(self):
        self.rows = [{'trade_date': '2026-09-%02d' % (1 + i // 2), 'pnl': 100.0} for i in range(20)]
        self.say('tearsheet capital=1L')
        self.assertEqual(len(self.qs_input['dates']), 10)
        self.assertAlmostEqual(self.qs_input['values'][0], 200.0 / 100000)


# ===================================================================================================================
# 7. NSE DATA (not verified against the live site)
# ===================================================================================================================
class TestNse(WireCase):
    def setUp(self):
        super().setUp()
        self.child_result['NSE'] = {'columns': ['DATE', 'OPEN', 'HIGH', 'LOW', 'CLOSE', 'VOLUME', 'X'], 'rows': [['2026-09-01', '24400', '24500', '24350', '24480', '1000', 'a'], ['2026-09-02', '24480', '24600', '24400', '24590', '1200', 'b']]}

    def test_a_stock_history_comes_as_a_csv_with_the_first_rows_shown(self):
        out = self.say('nse history RELIANCE from 2026-09-01 to 2026-09-30')
        self.assertIn('Asking NSE for the stock RELIANCE from 2026-09-01 to 2026-09-30', out)
        self.assertIn('not verified against it', out)
        self.assertIn('2 rows from NSE', out)
        self.assertIn('2026-09-01 | 24400 | 24500 | 24350 | 24480 | 1000', out)
        self.assertIn('Compare with your broker before acting on it', out)
        self.assertEqual(self.child_calls[0]['args'], ['stock', 'RELIANCE', '2026-09-01', '2026-09-30'])
        self.assertEqual(self.docs[0]['name'], 'nse_RELIANCE.csv')
        self.assertTrue(self.docs[0]['bytes'].decode().startswith('DATE,OPEN,HIGH,LOW,CLOSE,VOLUME,X'))
        self.assertEqual(self.passed, [])
        self.no_job_folders_left()

    def test_an_index_and_futures_with_the_right_arguments(self):
        self.say('nse index NIFTY 50 from 01-09-2026 to 30/09/2026')
        self.assertEqual(self.child_calls[-1]['args'], ['index', 'NIFTY 50', '2026-09-01', '2026-09-30'])
        self.say('nse futures NIFTY expiry 2026-10-28 from 2026-09-01 to 2026-09-30')
        self.assertEqual(self.child_calls[-1]['args'], ['futures', 'NIFTY', '2026-09-01', '2026-09-30', '2026-10-28', 'FUTIDX'])
        self.say('nse futures RELIANCE expiry 2026-10-28 from 2026-09-01 to 2026-09-30')
        self.assertEqual(self.child_calls[-1]['args'][-1], 'FUTSTK')

    def test_last_n_days_and_missing_pieces(self):
        self.say('nse history TCS last 10 days')
        a = self.child_calls[-1]['args']
        self.assertEqual((a[0], a[1]), ('stock', 'TCS'))
        self.assertEqual((self.m._n91_dt.date.fromisoformat(a[3]) - self.m._n91_dt.date.fromisoformat(a[2])).days, 10)
        self.child_calls.clear()
        self.assertIn('For futures also give the expiry date', self.say('nse futures NIFTY from 2026-09-01 to 2026-09-30'))
        self.assertIn('Give a start before the end, within one year', self.say('nse history TCS from 2026-09-30 to 2026-09-01'))
        self.assertIn('Give a start before the end, within one year', self.say('nse history TCS from 2024-01-01 to 2026-09-01'))
        self.assertEqual(self.child_calls, [])

    def test_no_rows_and_a_failing_site_are_explained(self):
        self.child_result['NSE'] = {'columns': ['DATE'], 'rows': []}
        self.assertIn('NSE returned no rows for that', self.say('nse history TCS from 2026-09-01 to 2026-09-30'))
        self.child_result['NSE'] = self.m._N92Err('failed', 'The helper stopped with an error (ConnectionError: blocked).')
        self.assertIn('The helper stopped with an error', self.say('nse history TCS from 2026-09-01 to 2026-09-30'))

    def test_dates_in_the_usual_forms(self):
        d = self.m._n92_date
        want = self.m._n91_dt.date(2026, 10, 4)
        for text in ('2026-10-04', '04-10-2026', '04/10/2026', '04 Oct 2026', '04 October 2026'):
            self.assertEqual(d(text), want, text)
        self.assertIsNone(d('tomorrow'))
        self.assertEqual(d('x', 'fallback'), 'fallback')


# ===================================================================================================================
# 8. VOICE NOTES WHEN GROQ FAILS
# ===================================================================================================================
class TestLocalVoice(WireCase):
    def setUp(self):
        super().setUp()
        self.voice = os.path.join(self.tmp.name, 'v.ogg')
        open(self.voice, 'wb').write(b'OggS fake')
        self.groq = {'text': None}
        self.start(self.m, '_N92_STT_PREV', lambda path, langs=('en', 'hi'): self.groq['text'])
        self.start(self.m, '_n92_have', lambda mod: mod == 'faster_whisper')
        self.m._N92_STT_NOTICE['ts'] = 0.0

    def test_groq_answers_so_nothing_local_runs(self):
        self.groq['text'] = 'hello from groq'
        self.assertEqual(self.m.groq_transcribe(self.voice), 'hello from groq')
        self.assertEqual(self.child_calls, [])

    def test_groq_fails_so_the_server_transcribes_and_the_owner_is_told_once(self):
        self.child_result['STT'] = {'text': 'नमस्ते nemo', 'lang': 'hi'}
        self.assertEqual(self.m.groq_transcribe(self.voice), 'नमस्ते nemo')
        call = self.child_calls[0]
        self.assertEqual(call['args'][0], self.voice)
        self.assertEqual(call['args'][2], 'base')
        self.assertTrue(call['args'][1].endswith('models'))
        self.assertEqual((call['need'], call['timeout']), (800, 240))
        notes = [t for c, t in self.sent if c == OWNER_ID and 'did it on the server' in t]
        self.assertEqual(len(notes), 1)
        self.m.groq_transcribe(self.voice)
        self.assertEqual(len([t for c, t in self.sent if 'did it on the server' in t]), 1, 'one notice per ten minutes')
        self.assertEqual(self.m._N92_STATS['stt_local'], 2)

    def test_the_model_size_follows_the_setting_and_stays_inside_the_known_ones(self):
        self.child_result['STT'] = {'text': 'x'}
        for setting, size, need in (('tiny', 'tiny', 500), ('small', 'small', 1500), ('huge', 'base', 800), ('', 'base', 800)):
            self.store['wire_stt_model'] = setting
            self.child_calls.clear()
            self.m.groq_transcribe(self.voice)
            self.assertEqual((self.child_calls[0]['args'][2], self.child_calls[0]['need']), (size, need), setting)

    def test_every_way_of_not_being_able_returns_nothing_instead_of_an_error(self):
        self.start(self.m, '_n92_have', lambda mod: False)
        self.assertIsNone(self.m.groq_transcribe(self.voice))
        self.start(self.m, '_n92_have', lambda mod: True)
        self.assertIsNone(self.m.groq_transcribe(os.path.join(self.tmp.name, 'missing.ogg')))
        for err in (self.m._N92Err('low_memory', 'x'), self.m._N92Err('busy', 'x'), self.m._N92Err('model_download', 'x'), RuntimeError('boom')):
            self.child_result['STT'] = err
            self.assertIsNone(self.m.groq_transcribe(self.voice), repr(err))
        self.child_result['STT'] = {'text': '  '}
        self.assertIsNone(self.m.groq_transcribe(self.voice))
        self.assertEqual(self.m._N92_STATS['stt_local'], 0)

    def test_the_wrapper_is_what_every_voice_path_calls(self):
        src = open(base.NEMO_FILE, encoding='utf-8').read()
        last = src.rindex('def groq_transcribe(')
        self.assertGreater(last, src.index('# NEMO 92 - WIRE'), 'the Wire wrapper is the last definition, so every caller gets the fallback')

    def test_the_older_regression_row_about_the_primary_model_still_sees_through_the_wrapper(self):
        self.assertTrue(callable(self.m.groq_transcribe.__wrapped__))
        self.assertIn('whisper-large-v3', str(self.m.groq_transcribe.__wrapped__.__code__.co_consts))
        suite = self.m.prime_regression_suite()
        row = [t for t in suite['tests'] if t.get('name') == 'v167-stt-large-v3-primary']
        self.assertTrue(row and row[0].get('ok'), row)


# ===================================================================================================================
# 9. EXTRA CHECKS NEXT TO THE UPDATE PRE-FLIGHT
# ===================================================================================================================
class TestUpdateChecks(WireCase):
    def install_fake_tools(self, which=('ruff', 'vulture', 'detect-secrets'), new_is_worse=True):
        scripts = {
            'ruff': "#!/bin/sh\ncase \"$*\" in *new.py*) printf '10\\tF401\\t[-] unused-import\\n%d\\tF821\\t[ ] undefined-name\\n'; exit 1;; *) printf '10\\tF401\\t[-] unused-import\\n13\\tF821\\t[ ] undefined-name\\n'; exit 1;; esac\\n" % (14 if new_is_worse else 13),
            'vulture': "#!/bin/sh\ncase \"$*\" in *new.py*) printf \"new.py:5: unsatisfiable 'if' condition (100%% confidence)\\nnew.py:9: unused import 'zlib' (100%% confidence)\\n\"; exit 3;; *) printf \"old.py:5: unsatisfiable 'if' condition (100%% confidence)\\n\"; exit 3;; esac\n",
            'detect-secrets': "#!/bin/sh\ncase \"$*\" in *new.py*) echo '{\"results\": {\"new.py\": [{\"type\": \"Basic Auth Credentials\", \"hashed_secret\": \"aaa\", \"line_number\": 1}, {\"type\": \"Secret Keyword\", \"hashed_secret\": \"bbb\", \"line_number\": 7}]}}'; exit 0;; *) echo '{\"results\": {\"old.py\": [{\"type\": \"Basic Auth Credentials\", \"hashed_secret\": \"aaa\", \"line_number\": 1}]}}'; exit 0;; esac\n",
        }
        for name in which:
            self.m._n91_ledger_add('pypi', name, '1.0', 'tool', name, {'scripts': [name], 'added': [name + '==1.0'], 'created_env': True})
            d = os.path.join(self.envs_dir, name, 'bin')
            os.makedirs(d)
            p = os.path.join(d, name)
            open(p, 'w').write(scripts[name].replace('\\n', '\n') if name == 'ruff' else scripts[name])
            os.chmod(p, 0o755)
        self.start(self.m, '_self_read', lambda: 'old code')

    def test_the_new_file_is_compared_with_the_running_one_tool_by_tool(self):
        self.install_fake_tools()
        lines = self.m._n92_update_checks('x' * 100001)
        text = '\n'.join(lines)
        self.assertIn('EXTRA CHECKS ON THE NEW FILE (information only; the pre-flight above decides)', text)
        self.assertIn('ruff (mistakes and unused code): undefined names 13 → 14 ⚠️ more than before: look at them · all findings 23 → 24 (+1)', text)
        self.assertIn("vulture (certain dead code): 1 → 2 ⚠️ · new kind: unused import 'zlib' (100% confidence)", text)
        self.assertIn('detect-secrets: 1 new possible secret ⚠️ look before you apply (known before: 1)', text)
        self.no_job_folders_left()
        self.assertEqual(self.m._N92_STATS['checks'], 1)

    def test_an_equal_file_gets_ticks(self):
        self.install_fake_tools(new_is_worse=False)
        text = '\n'.join(self.m._n92_update_checks('x' * 100001))
        self.assertIn('undefined names 13 → 13 ✅', text)

    def test_tools_that_are_not_installed_are_named_and_no_tools_means_silence(self):
        self.install_fake_tools(which=('ruff',))
        text = '\n'.join(self.m._n92_update_checks('x' * 100001))
        self.assertIn('vulture is not installed (say “install vulture”)', text)
        self.assertIn('detect-secrets is not installed (say “install detect-secrets”)', text)
        shutil.rmtree(self.envs_dir)
        self.assertEqual(self.m._n92_update_checks('x' * 100001), [])

    def test_a_tool_that_fails_is_reported_not_trusted(self):
        self.install_fake_tools(which=('ruff',))
        open(os.path.join(self.envs_dir, 'ruff', 'bin', 'ruff'), 'w').write('#!/bin/sh\nexit 2\n')
        self.assertIn('ruff: could not run', '\n'.join(self.m._n92_update_checks('x' * 100001)))

    def test_the_wrapper_sends_the_checks_first_and_then_runs_the_normal_update_exactly_once(self):
        self.install_fake_tools()
        calls = []
        self.start(self.m, '_N92_SELF_UPDATE_PREV', lambda chat_id: calls.append(('prev', len(self.sent))))
        self.m.LAST_CODE[self.cid] = {'name': 'x.py', 'code': 'x' * 100001}
        n = len(self.sent)
        self.m.self_update(self.cid)
        self.assertEqual(len(calls), 1)
        self.assertGreater(calls[0][1], n, 'the extra-checks message went out before the normal update flow')
        self.assertIn('EXTRA CHECKS', self.texts_to_owner(n))

    def test_nothing_changes_for_other_chats_small_files_or_a_crash(self):
        self.install_fake_tools()
        calls = []
        self.start(self.m, '_N92_SELF_UPDATE_PREV', lambda chat_id: calls.append(chat_id))
        self.m.LAST_CODE[self.cid] = {'name': 'x.py', 'code': 'short'}
        n = len(self.sent)
        self.m.self_update(self.cid)
        self.assertEqual((calls, len(self.sent)), ([self.cid], n))
        self.m.LAST_CODE[self.cid] = {'name': 'x.py', 'code': 'x' * 100001}
        self.m.self_update(5552)
        self.assertEqual(len(self.sent), n, 'only the owner gets the extra message')
        self.start(self.m, '_n92_update_checks', lambda code: (_ for _ in ()).throw(RuntimeError('boom')))
        self.m.self_update(self.cid)
        self.assertEqual(len(calls), 3, 'a crash in the extra checks never blocks the update flow')
        self.assertEqual(self.m._N92_STATS['errors'], 1)

    def test_tools_are_found_only_when_forge_installed_them_as_isolated_tools(self):
        self.assertEqual(self.m._n92_tool('ruff'), '')
        self.install_fake_tools(which=('ruff',))
        self.assertTrue(self.m._n92_tool('ruff').endswith('/ruff/bin/ruff'))
        self.m._n91_ledger_add('pypi', 'vulture', '1', 'runtime', '', {'scripts': ['vulture']})
        self.assertEqual(self.m._n92_tool('vulture'), '', 'a library added to my own Python is not an isolated tool')


# ===================================================================================================================
# 10. WHAT IS READY, THE FRONT DOOR, MCP, WRAPPERS
# ===================================================================================================================
class TestStatusAndFrontDoor(WireCase):
    def test_wire_lists_what_is_ready_and_what_to_say_to_get_the_rest(self):
        self.start(self.m, '_n92_have', lambda mod: mod in ('pdfplumber', 'trafilatura'))
        self.start(self.m._n91_shutil, 'which', lambda name: None)
        out = self.say('wire')
        for part in ('🔌 WIRE: what the installed libraries do for you', '✅ PDF tables → spreadsheet — “tables from this pdf”', '✅ Clean text of a web page — “read article <link>”',
                     '➖ Text in a picture (English, Hindi) — not ready: forge apt tesseract-ocr tesseract-ocr-eng tesseract-ocr-hin + install pytesseract into yourself',
                     '➖ Indicators on daily candles — not ready: install pandas-ta-classic into yourself', '➖ Voice notes on the server when Groq fails — not ready: install faster-whisper into yourself',
                     '➖ Extra update checks (ruff, vulture, detect-secrets): ruff not installed', '🖥 Server: memory', 'This run: 0 heavy jobs'):
            self.assertIn(part, out)
        for text in ('/wire', 'wire status', 'what can you do with the installed libraries?', 'what did you wire'):
            self.assertIn('WIRE: what the installed libraries do', self.say(text), text)
        self.assertEqual(self.passed, [])

    def test_only_the_owner_in_a_private_chat_can_use_any_of_it(self):
        self.child_result['PDF'] = {'pages': 1, 'text_chars': 100, 'tables': [{'page': 1, 'rows': [['a']]}]}
        self.child_result['IND'] = TestIndicators.IND
        guest = lambda **k: dict({'chat': {'id': 5552, 'type': 'private'}, 'from': {'id': 5552, 'first_name': 'Asha'}, 'message_id': 3}, **k)
        for kw in ({'text': 'wire'}, {'text': 'indicators NIFTY'}, {'text': 'tearsheet'}, {'text': 'read article https://example.com/a'}, {'text': 'nse history TCS last 5 days'},
                   {'caption': 'tables from this pdf', 'document': {'file_id': 'f', 'file_name': 'a.pdf', 'mime_type': 'application/pdf'}}, {'caption': 'read this picture', 'photo': [{'file_id': 'p'}]}):
            n = len(self.passed)
            self.m.handle(guest(**kw))
            self.assertEqual(len(self.passed), n + 1, kw)
        group = {'chat': {'id': -100, 'type': 'group'}, 'from': {'id': OWNER_ID}, 'text': 'wire', 'message_id': 1}
        n = len(self.passed)
        self.m.handle(group)
        self.assertEqual(len(self.passed), n + 1)
        self.assertEqual(self.child_calls, [])
        self.assertEqual([t for c, t in self.sent if c == 5552], [])

    def test_slash_commands_other_than_wire_are_never_taken(self):
        for text in ('/tables', '/read https://example.com', '/indicators NIFTY', '/tearsheet'):
            n = len(self.passed)
            self.m.handle(self.msg(text))
            self.assertEqual(len(self.passed), n + 1, text)

    def test_an_unexpected_error_in_the_front_door_never_blocks_normal_chat(self):
        self.start(self.m, '_n92_route', lambda cid, msg, text: (_ for _ in ()).throw(RuntimeError('boom')))
        n = len(self.passed)
        self.m.handle(self.msg('hello there'))
        self.assertEqual(len(self.passed), n + 1)
        self.assertEqual(self.m._N92_STATS['errors'], 1)

    def test_a_job_that_fails_unexpectedly_becomes_one_honest_sentence(self):
        self.child_result['IND'] = RuntimeError('boom')
        self.start(self.m, 'get_history', lambda t: {'closes': [1.0] * 60, 'highs': [1.0] * 60, 'lows': [1.0] * 60})
        out = self.say('indicators NIFTY')
        self.assertIn('Forge hit an unexpected problem (RuntimeError)', out)
        self.assertIn('Nothing was installed or changed', out)

    def test_wrappers_capabilities_status_abilities_and_the_menu(self):
        caps = self.m._n82_capabilities()
        for part in ('Wire 92:', '“tables from this pdf”', '“read this picture”', '“read article <link>”', '“indicators NIFTY”', '“tearsheet”', 'only when the server has memory'):
            self.assertIn(part, caps)
        self.assertIn('WIRE 92: heavy jobs 0', self.m._n83_status_text(self.cid))
        rows = [r for r in self.m._n88_abilities(self.cid) if r[1] == 'Wired libraries']
        self.assertEqual(len(rows), 1)
        self.assertIn('wire', [c[0] for c in self.m._N40_COMMANDS])
        self.assertIn('c:/wire', str(self.m._N40_MENUS['main']))

    def test_the_regression_rows_all_pass(self):
        res = self.m.prime_regression_suite()
        rows = [t for t in res['tests'] if t['name'].startswith('v92-')]
        self.assertEqual(len(rows), 8)
        self.assertEqual([t['name'] for t in rows if not t['ok']], [])

    def test_the_boot_creates_the_work_area_and_sweeps_old_folders(self):
        job = self.m._n92_new_job()
        os.utime(job, (time.time() - 3 * 86400,) * 2)
        self.start(self.m, '_N92_MAIN_PREV', lambda: 'ran')
        self.assertEqual(self.m.main(), 'ran')
        self.assertFalse(os.path.exists(job))


class TestMcpStopsWithItsTree(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if base.m is None:
            base.setUpModule()
        cls.m = base.m

    def make_client(self, pid=4321):
        c = self.m.MCPClient('fs', {'transport': 'stdio'})
        c.proc = mock.Mock(pid=pid)
        return c

    def test_stop_ends_the_whole_server_tree_not_only_the_wrapper(self):
        import signal
        c = self.make_client()
        with mock.patch.object(self.m.os, 'getpgid', return_value=4321), mock.patch.object(self.m.os, 'getpgrp', return_value=1000), mock.patch.object(self.m.os, 'killpg') as killpg:
            c.stop()
        killpg.assert_called_once_with(4321, signal.SIGTERM)
        c.proc.terminate.assert_not_called()
        self.assertFalse(c.alive)

    def test_it_never_signals_the_group_nemo_himself_is_in(self):
        c = self.make_client()
        with mock.patch.object(self.m.os, 'getpgid', return_value=1000), mock.patch.object(self.m.os, 'getpgrp', return_value=1000), mock.patch.object(self.m.os, 'killpg') as killpg:
            c.stop()
        killpg.assert_not_called()
        c.proc.terminate.assert_called_once()

    def test_if_the_group_cannot_be_read_it_falls_back_to_the_plain_stop(self):
        c = self.make_client()
        with mock.patch.object(self.m.os, 'getpgid', side_effect=ProcessLookupError()):
            c.stop()
        c.proc.terminate.assert_called_once()

    def test_a_new_server_gets_its_own_session_so_it_has_its_own_group(self):
        import inspect
        src = inspect.getsource(self.m.MCPClient.start)
        self.assertIn('start_new_session=True', src)
        self.assertEqual(src.count('subprocess.Popen('), src.count('start_new_session=True') + src.count('# no new session') or src.count('subprocess.Popen('))

    def test_every_connected_server_is_stopped_when_nemo_exits(self):
        a, b = mock.Mock(), mock.Mock()
        b.stop.side_effect = RuntimeError('already gone')
        with mock.patch.dict(self.m.MCP_CLIENTS, {'a': a, 'b': b}, clear=True):
            self.m._n92_stop_mcp()
        a.stop.assert_called_once()
        b.stop.assert_called_once()
        self.assertIn('_n92_atexit.register(_n92_stop_mcp)', wire_source())


# ===================================================================================================================
# 11. STRUCTURE
# ===================================================================================================================
class TestStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if base.m is None:
            base.setUpModule()
        cls.m = base.m
        cls.src = open(base.NEMO_FILE, encoding='utf-8').read()
        cls.layer = wire_source()
        cls.tree = ast.parse(cls.layer)

    def calls(self):
        out = []
        for node in ast.walk(self.tree):
            if isinstance(node, ast.Call):
                f = node.func
                out.append(('%s.%s' % (getattr(f.value, 'id', '?'), f.attr)) if isinstance(f, ast.Attribute) else getattr(f, 'id', '?'))
        return out

    def test_the_version_and_the_documentation(self):
        self.assertGreaterEqual(float(self.m.VERSION), 92.0)
        self.assertEqual(re.findall(r'^VERSION\s*=\s*["\']([^"\']+)["\']', self.src, re.M)[-1], self.m.VERSION)
        self.assertIn('v92.0 - WIRE', self.src[:12000])
        self.assertIn('+ v91.5 - FORGE', self.src[:12000])
        self.assertTrue(self.src.index('# NEMO 91 - FORGE') < self.src.index('# NEMO 92 - WIRE'))
        self.assertTrue(self.src.rindex("if __name__") > self.src.index('# NEMO 92 - WIRE'))

    def test_the_new_layer_is_protected_from_live_self_editing(self):
        for name in ('_n92_child', '_n92_guard', '_n92_front', '_n92_update_checks'):
            self.assertFalse(self.m._n79_editable(name), name)

    def test_the_hooks_are_the_outermost_ones(self):
        for name in ('handle', 'main', 'self_update', 'groq_transcribe', 'prime_regression_suite', '_n82_capabilities', '_n83_status_text', '_n88_abilities'):
            self.assertGreater(self.src.rindex('def %s(' % name), self.src.index('# NEMO 92 - WIRE'), name)
        self.assertIsNot(self.m.handle, self.m._N92_HANDLE_PREV)

    def test_no_program_is_started_here_and_there_is_no_shell(self):
        calls = self.calls()
        for bad in ('_n91_sp.run', 'subprocess.run', 'subprocess.Popen', 'os.system', '_n91_os.system', '_n91_os.popen', 'eval', 'exec', 'compile'):
            self.assertNotIn(bad, calls, bad)
        for node in ast.walk(self.tree):
            if isinstance(node, ast.keyword) and node.arg == 'shell':
                self.fail('shell= is never used')
        self.assertIn('_n91_run', calls, 'everything starts through the one place Forge already has')

    def test_the_layer_opens_no_address_itself_and_never_touches_the_broker_or_the_trading_guards(self):
        calls = self.calls()
        self.assertEqual([c for c in calls if c.startswith('requests.')], [])
        for bad in ('place_order', 'fyers_data', 'fyers_history', 'BROKER', 'can_enter', 'must_square_off', 'HOLIDAYS', 'client_id', 'access_token', 'BOT_TOKEN', 'save_secret(', 'OWNER =', 'OWNER.update', 'OWNER.pop'):
            self.assertNotIn(bad, self.layer, bad)
        self.assertEqual(re.findall(r'OWNER\[[^\]]*\]\s*=[^=]', self.layer), [], 'the owner lock is only read here, never written')
        self.assertIn('_n85_fetch(', self.layer, 'web pages are opened by the existing safe fetcher')

    def test_the_helper_programs_are_small_valid_and_print_exactly_one_json_line(self):
        names = [n for n in dir(self.m) if n.startswith('_N92_CODE_')]
        self.assertEqual(sorted(names), ['_N92_CODE_ARTICLE', '_N92_CODE_IND', '_N92_CODE_NSE', '_N92_CODE_OCR', '_N92_CODE_PDF', '_N92_CODE_QS', '_N92_CODE_STT'])
        for n in names:
            code = getattr(self.m, n)
            ast.parse(code)
            self.assertEqual(code.count('print(json.dumps('), 1, n)
            for bad in ('os.environ', 'requests', 'subprocess', 'open(sys.argv[1], \'w\'', 'eval(', 'exec('):
                self.assertNotIn(bad, code, (n, bad))
            self.assertLess(len(code), 3500, n)

    def test_tokens_and_keys_are_never_logged_or_sent_to_a_helper(self):
        for call in re.findall(r'(?:act_log|v11_audit|_n68_audit)\([^\n]*', self.layer):
            self.assertNotIn('token', call.lower(), call)
        self.assertNotIn('print', self.calls())
        self.assertIn('_n91_env()', self.layer, 'helpers get Forge\'s minimal environment')

    def test_job_folders_are_only_deleted_through_the_contained_helper(self):
        for m_ in re.finditer(r'rmtree\(([^,)]+)', self.layer):
            self.assertEqual(m_.group(1).strip(), 'p', m_.group(0))
        self.assertIn("basename(p).startswith('J92-')", self.layer)


# ===================================================================================================================
# 12. THE REAL LIBRARIES (only when NEMO_WIRE_PY points at a Python that has them)
# ===================================================================================================================
REAL_PY = os.environ.get('NEMO_WIRE_PY', '')


def py_has(mod):
    if not REAL_PY:
        return False
    return subprocess.run([REAL_PY, '-c', 'import ' + mod], capture_output=True).returncode == 0


@unittest.skipUnless(REAL_PY, 'set NEMO_WIRE_PY to a Python that has the libraries to run the real-library tests')
class TestRealLibraries(ForgeCase):
    def setUp(self):
        super().setUp()
        self.start(self.m, '_n92_python', lambda: REAL_PY)
        self.start(self.m, '_n92_guard', lambda need: (True, ''))
        self.job = self.m._n92_new_job()

    def tearDown(self):
        self.m._n92_rm_job(self.job)
        super().tearDown()

    @unittest.skipUnless(py_has('pdfplumber'), 'pdfplumber is not in that Python')
    def test_pdfplumber_returns_the_table_cell_for_cell(self):
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font('Helvetica', size=11)
        for row in (('Contract', 'Qty', 'Price'), ('NIFTY OCT FUT', '75', '24,590.50')):
            for w, t in zip((50, 25, 30), row):
                pdf.cell(w, 8, t, border=1)
            pdf.ln(8)
        path = os.path.join(self.job, 's.pdf')
        pdf.output(path)
        res = self.m._n92_child(self.m._N92_CODE_PDF, [path, 5], need_mb=1, job=self.job)
        self.assertEqual(res['pages'], 1)
        self.assertEqual(res['tables'][0]['rows'], [['Contract', 'Qty', 'Price'], ['NIFTY OCT FUT', '75', '24,590.50']])

    @unittest.skipUnless(py_has('pytesseract') and shutil.which('tesseract'), 'pytesseract or tesseract is not available')
    def test_tesseract_reads_a_picture(self):
        from PIL import Image, ImageDraw, ImageFont
        im = Image.new('RGB', (900, 120), 'white')
        ImageDraw.Draw(im).text((20, 30), 'NIFTY OCT FUT Qty 75 Price 24590.50', fill='black', font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 36))
        path = os.path.join(self.job, 'p.png')
        im.save(path)
        res = self.m._n92_child(self.m._N92_CODE_OCR, [path, 'eng'], need_mb=1, job=self.job)
        self.assertIn('NIFTY OCT FUT Qty 75 Price 24590.50', res['text'])

    @unittest.skipUnless(py_has('trafilatura'), 'trafilatura is not in that Python')
    def test_trafilatura_drops_the_menu_and_keeps_the_article(self):
        html = '<html><head><title>Oil</title></head><body><nav>Home | Markets | Login</nav><article><h1>Oil jumps</h1><p>' + 'Brent rose 4% on Friday after a pipeline outage. ' * 8 + '</p></article><footer>Cookie policy</footer></body></html>'
        path = os.path.join(self.job, 'p.html')
        open(path, 'w').write(html)
        res = self.m._n92_child(self.m._N92_CODE_ARTICLE, [path], need_mb=1, job=self.job)
        self.assertIn('Brent rose 4% on Friday', res['text'])
        self.assertNotIn('Login', res['text'])

    @unittest.skipUnless(py_has('pandas_ta_classic'), 'pandas-ta-classic is not in that Python')
    def test_the_indicators_come_out_as_numbers(self):
        import random
        rnd = random.Random(1)
        closes = [24500.0]
        for _ in range(99):
            closes.append(closes[-1] + rnd.gauss(0, 30))
        path = os.path.join(self.job, 'c.json')
        json.dump({'closes': closes, 'highs': [c + 20 for c in closes], 'lows': [c - 20 for c in closes]}, open(path, 'w'))
        res = self.m._n92_child(self.m._N92_CODE_IND, [path], need_mb=1, job=self.job)
        for k in ('rsi', 'atr', 'adx', 'st_line', 'st_dir', 'bb_low', 'bb_up', 'ema20', 'ema50', 'macd_hist'):
            self.assertIsInstance(res[k], (int, float), k)
        self.assertTrue(0 <= res['rsi'] <= 100)
        self.assertIn('RSI(14)', self.m._n92_ind_text('X', res))

    @unittest.skipUnless(py_has('quantstats'), 'quantstats is not in that Python')
    def test_quantstats_writes_a_report(self):
        path = os.path.join(self.job, 'in.json')
        days = ['2026-09-%02d' % d for d in range(1, 21)]
        json.dump({'dates': days, 'values': [0.002 * ((-1) ** i) + 0.0005 for i in range(20)], 'title': 'T'}, open(path, 'w'))
        out = os.path.join(self.job, 'r.html')
        res = self.m._n92_child(self.m._N92_CODE_QS, [path, out], need_mb=1, job=self.job)
        self.assertEqual(res['days'], 20)
        self.assertGreater(os.path.getsize(out), 50000)
        self.assertIsInstance(res['max_drawdown'], float)


if __name__ == '__main__':
    unittest.main()
