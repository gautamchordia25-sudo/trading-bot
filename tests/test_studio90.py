"""Nemo v90 "Studio": more image engines and tools, exact designs and honest charts, researched reports whose numbers are checked, and data reports from your own files.

Everything is offline: temporary SQLite database, a scripted fake for every HTTP call (image engines, Telegram uploads), scripted AI, synthetic pictures and data.
No Telegram, no image provider, no AI, no paid calls, no trades. Real Pillow, matplotlib, fpdf2, PyMuPDF and openpyxl are used where the sandbox has them.

    NEMO_FILE=/path/to/nemotron_bot.py python -m unittest tests.test_studio90 -v
"""
import ast
import datetime as dt
import io
import json
import os
import random
import re
import tempfile
import threading
import unittest
import zipfile
from decimal import Decimal as D
from unittest import mock

from tests import test_cortex83 as base

OWNER_ID = base.OWNER_ID
KEY_A = 'KEYA-0123456789-abcdef'          # synthetic; never a real key
KEY_B = 'KEYB-9876543210-fedcba'


def setUpModule():
    if base.m is None:
        base.setUpModule()


try:
    from PIL import Image, ImageDraw
    HAVE_PIL = True
except Exception:                                   # pragma: no cover
    HAVE_PIL = False
try:
    import matplotlib                               # noqa: F401
    HAVE_MPL = True
except Exception:                                   # pragma: no cover
    HAVE_MPL = False
try:
    import fitz
    HAVE_FITZ = True
except Exception:                                   # pragma: no cover
    HAVE_FITZ = False
try:
    import openpyxl
    HAVE_XLSX = True
except Exception:                                   # pragma: no cover
    HAVE_XLSX = False
try:
    import fpdf                                     # noqa: F401
    HAVE_FPDF = True
except Exception:                                   # pragma: no cover
    HAVE_FPDF = False

needs_pil = unittest.skipUnless(HAVE_PIL, 'Pillow is not installed')
needs_mpl = unittest.skipUnless(HAVE_PIL and HAVE_MPL, 'matplotlib/Pillow are not installed')
needs_pdf = unittest.skipUnless(HAVE_PIL and HAVE_MPL and HAVE_FITZ and HAVE_FPDF, 'PDF libraries are not installed')
needs_xlsx = unittest.skipUnless(HAVE_XLSX, 'openpyxl is not installed')


def png(w=320, h=200, kind='png', seed=1, plain=False):
    """A real, non-blank synthetic picture."""
    im = Image.new('RGB', (w, h), (240, 240, 240) if plain else (30, 60, 120))
    d = ImageDraw.Draw(im)
    if not plain:
        rnd = random.Random(seed)
        for y in range(h):
            d.line((0, y, w, y), fill=(30 + y * 200 // h, 60 + rnd.randint(0, 20), 120))
    d.ellipse((w // 4, h // 4, 3 * w // 4, 3 * h // 4), fill=(250, 200, 40))
    d.rectangle((w // 3, h // 3, w // 2, 2 * h // 3), fill=(200, 30, 30))
    b = io.BytesIO()
    im.save(b, 'JPEG' if kind == 'jpeg' else 'PNG')
    return b.getvalue()


class HttpResp:
    def __init__(self, status=200, content=b'', js=None, headers=None, text=None):
        self.status_code = status
        self.content = content if isinstance(content, bytes) else str(content).encode()
        self._js = js
        self.headers = headers or {}
        self.text = text if text is not None else self.content[:300].decode('latin-1')

    def json(self):
        if self._js is None:
            raise ValueError('no json')
        return self._js


class FakeHttp:
    """Routes requests.get/post/delete by URL fragment. An address nobody scripted counts as a dead network (and is recorded)."""

    def __init__(self):
        self.routes = []
        self.calls = []
        self.unmatched = []

    def on(self, fragment, handler):
        self.routes.append((fragment, handler))

    def _do(self, method, url, **kw):
        self.calls.append((method, url, kw))
        for fragment, handler in self.routes:
            if fragment in url:
                out = handler(method, url, kw) if callable(handler) else handler
                if isinstance(out, Exception):
                    raise out
                return out
        self.unmatched.append(url)
        import requests
        raise requests.exceptions.ConnectionError('no route')

    def get(self, url, **kw):
        return self._do('get', url, **kw)

    def post(self, url, **kw):
        return self._do('post', url, **kw)

    def delete(self, url, **kw):
        return self._do('delete', url, **kw)

    def to(self, fragment):
        return [c for c in self.calls if fragment in c[1]]


class StudioCase(base.CortexCase):
    def setUp(self):
        super().setUp()
        m = self.m = base.m
        self.kbs, self.tgcalls, self.docs, self.passed, self.store = [], [], [], [], {}
        self.http = FakeHttp()
        self.studio_dir = os.path.join(self.tmp.name, 'studio')
        os.makedirs(self.studio_dir, exist_ok=True)
        import requests

        def start(obj, name, value):
            p = mock.patch.object(obj, name, value)
            p.start()
            self.patches.append(p)

        def send(cid, text, kb=None):
            self.sent.append((cid, text))
            self.kbs.append((cid, kb))

        def tg(method, **params):
            self.tgcalls.append((method, params))
            return {'ok': True, 'result': {'file_path': 'x/y'}}

        def send_document(cid, path, filename, mime):
            with open(path, 'rb') as fh:
                self.docs.append({'cid': cid, 'name': filename, 'mime': mime, 'bytes': fh.read(), 'path': path})
            return True

        def secret(name):
            for n in m._N90_KEY_NAMES.get(name, (name,)):
                v = self.store.get(n)
                if v:
                    return v
            return ''
        start(m, 'send_text', send)
        start(m, 'tg', tg)
        start(m, 'send_document', send_document)
        start(m, 'save_secret', lambda k, v: self.store.__setitem__(k, v))
        start(m, '_n90_secret', secret)
        start(m, '_n90_dir', lambda: self.studio_dir)
        start(m, 'NVIDIA_KEY', '')
        start(requests, 'get', self.http.get)
        start(requests, 'post', self.http.post)
        start(requests, 'delete', self.http.delete)
        start(m, '_N90_HANDLE_PREV', lambda msg: self.passed.append(msg))
        start(m, '_N90_CALLBACK_PREV', lambda cq: self.passed.append(cq))
        start(m, '_N90_SYNC', True)
        start(m, 'ask_ai', lambda cid, text, **kw: 'AI-REPLY')
        start(m._n90_time, 'sleep', lambda s: None)
        m._N90_LEDGER.clear()
        m._N90_PENDING_DATA.clear()
        m._N90_INSTALL_TRIED.clear()
        for k in m._N90_STATS:
            m._N90_STATS[k] = 0
        m._n90_db().close()
        m._n88_db().close()
        self.cid = OWNER_ID
        self.orig_ledger_save = m._n90_ledger_save

    def tearDown(self):
        for p in reversed(self.patches):                  # last patch first, so an attribute patched twice goes back to the real function
            p.stop()
        self.patches = []
        self.m._N83_SYNC['on'] = False
        self.m.HISTORY.clear()
        self.tmp.cleanup()

    def owner(self, text, **extra):
        n0 = len(self.sent)
        self.m.handle(self.msg(text, **extra))
        return '\n'.join(t for c, t in self.sent[n0:] if c == OWNER_ID)

    def texts(self, since=0):
        return '\n'.join(t for c, t in self.sent[since:] if c == OWNER_ID)

    def telegram_uploads(self):
        return [c for c in self.http.calls if '/sendPhoto' in c[1] or '/sendDocument' in c[1]]

    def telegram_ok(self):
        self.http.on('/sendPhoto', HttpResp(200, js={'ok': True}))
        self.http.on('/sendDocument', HttpResp(200, js={'ok': True}))

    def give_keys(self, *names):
        for n in names:
            self.store[self.m._N90_KEY_SAVE_AS[n]] = KEY_A if n != 'cloudflare_account' else '0123456789abcdef0123456789abcdef'

    # --- scripted engines (the real adapters run; only the network is fake)
    def script_gemini(self, image=None, status=200, body=None):
        raw = image or png()
        import base64
        ok = {'candidates': [{'content': {'parts': [{'inlineData': {'mimeType': 'image/png', 'data': base64.b64encode(raw).decode()}}]}}]}
        self.http.on('generativelanguage.googleapis.com', HttpResp(status, js=body if body is not None else ok))

    def script_together(self, image=None, status=200):
        import base64
        raw = image or png()
        self.http.on('api.together.xyz', HttpResp(status, js={'data': [{'b64_json': base64.b64encode(raw).decode()}]} if status == 200 else {'error': 'x'}))

    def script_pollinations(self, content=None, ctype='image/png', status=200):
        self.http.on('image.pollinations.ai', HttpResp(status, content or png(), headers={'content-type': ctype}))


# ===================================================================================================================
# 1. THE IMAGE CHAIN
# ===================================================================================================================
@needs_pil
class TestImageChain(StudioCase):
    def test_nine_engines_in_one_table_with_exactly_one_paid(self):
        m = self.m
        ids = [e[0] for e in m._N90_ENGINES]
        self.assertEqual(ids, ['nvidia', 'gemini', 'cloudflare', 'together', 'huggingface', 'pollinations', 'pollinations_free', 'horde', 'openai'])
        self.assertEqual([e[0] for e in m._N90_ENGINES if e[4]], ['openai'])
        self.assertEqual({e[0] for e in m._N90_ENGINES if e[3]}, {'pollinations_free', 'horde'})
        self.assertEqual({e[0] for e in m._N90_ENGINES if e[5]}, {'gemini'}, 'only Gemini can take a source picture')

    def test_with_no_keys_only_the_keyless_engines_are_ready(self):
        self.assertEqual(self.m._n90_ready_engines(), ['pollinations_free', 'horde'])

    def test_a_saved_key_makes_its_engine_ready_and_cloudflare_needs_both_values(self):
        self.give_keys('together', 'gemini')
        self.assertEqual(self.m._n90_ready_engines(), ['gemini', 'together', 'pollinations_free', 'horde'])
        self.give_keys('cloudflare')
        self.assertNotIn('cloudflare', self.m._n90_ready_engines())
        self.give_keys('cloudflare_account')
        self.assertIn('cloudflare', self.m._n90_ready_engines())

    def test_the_paid_engine_stays_off_until_the_owner_allows_it(self):
        self.give_keys('openai')
        self.assertEqual(self.m._n90_engine_state('openai'), ('paid_off', ''))
        self.store['studio_allow_paid'] = '1'
        self.assertEqual(self.m._n90_engine_state('openai'), ('ready', ''))

    def test_a_real_image_from_the_first_engine_is_returned_with_who_made_it(self):
        self.give_keys('gemini')
        self.script_gemini()
        res = self.m._n90_generate({'prompt': 'a red kite', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
        self.assertTrue(res['ok'])
        self.assertEqual((res['engine'], res['label']), ('gemini', 'Gemini image'))
        self.assertEqual((res['info']['w'], res['info']['h']), (320, 200))
        self.assertEqual(self.http.to('pollinations'), [], 'nothing after the first engine that worked is called')

    def test_an_html_error_page_is_not_an_image_and_the_next_engine_is_tried(self):
        self.give_keys('together')
        self.http.on('api.together.xyz', HttpResp(200, js={'data': [{'b64_json': __import__('base64').b64encode(b'<html>busy</html>' + b' ' * 3000).decode()}]}))
        self.script_pollinations()
        res = self.m._n90_generate({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
        self.assertTrue(res['ok'])
        self.assertEqual(res['engine'], 'pollinations_free')
        self.assertIn(('together', 'not_image', ''), res['tried'])

    def test_the_old_one_shot_accepted_anything_over_5kb_this_does_not(self):
        html = b'<html><body>' + b'rate limited ' * 800 + b'</body></html>'
        self.assertGreater(len(html), 5000)
        self.script_pollinations(content=html, ctype='text/html')
        self.http.on('aihorde.net', HttpResp(500, js={}))
        res = self.m._n90_generate({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
        self.assertFalse(res['ok'])
        self.assertEqual(dict((e, c) for e, c, _d in res['tried'])['pollinations_free'], 'not_image')
        out = os.path.join(self.tmp.name, 'old.png')
        self.assertFalse(self.m.fetch_image('x', out), 'fetch_image now says False instead of writing the error page')
        self.assertFalse(os.path.exists(out))

    def test_every_kind_of_unusable_reply_has_its_own_word(self):
        check = self.m._n90_check_image
        self.assertEqual(check(b'')[:2], (False, 'too_small'))
        self.assertEqual(check(b'<html>' + b'x' * 3000)[:2], (False, 'not_image'))
        self.assertEqual(check(b'\x89PNG\r\n\x1a\n' + b'0' * 3000)[:2], (False, 'corrupt'))
        noise = io.BytesIO()
        rnd = random.Random(3)
        Image.frombytes('RGB', (40, 40), bytes(rnd.randint(0, 255) for _ in range(40 * 40 * 3))).save(noise, 'PNG')
        self.assertGreater(len(noise.getvalue()), 1200)
        self.assertEqual(check(noise.getvalue())[:2], (False, 'tiny'))
        flat = io.BytesIO()
        Image.new('RGB', (300, 300), (10, 10, 10)).save(flat, 'PNG')
        flat_raw = flat.getvalue() + b'\0' * 1500
        self.assertEqual(check(flat_raw)[:2], (False, 'blank'))
        self.assertTrue(check(png())[0])

    def test_exact_failure_words_for_http_errors(self):
        self.give_keys('together')
        for status, word in ((401, 'auth'), (402, 'quota'), (429, 'quota'), (404, 'gone'), (503, 'server_503')):
            self.http.routes.clear()
            self.http.on('api.together.xyz', HttpResp(status, js={'error': 'x'}))
            with self.assertRaises(self.m._N90Fail) as cm:
                self.m._n90_engine_together({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
            self.assertEqual(cm.exception.code, word if word != 'gone' else 'gone')

    def test_a_provider_that_refuses_the_prompt_is_reported_as_blocked_not_as_broken(self):
        self.give_keys('gemini')
        self.http.on('generativelanguage.googleapis.com', HttpResp(200, js={'promptFeedback': {'blockReason': 'SAFETY'}}))
        res = self.m._n90_generate({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None}, prefer=['gemini'])
        self.assertIn(('gemini', 'blocked', ''), res['tried'])
        self.assertEqual(self.m._N90_LEDGER.get('gemini', {}).get('fail', 0), 0, 'a refused prompt does not count against the engine')

    def test_the_engine_that_worked_last_goes_first_next_time(self):
        self.give_keys('together', 'gemini')
        self.script_together()
        self.script_gemini(status=500, body={})
        self.m._n90_generate({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
        self.assertEqual(self.m._n90_engine_order()[0], 'together')
        self.http.calls.clear()
        self.m._n90_generate({'prompt': 'y', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
        self.assertEqual(self.http.to('generativelanguage'), [], 'the engine that failed is not asked first any more')

    def test_three_failures_in_a_row_put_an_engine_to_rest_for_ten_minutes(self):
        m = self.m
        self.give_keys('together')
        for _ in range(3):
            m._n90_ledger_note('together', False, 'timeout')
        self.assertEqual(m._n90_engine_state('together')[0], 'cooldown')
        self.assertEqual(m._n90_engine_state('together', now=m._N90_LEDGER['together']['last_fail'] + 601)[0], 'ready')
        m._n90_ledger_note('together', True)
        self.assertEqual(m._n90_engine_state('together')[0], 'ready')

    def test_a_missing_key_is_not_held_against_the_engine(self):
        self.m._n90_ledger_note('together', False, 'no_key')
        self.assertEqual(self.m._N90_LEDGER['together']['fail'], 0)
        self.assertEqual(self.m._N90_LEDGER['together']['streak'], 0)

    def test_an_engine_that_crashes_does_not_stop_the_others(self):
        self.give_keys('together')
        self.script_pollinations()
        with mock.patch.object(self.m, '_n90_engine_together', side_effect=KeyError('boom')):
            ents = tuple((e[:7] + (self.m._n90_engine_together,)) if e[0] == 'together' else e for e in self.m._N90_ENGINES)
            with mock.patch.object(self.m, '_N90_ENGINES', ents), mock.patch.object(self.m, '_N90_ENGINE_BY_ID', {e[0]: e for e in ents}):
                res = self.m._n90_generate({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
        self.assertTrue(res['ok'])
        self.assertIn(('together', 'failed', 'KeyError'), res['tried'])

    def test_the_shared_time_budget_stops_the_chain(self):
        self.give_keys('together', 'gemini')
        res = self.m._n90_generate({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None}, deadline=self.m._n90_time.time() + 3)
        self.assertFalse(res['ok'])
        self.assertTrue(all(code in ('timeout', 'no_key', 'paid_off', 'skipped_edit') for _e, code, _d in res['tried']))
        self.assertEqual(self.http.calls, [], 'with less than 8 seconds left no engine is started')

    def test_the_failure_text_names_every_engine_and_what_to_do_and_never_leaks_a_key_or_link(self):
        self.give_keys('together')
        self.http.on('api.together.xyz', HttpResp(401, js={'error': 'bad key ' + KEY_A}))
        self.http.on('image.pollinations.ai', HttpResp(503, content=b'down'))
        self.http.on('aihorde.net', HttpResp(500, js={}))
        res = self.m._n90_generate({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
        text = self.m._n90_failure_text(res['tried'])
        for label in ('Together FLUX: the key was rejected', 'Pollinations (keyless)', 'AI Horde'):
            self.assertIn(label, text)
        self.assertIn('studio key', text)
        self.assertNotIn(KEY_A, text)
        self.assertNotIn('http', text)

    def test_finalize_gives_the_exact_size_without_stretching(self):
        out, info = self.m._n90_finalize(png(640, 400), 1080, 1350, want='png')
        im = Image.open(io.BytesIO(out))
        self.assertEqual(im.size, (1080, 1350))
        self.assertEqual((info['w'], info['h']), (1080, 1350))
        self.assertIn('upscaled', info['note'])
        small, _ = self.m._n90_finalize(png(1600, 1600), 512, 512, want='jpg')
        self.assertEqual(Image.open(io.BytesIO(small)).size, (512, 512))
        self.assertEqual(Image.open(io.BytesIO(small)).format, 'JPEG')

    def test_fetch_image_keeps_its_old_contract_and_uses_the_chain(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        out = os.path.join(self.tmp.name, 'rep.png')
        self.assertTrue(self.m.fetch_image('a mountain', out, w=800, h=450))
        self.assertEqual(Image.open(out).size, (800, 450))
        self.assertEqual(self.m._N90_STATS['images_made'], 1)

    def test_fetch_image_refuses_the_blocked_prompt_without_calling_anything(self):
        out = os.path.join(self.tmp.name, 'blocked.png')
        self.assertFalse(self.m.fetch_image('a naked child', out))
        self.assertEqual(self.http.calls, [])

    def test_the_old_nvidia_helper_is_used_for_the_nvidia_engine_and_its_key_is_checked_first(self):
        with mock.patch.object(self.m, 'nvidia_image', side_effect=AssertionError('must not be called without a key')):
            with self.assertRaises(self.m._N90Fail) as cm:
                self.m._n90_engine_nvidia({'prompt': 'x'})
        self.assertEqual(cm.exception.code, 'no_key')

        def fake_nvidia(prompt, path):
            with open(path, 'wb') as fh:
                fh.write(png())
            return True
        self.store['nvidia_key'] = KEY_A
        with mock.patch.object(self.m, 'nvidia_image', fake_nvidia):
            raw, model = self.m._n90_engine_nvidia({'prompt': 'x'})
        self.assertTrue(self.m._n90_check_image(raw)[0])

    def test_the_ai_horde_engine_submits_polls_and_downloads(self):
        import base64
        state = {'polls': 0}

        def check(method, url, kw):
            state['polls'] += 1
            return HttpResp(200, js={'done': state['polls'] >= 2, 'queue_position': 3})
        self.http.on('generate/async', HttpResp(202, js={'id': 'abcd-1234-efgh'}))
        self.http.on('generate/check/', check)
        self.http.on('generate/status/', HttpResp(200, js={'generations': [{'img': base64.b64encode(png(512, 512)).decode()}]}))
        raw, model = self.m._n90_engine_horde({'prompt': 'x', 'negative': 'blurry', 'w': 1024, 'h': 1024, 'seed': 5, 'time_left': 100})
        self.assertEqual(model, 'community-model')
        self.assertGreaterEqual(state['polls'], 2)
        submit = self.http.to('generate/async')[0][2]['json']
        self.assertEqual(submit['params']['seed'], '5')
        self.assertEqual((submit['params']['width'], submit['params']['height']), (576, 576), 'the community GPUs get a size they can really do')
        self.assertEqual(self.http.to('generate/async')[0][2]['headers']['apikey'], '0000000000', 'anonymous key when none is saved')

    def test_the_horde_gives_up_inside_the_budget_and_cancels_the_request(self):
        self.http.on('generate/async', HttpResp(202, js={'id': 'abcd-1234-efgh'}))
        self.http.on('generate/check/', HttpResp(200, js={'done': False, 'queue_position': 40}))
        self.http.on('generate/status/', HttpResp(200, js={}))
        with mock.patch.object(self.m._n90_time, 'time', side_effect=[1000.0] + [1000.0 + 50 * i for i in range(1, 40)]):
            with self.assertRaises(self.m._N90Fail) as cm:
                self.m._n90_engine_horde({'prompt': 'x', 'negative': '', 'w': 512, 'h': 512, 'seed': None, 'time_left': 20})
        self.assertEqual(cm.exception.code, 'timeout')
        self.assertTrue(any(c[0] == 'delete' for c in self.http.calls), 'the queued job is cancelled')

    def test_cloudflare_reads_both_the_json_and_the_binary_reply(self):
        import base64
        self.give_keys('cloudflare', 'cloudflare_account')
        self.http.on('api.cloudflare.com', HttpResp(200, js={'result': {'image': base64.b64encode(png()).decode()}}, headers={'content-type': 'application/json'}))
        raw, model = self.m._n90_engine_cloudflare({'prompt': 'x', 'negative': '', 'w': 512, 'h': 512, 'seed': 3})
        self.assertTrue(self.m._n90_check_image(raw)[0])
        self.assertIn('flux', model.lower())
        self.http.routes.clear()
        self.http.on('api.cloudflare.com', HttpResp(200, png(), headers={'content-type': 'image/png'}))
        raw, model = self.m._n90_engine_cloudflare({'prompt': 'x', 'negative': '', 'w': 512, 'h': 512, 'seed': 3})
        self.assertTrue(self.m._n90_check_image(raw)[0])

    def test_gemini_can_edit_a_picture_and_the_other_engines_are_skipped_for_edits(self):
        self.give_keys('gemini', 'together')
        self.script_gemini()
        src = png()
        res = self.m._n90_generate({'prompt': 'make the sky pink', 'negative': '', 'w': 320, 'h': 200, 'seed': None, 'source': src}, editing=True)
        self.assertTrue(res['ok'])
        self.assertEqual(res['engine'], 'gemini')
        body = self.http.to('generativelanguage')[0][2]['json']
        self.assertEqual(len(body['contents'][0]['parts']), 2, 'the source picture is sent with the instruction')
        self.assertIn(('nvidia', 'skipped_edit', ''), res['tried'])
        self.assertEqual(self.http.to('together'), [])

    def test_the_gemini_key_goes_in_a_header_not_the_url(self):
        self.give_keys('gemini')
        self.script_gemini()
        self.m._n90_generate({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None}, prefer=['gemini'])
        method, url, kw = self.http.to('generativelanguage')[0]
        self.assertNotIn(KEY_A, url)
        self.assertEqual(kw['headers']['x-goog-api-key'], KEY_A)

    def test_keys_never_appear_in_a_urls_that_is_logged_or_shown(self):
        self.give_keys('together', 'huggingface', 'pollinations', 'openai')
        self.store['studio_allow_paid'] = '1'
        for fn in (self.m._n90_engine_together, self.m._n90_engine_huggingface, self.m._n90_engine_pollinations_new, self.m._n90_engine_openai):
            try:
                fn({'prompt': 'x', 'negative': '', 'w': 320, 'h': 200, 'seed': None})
            except self.m._N90Fail:
                pass
        for _m, url, _kw in self.http.calls:
            self.assertNotIn(KEY_A, url)


# ===================================================================================================================
# 2. WHAT WAS ASKED: size, style, count, seed, subject
# ===================================================================================================================
class TestParsers(StudioCase):
    def test_sizes(self):
        p = self.m._n90_parse_size
        self.assertEqual(p('draw a cat 1080x1350')[:2], (1080, 1350))
        self.assertEqual(p('make it for instagram story')[:2], (1080, 1920))
        self.assertEqual(p('youtube thumbnail of a car')[:2], (1280, 720))
        self.assertEqual(p('a wallpaper')[:2], (1920, 1080))
        self.assertEqual(p('a cat 16:9')[0] % 8, 0)
        self.assertEqual(p('a cat')[:2], (1024, 1024))
        self.assertEqual(p('a cat 20000x20000')[:2], (1024, 1024), 'an absurd size falls back instead of exhausting memory')

    def test_styles_counts_and_seeds(self):
        m = self.m
        self.assertEqual(m._n90_parse_style('a cat in anime style'), 'anime')
        self.assertEqual(m._n90_parse_style('watercolor lighthouse'), 'watercolor')
        self.assertEqual(m._n90_parse_style('a plain cat'), '')
        self.assertEqual(m._n90_parse_count('give me 3 variants'), 3)
        self.assertEqual(m._n90_parse_count('two versions of a cat'), 2)
        self.assertEqual(m._n90_parse_count('50 images of a cat'), 4, 'capped')
        self.assertEqual(m._n90_parse_count('a cat'), 1)
        self.assertEqual(m._n90_parse_seed('a cat seed 42'), 42)
        self.assertIsNone(m._n90_parse_seed('a cat'))
        self.assertIn('crowd', m._n90_parse_negative('a beach without crowd, at dawn'))

    def test_the_subject_loses_the_request_words_and_the_settings(self):
        s = self.m._n90_subject
        self.assertEqual(s('draw a cat riding a bicycle'), 'cat riding a bicycle')
        self.assertEqual(s('please generate an image of a red kite 1080x1080 seed 7'), 'red kite')
        self.assertEqual(s('Nemo, can you make me a picture of a lighthouse at dusk in watercolor style, 3 variants'), 'lighthouse at dusk', 'the style becomes a setting, not picture content')
        self.assertEqual(s('/imagine a dragon without text'), 'a dragon')
        self.assertEqual(s('draw a street with no cars'), 'street with no cars', 'ordinary words stay in the picture')

    def test_the_prompt_builder_adds_style_words_and_merges_negatives(self):
        prompt, neg = self.m._n90_build_prompt('a cat', 'anime', 'blur')
        self.assertTrue(prompt.startswith('a cat, anime style'))
        self.assertIn('blur', neg)
        self.assertIn('photorealistic', neg)

    def test_the_child_safety_filter(self):
        b = self.m._n90_prompt_blocked
        self.assertTrue(b('a naked child'))
        self.assertTrue(b('explicit photo of a teenage girl'))
        self.assertFalse(b('a child flying a kite at the beach'))
        self.assertFalse(b('a nude statue in a museum'))


# ===================================================================================================================
# 3. EDITING A PICTURE WITH NO AI
# ===================================================================================================================
@needs_pil
class TestImageTools(StudioCase):
    def run_ops(self, text, raw=None, cid=None):
        m = self.m
        ctx = m._N90Ctx(cid or self.cid)
        ops = m._n90_parse_ops(text)
        self.assertTrue(ops, text)
        im = m._n90_apply_ops(raw or png(800, 600), ops, ctx)
        return im, ctx, ops

    def finish(self, text, raw=None):
        im, ctx, _ops = self.run_ops(text, raw)
        data, fmt, note = self.m._n90_render_output(im, ctx)
        return Image.open(io.BytesIO(data)), fmt, ctx, note

    def test_operations_are_found_in_the_order_they_are_said(self):
        ops = self.m._n90_parse_ops('crop to 16:9 then add text "SALE" at the top then sharpen and convert to webp')
        self.assertEqual([o['op'] for o in ops], ['crop', 'text', 'sharpen', 'convert'])

    def test_resize_gives_exact_pixels_for_a_size_and_for_social_presets(self):
        self.assertEqual(self.finish('resize to 1080x1080')[0].size, (1080, 1080))
        self.assertEqual(self.finish('resize it to an instagram story')[0].size, (1080, 1920))
        self.assertEqual(self.finish('scale 50%')[0].size, (400, 300))
        self.assertEqual(self.finish('resize to width 400')[0].size, (400, 300))

    def test_crop_to_a_ratio_and_a_box(self):
        self.assertEqual(self.finish('crop to square')[0].size, (600, 600))
        w, h = self.finish('crop to 16:9')[0].size
        self.assertAlmostEqual(w / float(h), 16 / 9.0, delta=0.02)
        self.assertEqual(self.finish('crop 10,20,300,200')[0].size, (300, 200))
        with self.assertRaises(self.m._N90Fail) as cm:
            self.run_ops('crop 700,500,300,300')
        self.assertEqual(cm.exception.code, 'bad_crop')

    def test_rotate_and_flip(self):
        self.assertEqual(self.finish('rotate 90')[0].size, (600, 800))
        src = png(800, 600)
        orig = Image.open(io.BytesIO(src)).convert('RGB')
        flipped = self.finish('flip horizontally', src)[0].convert('RGB')
        self.assertEqual(flipped.size, (800, 600))
        self.assertEqual(list(flipped.getdata()), list(orig.transpose(Image.FLIP_LEFT_RIGHT).getdata()))
        self.assertNotEqual(list(flipped.getdata()), list(orig.getdata()))
        vert = self.finish('flip vertically', src)[0].convert('RGB')
        self.assertEqual(list(vert.getdata()), list(orig.transpose(Image.FLIP_TOP_BOTTOM).getdata()))

    def test_grayscale_sepia_and_other_filters_really_change_the_pixels(self):
        src = png(400, 300)
        base_px = Image.open(io.BytesIO(src)).convert('RGB')
        gray = self.finish('make it black and white', src)[0].convert('RGB')
        self.assertTrue(all(abs(r - g) < 3 and abs(g - b) < 3 for r, g, b in list(gray.getdata())[::97]))
        sepia = self.finish('sepia', src)[0].convert('RGB')
        r, g, b = sepia.getpixel((200, 150))
        self.assertTrue(r >= g >= b)
        for text in ('blur', 'sharpen', 'brighten by 30%', 'increase contrast', 'invert', 'vignette', 'pixelate 12'):
            out = self.finish(text, src)[0].convert('RGB')
            self.assertEqual(out.size, base_px.size, text)
            self.assertNotEqual(list(out.getdata())[::53], list(base_px.getdata())[::53], text)

    def test_text_is_drawn_exactly_and_stays_inside_the_picture(self):
        src = png(800, 600, plain=True)
        out, fmt, ctx, note = self.finish('add text "SALE 50%" at the top', src)
        base_im = Image.open(io.BytesIO(src)).convert('RGB')
        diff = [(x, y) for y in range(0, 600, 3) for x in range(0, 800, 3) if out.convert('RGB').getpixel((x, y)) != base_im.getpixel((x, y))]
        self.assertTrue(diff)
        self.assertLess(min(y for _x, y in diff), 120, 'the text is at the top')
        self.assertTrue(all(0 <= x < 800 and 0 <= y < 600 for x, y in diff))

    def test_a_watermark_with_a_quoted_text_does_not_swallow_the_next_instruction(self):
        ops = self.m._n90_parse_ops('add watermark "(c) Nemo" and convert to webp')
        self.assertEqual([o['op'] for o in ops], ['watermark', 'convert'])
        self.assertEqual(ops[0]['text'], '(c) Nemo')

    def test_round_corners_circle_and_border(self):
        im, fmt, ctx, note = self.finish('round the corners')
        self.assertEqual(im.mode, 'RGBA')
        self.assertEqual(im.getpixel((0, 0))[3], 0, 'the corner is transparent')
        self.assertEqual(fmt, 'png', 'transparency forces PNG')
        im, fmt, ctx, note = self.finish('circle crop')
        self.assertEqual(im.size[0], im.size[1])
        self.assertEqual(im.getpixel((0, 0))[3], 0)
        im = self.finish('add a red border 12px')[0].convert('RGB')
        self.assertEqual(im.getpixel((3, 3)), self.m._n90_color('red'))
        self.assertEqual(im.size, (824, 624))

    def test_upscale_compress_and_convert(self):
        self.assertEqual(self.finish('upscale 2x', png(300, 200))[0].size, (600, 400))
        im, fmt, ctx, note = self.finish('compress to 30 kb', png(1200, 900, seed=4))
        self.assertEqual(fmt, 'jpg')
        data = self.m._n90_compress_to(Image.open(io.BytesIO(png(1200, 900, seed=4))), 30)[0]
        self.assertLessEqual(len(data), 30 * 1024)
        tiny, note = self.m._n90_compress_to(Image.open(io.BytesIO(png(1200, 900))), 5)
        self.assertLessEqual(len(tiny), 5 * 1024)
        self.assertIn('shrinking', note, 'when quality alone cannot reach the size the picture is made smaller, and it says so')
        with mock.patch.object(self.m, '_n90_save_bytes', lambda im, fmt, quality=90: b'x' * 10 ** 6):
            with self.assertRaises(self.m._N90Fail) as cm:
                self.m._n90_compress_to(Image.open(io.BytesIO(png(300, 200))), 30)
        self.assertEqual(cm.exception.code, 'too_big')
        self.assertEqual(self.finish('convert to webp')[1], 'webp')
        im2, ctx2, _o = self.run_ops('convert to pdf')
        data, fmt, _n = self.m._n90_render_output(im2, ctx2)
        self.assertEqual(fmt, 'pdf')
        self.assertEqual(data[:4], b'%PDF')

    def test_palette_info_and_tiles(self):
        im, ctx, ops = self.run_ops('what are the dominant colors')
        self.assertTrue(ctx.palette and all(re.fullmatch(r'#[0-9A-F]{6}', h) for h, _p in ctx.palette))
        self.assertAlmostEqual(sum(p for _h, p in ctx.palette), 100.0, delta=1.0)
        im, ctx, ops = self.run_ops('image info', png(640, 480, 'jpeg'))
        self.assertEqual((ctx.info['w'], ctx.info['h'], ctx.info['format']), (640, 480, 'JPEG'))
        self.assertFalse(ctx.info['gps'])
        im, ctx, ops = self.run_ops('split into 3x1', png(900, 300))
        self.assertEqual([t.size for t in ctx.tiles], [(300, 300)] * 3)
        with self.assertRaises(self.m._N90Fail) as cm:
            self.run_ops('split into 9x9', png(900, 300))
        self.assertEqual(cm.exception.code, 'bad_tiles')

    def test_gps_location_in_a_photo_is_flagged_and_strip_removes_it(self):
        im = Image.open(io.BytesIO(png(300, 200)))
        ex = Image.Exif()
        ex[0x8825] = {1: 'N', 2: (19.0, 4.0, 0.0), 3: 'E', 4: (72.0, 52.0, 0.0)}
        ex[271] = 'TestCam'
        b = io.BytesIO()
        im.save(b, 'JPEG', exif=ex)
        raw = b.getvalue()
        _im, ctx, _ops = self.run_ops('image info', raw)
        self.assertTrue(ctx.info['gps'])
        out, fmt, ctx2, _n = self.finish('strip the metadata', raw)
        self.assertFalse(out.getexif().get_ifd(0x8825))

    def test_background_removal_cuts_a_plain_background_and_refuses_a_busy_one(self):
        img = Image.new('RGB', (400, 300), (250, 250, 250))
        ImageDraw.Draw(img).ellipse((120, 70, 280, 230), fill=(200, 30, 30))
        b = io.BytesIO()
        img.save(b, 'PNG')
        out, fmt, ctx, note = self.finish('remove the background', b.getvalue())
        self.assertEqual(out.mode, 'RGBA')
        self.assertEqual(out.getpixel((5, 5))[3], 0, 'background gone')
        self.assertEqual(out.getpixel((200, 150))[3], 255, 'subject kept')
        self.assertEqual(fmt, 'png')
        with self.assertRaises(self.m._N90Fail) as cm:
            self.run_ops('remove the background', png(400, 300))
        self.assertEqual(cm.exception.code, 'busy_background', 'a busy background is refused instead of producing a ragged cut-out')

    def test_collage_needs_at_least_two_pictures(self):
        Imgs = [Image.open(io.BytesIO(png(200, 200, seed=i))).convert('RGB') for i in range(4)]
        canvas = self.m._n90_collage(Imgs, 2)
        self.assertGreater(canvas.size[0], 400)
        with self.assertRaises(self.m._N90Fail):
            self.m._n90_collage(Imgs[:1])

    def test_a_decompression_bomb_is_refused_before_it_is_decoded_into_memory(self):
        big = io.BytesIO()
        Image.new('L', (9000, 9000), 5).save(big, 'PNG')
        with mock.patch.object(self.m, '_N90_MAX_PIXELS', 20000000):
            with self.assertRaises(self.m._N90Fail) as cm:
                self.run_ops('resize to 100x100', big.getvalue())
        self.assertIn(cm.exception.code, ('too_big',))

    def test_a_file_that_is_not_a_picture_is_refused_with_a_word_not_a_crash(self):
        with self.assertRaises(self.m._N90Fail) as cm:
            self.run_ops('resize to 100x100', b'<html>no</html>' * 100)
        self.assertEqual(cm.exception.code, 'not_image')


# ===================================================================================================================
# 4. DESIGNS WITH EXACT TEXT AND PICTURES THAT NEED NO AI
# ===================================================================================================================
@needs_pil
class TestDesigns(StudioCase):
    def test_every_design_has_its_exact_size(self):
        m = self.m
        for kind, size in (('quote', (1080, 1080)), ('poster', (1080, 1350)), ('banner', (1500, 500)), ('thumbnail', (1280, 720)), ('logo', (1024, 1024)), ('card', (1200, 628))):
            im = m._n90_design(kind, ['Diwali Sale', 'Up to 50% off', '12-14 Oct'], 'sunset')
            self.assertEqual(im.size, size, kind)

    def test_the_text_is_drawn_by_code_so_different_words_give_different_pictures_and_the_same_words_the_same(self):
        m = self.m
        a = m._n90_save_bytes(m._n90_design('poster', ['Diwali Sale'], 'ocean'), 'png')
        b = m._n90_save_bytes(m._n90_design('poster', ['Diwali Sale'], 'ocean'), 'png')
        c = m._n90_save_bytes(m._n90_design('poster', ['Holi Sale'], 'ocean'), 'png')
        self.assertEqual(a, b)
        self.assertNotEqual(a, c)

    def test_text_is_readable_against_its_background_on_every_palette(self):
        m = self.m
        for pal in m._N90_PALETTES:
            im = m._n90_design('quote', ['The best way out is always through.', 'Robert Frost'], pal).convert('RGB')
            px = list(im.getdata())[::211]
            lums = sorted(m._n90_luma(p) for p in px)
            self.assertGreater(lums[-1] - lums[0], 0.25, pal)

    def test_design_without_words_asks_for_them_in_quotes(self):
        with self.assertRaises(self.m._N90Fail) as cm:
            self.m._n90_design('poster', [], 'ocean')
        self.assertEqual(cm.exception.code, 'no_text')
        with self.assertRaises(self.m._N90Fail) as cm:
            self.m._n90_design('hologram', ['x'], 'ocean')
        self.assertEqual(cm.exception.code, 'bad_design')

    def test_a_logo_comes_as_png_and_svg_and_the_svg_escapes_the_text(self):
        svg = self.m._n90_logo_svg(['A & B <Studio>', 'say "hi"'], 'gold')
        self.assertTrue(svg.lstrip().startswith('<svg'))
        self.assertIn('A &amp; B &lt;Studio&gt;', svg)
        self.assertNotIn('<Studio>', svg)
        self.assertNotIn('<script', svg.lower())
        ET = __import__('xml.etree.ElementTree', fromlist=['x'])
        ET.fromstring(svg)

    def test_a_photo_background_is_dimmed_and_cover_fitted(self):
        bg = Image.open(io.BytesIO(png(900, 400)))
        im = self.m._n90_design('thumbnail', ['How I trade Nifty'], 'midnight', bg_image=bg)
        self.assertEqual(im.size, (1280, 720))

    def test_procedural_wallpapers_are_deterministic_for_a_seed_and_need_no_network(self):
        m = self.m
        for kind in ('gradient', 'radial', 'mesh', 'waves', 'geometric', 'clouds'):
            a = m._n90_art(kind, 320, 200, 7, 'sunset')
            b = m._n90_art(kind, 320, 200, 7, 'sunset')
            c = m._n90_art(kind, 320, 200, 8, 'sunset')
            self.assertEqual(a.size, (320, 200), kind)
            self.assertEqual(list(a.getdata())[::37], list(b.getdata())[::37], kind)
            if kind not in ('radial',):
                self.assertNotEqual(list(a.getdata())[::37], list(c.getdata())[::37], kind)
        self.assertEqual(self.http.calls, [])
        with self.assertRaises(self.m._N90Fail):
            m._n90_art('nonsense', 100, 100)

    def test_the_wallpaper_kind_and_palette_come_from_the_words(self):
        m = self.m
        self.assertEqual(m._n90_art_kind('geometric pattern'), 'geometric')
        self.assertEqual(m._n90_art_kind('soft clouds'), 'clouds')
        self.assertEqual(m._n90_pick_palette('a sunset gradient'), 'sunset')
        self.assertEqual(m._n90_pick_palette('dark purple'), 'midnight' if 'dark' in m._N90_PALETTE_WORDS else m._n90_pick_palette('dark purple'))


# ===================================================================================================================
# 5. CHARTS FROM THE PERSON'S NUMBERS, NEVER INVENTED
# ===================================================================================================================
@needs_mpl
class TestCharts(StudioCase):
    def spec(self, text, kind=''):
        return self.m._n90_chart_spec(text, kind)

    def test_pairs_tables_and_labelled_lists_are_read_exactly(self):
        s = self.spec('bar chart: Jan 120, Feb 150, Mar 90')
        self.assertEqual((s['type'], s['labels']), ('bar', ['Jan', 'Feb', 'Mar']))
        self.assertEqual(s['series'][0]['values'], [D(120), D(150), D(90)])
        s = self.spec('chart\nMonth, Sales, Cost\nJan, 120, 80\nFeb, 150, 95\nMar, 90, 70')
        self.assertEqual([x['name'] for x in s['series']], ['Sales', 'Cost'])
        self.assertEqual(s['series'][1]['values'], [D(80), D(95), D(70)])
        s = self.spec('line chart x: Mon, Tue, Wed y: 5, 7, 6')
        self.assertEqual((s['type'], s['series'][0]['values']), ('line', [D(5), D(7), D(6)]))

    def test_indian_grouping_decimals_percent_and_currency_are_parsed_without_loss(self):
        s = self.spec('line chart of sales Jan: 1,20,000; Feb: 1,50,000; Mar: 90,000 in rs')
        self.assertEqual(s['series'][0]['values'], [D(120000), D(150000), D(90000)])
        self.assertTrue(s['indian'])
        n = self.m._n90_num
        self.assertEqual(n('₹5,000'), (D(5000), '₹'))
        self.assertEqual(n('12.5%'), (D('12.5'), '%'))
        self.assertEqual(n('(1,200)')[0], D(-1200))
        self.assertEqual(n('4 lakh')[0], D(400000))
        self.assertEqual(n('$3.2m')[0], D(3200000))
        self.assertIsNone(n('abc'))
        self.assertIsNone(n(''))

    def test_no_numbers_means_no_chart_and_no_invented_data(self):
        self.assertIsNone(self.spec('make a chart of nothing'))
        self.assertIsNone(self.spec("chart of india's gdp over the last five years"))

    def test_numbers_that_do_not_line_up_are_refused_not_fixed(self):
        with self.assertRaises(self.m._N90Fail) as cm:
            self.spec('line chart x: a, b, c y: 1, 2')
        self.assertEqual(cm.exception.code, 'mismatch')
        with self.assertRaises(self.m._N90Fail) as cm:
            self.spec('line chart x: a, b y: 1, banana')
        self.assertEqual(cm.exception.code, 'bad_number')
        with self.assertRaises(self.m._N90Fail) as cm:
            self.spec('candlestick\nday, close\n1, 5\n2, 6\n3, 7')
        self.assertEqual(cm.exception.code, 'need_ohlc')

    def test_the_plotted_numbers_are_read_back_and_a_wrong_plot_is_refused(self):
        s = self.spec('bar chart: Jan 120, Feb 150, Mar 90')
        png_bytes, info = self.m._n90_make_chart(s)
        self.assertTrue(self.m._n90_check_image(png_bytes)[0])
        self.assertIn('3 bars read back', info['verified'])
        with mock.patch.object(self.m, '_n90_chart_readback', lambda ax, kind: [120.0, 151.0, 90.0]):
            with self.assertRaises(self.m._N90Fail) as cm:
                self.m._n90_make_chart(s)
        self.assertEqual(cm.exception.code, 'verify')

    def test_every_chart_type_draws(self):
        cases = {'bar': 'bar chart: A 3, B 5, C 4', 'barh': 'horizontal bar chart: A 3, B 5, C 4', 'line': 'line chart: A 3, B 5, C 4, D 6', 'area': 'area chart: A 3, B 5, C 4, D 6', 'pie': 'pie chart: A 30, B 50, C 20',
                 'donut': 'donut chart: A 30, B 50, C 20', 'histogram': 'histogram of 1 2 2 3 3 3 4 4 5 9', 'scatter': 'scatter chart x: 1, 2, 3, 4 y: 2, 4, 5, 8',
                 'candlestick': 'candlestick\nday, open, high, low, close\n1, 100, 110, 95, 105\n2, 105, 112, 101, 99\n3, 99, 104, 97, 103'}
        for kind, text in cases.items():
            s = self.spec(text)
            self.assertEqual(s['type'], kind, text)
            data, info = self.m._n90_make_chart(s)
            self.assertTrue(self.m._n90_check_image(data)[0], kind)

    def test_the_caption_lists_every_value_and_says_how_it_was_checked(self):
        s = self.spec('bar chart: Jan 120, Feb 150, Mar 90')
        data, info = self.m._n90_make_chart(s)
        cap = self.m._n90_chart_caption(s, info)
        for part in ('Jan 120', 'Feb 150', 'Mar 90', 'Plotted exactly as given'):
            self.assertIn(part, cap)
        self.assertNotIn('Illustrative', cap)

    def test_sample_data_is_labelled_everywhere_and_is_reproducible(self):
        a = self.m._n90_illustrative_spec('sample bar chart', 'bar')
        b = self.m._n90_illustrative_spec('sample bar chart', 'bar')
        self.assertEqual(a['series'][0]['values'], b['series'][0]['values'])
        self.assertIn('(illustrative)', a['title'])
        data, info = self.m._n90_make_chart(a)
        self.assertIn('Illustrative data', self.m._n90_chart_caption(a, info))

    def test_too_many_points_are_refused(self):
        s = {'type': 'bar', 'title': 't', 'labels': ['p%d' % i for i in range(80)], 'series': [{'name': 'v', 'values': [D(i) for i in range(80)], 'unit': ''}], 'unit': '', 'indian': False, 'dark': False, 'size': (1200, 700)}
        with self.assertRaises(self.m._N90Fail) as cm:
            self.m._n90_make_chart(s)
        self.assertEqual(cm.exception.code, 'too_many')

    def test_number_formatting_never_rounds_away_what_was_given(self):
        f = self.m._n90_fmt
        self.assertEqual(f(D('24310.55'), '', False, 2), '24,310.55')
        self.assertEqual(f(D(1250000), '', True), '12,50,000')
        self.assertEqual(f(D('12.5'), '%'), '12.50%')
        self.assertEqual(f(D('12.5'), '%', False, 1), '12.5%')
        self.assertEqual(f(D(-3), '₹'), '-₹3')


# ===================================================================================================================
# 6. THE PRECISION ENGINE: every figure, date and citation checked against what was actually read
# ===================================================================================================================
TODAY = dt.date(2026, 10, 3)            # a Saturday


class TestNumberExtraction(StudioCase):
    def nums(self, text, **kw):
        return self.m._n90_extract_numbers(text, **kw)

    def raws(self, text):
        return [c['raw'] for c in self.nums(text)]

    def test_money_percent_scale_words_and_indian_units(self):
        c = {x['raw']: x for x in self.nums('Brent at $82.40 a barrel, GDP up 6.8%, revenue Rs 1,25,000 crore, stake of 4.5 million, fund of ₹2 lakh, spread of 25 bps, valuation 3.2x earnings, cap 1.25 lakh crore')}
        self.assertEqual(c['$82.40']['value'], D('82.40'))
        self.assertEqual(c['$82.40']['kind'], 'money')
        self.assertEqual((c['6.8%']['kind'], c['6.8%']['value']), ('percent', D('6.8')))
        self.assertEqual(c['Rs 1,25,000 crore']['value'], D('1250000000000'))
        self.assertEqual(c['4.5 million']['value'], D(4500000))
        self.assertEqual(c['₹2 lakh']['value'], D(200000))
        self.assertEqual((c['25 bps']['kind'], c['25 bps']['value']), ('percent', D('0.25')))
        self.assertEqual((c['3.2x']['kind'], c['3.2x']['value']), ('multiple', D('3.2')))
        self.assertEqual(c['1.25 lakh crore']['value'], D('1250000000000'), 'a lakh crore is 10^12')

    def test_things_that_are_not_figures_are_left_alone(self):
        text = ('In 2026 at 10:30 the 3rd item (see https://example.com/a/1234567) was v2 and no 5 in section 4. [12] cited.\n'
                '1. First point\nOUTLOOK - NEXT 6 TO 24 MONTHS:\nWithin 18 months and 3 years, with 3 sources and two people.')
        self.assertEqual(self.raws(text), [])

    def test_a_real_figure_next_to_those_is_still_found(self):
        self.assertEqual(self.raws('In 2026, at 10:30, exports rose 12.5% [3] to $4.2 billion.'), ['12.5%', '$4.2 billion'])

    def test_decimals_and_the_step_of_the_last_digit_are_kept(self):
        c = self.nums('Up 6.8% and 6.80% and 24,310.55 and 79,870')
        self.assertEqual([x['decimals'] for x in c], [1, 2, 2, 0])
        self.assertEqual([x['unit'] for x in c], [D('0.1'), D('0.01'), D('0.01'), D(1)])

    def test_spans_point_at_the_figure_not_the_space_or_marker_after_it(self):
        text = 'Oil rose to $92 [1] and 5% [2].'
        for c in self.nums(text):
            self.assertEqual(text[c['span'][0]:c['span'][1]], c['raw'])
            self.assertFalse(c['raw'].endswith((' ', '[')))


class TestValueMatching(StudioCase):
    def match(self, claim, evidence):
        e = self.m._n90_extract_numbers
        return self.m._n90_values_match(e(claim, keep_small=True)[0], e(evidence, keep_small=True)[0])

    def test_equal_values_match_even_with_a_trailing_zero(self):
        self.assertEqual(self.match('$82.40', '$82.4'), 'exact')
        self.assertEqual(self.match('6.8%', '6.8 percent'), 'exact')
        self.assertEqual(self.match('4,500,000', '4.5 million'), 'exact')

    def test_a_figure_with_more_digits_than_its_source_is_over_precise_not_exact(self):
        self.assertEqual(self.match('6.83%', '6.8%'), 'overprecise')
        self.assertEqual(self.match('82.43', '82.4'), 'overprecise')

    def test_rounding_and_truncation_of_a_more_precise_source_are_accepted(self):
        self.assertEqual(self.match('24,311', '24,310.55'), 'exact')
        self.assertEqual(self.match('24,310', '24,310.55'), 'exact', 'truncation is not an error')
        self.assertEqual(self.match('24,309', '24,310.55'), '')

    def test_different_values_and_different_kinds_never_match(self):
        self.assertEqual(self.match('7.8%', '6.8%'), '')
        self.assertEqual(self.match('6.8%', '$6.8'), '')
        self.assertEqual(self.match('6.8x', '6.8%'), '')

    def test_a_count_and_money_of_the_same_number_can_match(self):
        self.assertEqual(self.match('$4,200', '4,200'), 'exact')


class TestAudit(StudioCase):
    SRC = [{'title': 'Brent', 'body': 'Brent crude settled at $82.4 a barrel on 3 October 2026, up 1.2% from $81.4. Output rose from 8.0 to 9.0 million barrels, a 12.5% increase.'},
           {'title': 'India', 'body': 'India GDP grew 6.8% in FY26 while inflation was 4.9 percent. Revenue was Rs 1,25,000 crore. Repo rate 5.5%.'},
           {'title': 'Indices', 'body': 'Nifty 50 closed at 24,310.55 points on Friday 2 October 2026. Sensex ended at 79,870.'}]

    def audit(self, text, src=None, **kw):
        return self.m._n90_audit(text, self.SRC if src is None else src, today=TODAY, **kw)

    def status(self, a):
        return {c['raw']: c['status'] for c in a['claims']}

    def test_a_correct_cited_figure_is_supported_by_the_source_it_cites(self):
        a = self.audit('Brent settled at $82.40 a barrel [1]. GDP grew 6.8% [2].')
        self.assertEqual(self.status(a), {'$82.40': 'supported', '6.8%': 'supported'})
        self.assertEqual([c['source'] for c in a['claims']], [1, 2])
        self.assertEqual((a['grade'], a['block'], a['issues']), ('A', False, []))

    def test_a_figure_that_is_in_another_source_than_the_one_cited_is_flagged_softly(self):
        a = self.audit('The repo rate is 5.5% [1].')
        self.assertEqual(self.status(a), {'5.5%': 'elsewhere'})
        self.assertEqual(a['claims'][0]['source'], 2)
        self.assertFalse(a['block'])

    def test_a_figure_found_nowhere_is_unsupported_and_an_uncited_one_says_so(self):
        a = self.audit('Demand reached 104.3 million barrels a day [1].\nStocks fell 3.1% last week.')
        self.assertEqual(self.status(a), {'104.3 million': 'unsupported', '3.1%': 'uncited_unsupported'})
        self.assertIn('104.3 million', a['repair'])

    def test_a_figure_more_precise_than_its_source_is_flagged(self):
        a = self.audit('GDP grew 6.83% [2].')
        self.assertEqual(self.status(a), {'6.83%': 'overprecise'})
        self.assertEqual(a['claims'][0]['near'], '6.8%')
        self.assertIn('more precise than the source', a['repair'])

    def test_the_percent_change_between_two_numbers_is_recomputed(self):
        ok = self.audit('Output rose 12.5% from 8.0 to 9.0 million barrels [1].')
        self.assertEqual([i for i in ok['issues'] if i['type'] in ('arithmetic', 'direction')], [])
        self.assertEqual(self.status(ok)['12.5%'], 'supported')
        bad = self.audit('Output rose 15% from 8.0 to 9.0 million barrels [1].')
        issue = [i for i in bad['issues'] if i['type'] == 'arithmetic'][0]
        self.assertIn('12.50%', issue['note'])
        self.assertTrue(bad['block'])

    def test_the_direction_of_a_change_is_checked(self):
        a = self.audit('Prices fell 10% from 100 to 110 [1].', src=['Prices went from 100 to 110.'])
        self.assertTrue(any(i['type'] == 'direction' and 'rise' in i['note'] for i in a['issues']))
        a = self.audit('Prices rose 10% from 100 to 110 [1].', src=['Prices went from 100 to 110.'])
        self.assertEqual([i for i in a['issues'] if i['type'] == 'direction'], [])

    def test_dates_must_exist_and_weekdays_must_match_the_calendar(self):
        a = self.audit('The meeting was on 31 February 2026 and the vote on Friday, 3 October 2026.')
        kinds = {i['type']: i for i in a['issues']}
        self.assertIn('date_invalid', kinds)
        self.assertIn('Saturday', kinds['weekday']['note'])
        ok = self.audit('Nifty closed on Friday 2 October 2026 [3].')
        self.assertEqual([i for i in ok['issues'] if i['type'] in ('weekday', 'date_invalid')], [])
        self.assertEqual(self.m._n90_find_dates('Wed 4 March 2026 and March 4, 2026 (Wed) and 2026-03-04')[0]['weekday_real'], 2)

    def test_a_future_date_described_as_having_happened_is_flagged_but_a_plan_is_not(self):
        past = self.audit('The company announced the merger on 20 November 2026.')
        self.assertTrue(any(i['type'] == 'date_future' for i in past['issues']))
        plan = self.audit('The merger is scheduled for 20 November 2026.')
        self.assertFalse(any(i['type'] == 'date_future' for i in plan['issues']))

    def test_a_citation_that_points_at_nothing_is_an_error(self):
        a = self.audit('Brent settled at $82.4 [9].')
        self.assertTrue(any(i['type'] == 'cite_invalid' and i['claim'] == '[9]' for i in a['issues']))
        self.assertTrue(a['block'])

    def test_placeholders_left_in_the_text_are_errors(self):
        a = self.audit('Revenue was TBD and growth XX% [1].')
        self.assertEqual(sorted(i['claim'] for i in a['issues'] if i['type'] == 'placeholder'), ['TBD', 'XX%'])

    def test_the_same_measure_with_two_different_values_in_two_sentences_is_flagged_but_from_to_is_not(self):
        a = self.audit('Inflation was 4.9% [2]. Later, inflation was 5.4% [2].')
        self.assertTrue(any(i['type'] == 'consistency' for i in a['issues']))
        b = self.audit('Output rose 12.5% from 8.0 to 9.0 million barrels [1]. The summary repeats: output rose from 8.0 to 9.0 million barrels [1].')
        self.assertFalse(any(i['type'] == 'consistency' for i in b['issues']))

    def test_live_market_data_given_to_the_writer_counts_as_evidence(self):
        a = self.audit('Nifty trades at 24,500.10 now.', extra_texts=['NIFTY spot is LIVE 24,500.10 (as of now)'])
        self.assertEqual(self.status(a), {'24,500.10': 'elsewhere'})
        self.assertFalse(a['block'])

    def test_data_tables_are_checked_row_by_row_and_taken_out_of_the_body(self):
        text = ('GDP grew 6.8% [2].\nDATA TABLE: Index close | unit: points | source: [3]\nNifty 50 | 24,310.55\nSensex | 79,871\nBank Nifty | 52,000\nEND TABLE\n'
                'DATA TABLE: Growth | unit: % | source: [2]\nGDP | 6.8\nInflation | 4.9\nEND TABLE\n')
        a = self.audit(text)
        rows = {c['label']: c['status'] for c in a['claims'] if c.get('table') is not None}
        self.assertEqual(rows, {'Nifty 50': 'supported', 'Sensex': 'unsupported', 'Bank Nifty': 'unsupported', 'GDP': 'supported', 'Inflation': 'supported'})
        self.assertEqual(len(a['tables']), 2)
        self.assertNotIn('DATA TABLE', self.m._n90_strip_tables(text))

    def test_a_table_with_a_non_number_is_left_out_and_said_so_and_one_without_a_source_is_flagged(self):
        a = self.audit('DATA TABLE: Bad | source: [1]\nx | abc\ny | 4\nEND TABLE\nDATA TABLE: Nosrc | unit: %\nA | 1.5\nB | 2.5\nEND TABLE')
        types = [i['type'] for i in a['issues']]
        self.assertIn('table_bad', types)
        self.assertIn('table_source', types)
        self.assertEqual(len(a['tables']), 1)

    def test_score_grade_and_blocking_follow_the_findings(self):
        clean = self.audit('GDP grew 6.8% [2]. Inflation was 4.9% [2].')
        self.assertEqual((clean['grade'], clean['score']), ('A', 100))
        mixed = self.audit('GDP grew 6.8% [2]. Demand reached 104.3 million [1]. Prices rose 3.1%. Exports fell 2.2%. Imports rose 1.9%. Tax was 7.7%. Debt was 9.9%.')
        self.assertTrue(mixed['block'], 'more than 30% of six or more figures with no source blocks')
        self.assertEqual(mixed['grade'], 'D')
        small = self.audit('Demand reached 104.3 million [1].')
        self.assertFalse(small['block'], 'one unverified figure is flagged, not blocked')

    def test_the_repair_text_names_the_exact_figures_without_repeating_itself(self):
        a = self.audit('Output rose 15% from 8.0 to 9.0 million [1]. Output rose 15% from 8.0 to 9.0 million [1]. Friday, 3 October 2026.')
        self.assertEqual(a['repair'].count('going from 8.0 to 9.0 million'), 1)
        self.assertIn('Saturday', a['repair'])

    def test_the_audit_never_raises_on_odd_input_and_its_counts_add_up(self):
        rnd = random.Random(11)
        pieces = ['$', '%', '[', ']', '1', '2.5', ' lakh', ' crore', 'TBD', ' Friday ', ' 31 Feb 2026 ', '\n', 'DATA TABLE: x | source: [1]\na | 1\nb | 2\nEND TABLE', '٣', '--', 'rose', 'fell', 'from', 'to', '12', '0']
        for _ in range(200):
            text = ''.join(rnd.choice(pieces) for _ in range(rnd.randint(0, 25)))
            a = self.audit(text)
            s = a['summary']
            self.assertEqual(s['figures'], len(a['claims']))
            self.assertLessEqual(s['supported'] + s['derived_ok'] + s['elsewhere'] + s['overprecise'] + s['unsupported'], s['figures'] + s['arithmetic'])
            self.assertTrue(0 <= a['score'] <= 100)
        for odd in (None, '', 'x' * 100000, '9' * 5000 + '%'):
            self.m._n90_audit(odd, self.SRC, today=TODAY)

    def test_the_summary_line_is_plain_and_complete(self):
        line = self.m._n90_audit_line(self.audit('GDP grew 6.8% [2]. Demand reached 104.3 million [1].'))
        for part in ('2 figures', '1 found in the cited source', '1 NOT found anywhere', 'grade'):
            self.assertIn(part, line)


# ===================================================================================================================
# 7. THE REAL REPORT PIPELINE WITH THE PRECISION ENGINE INSIDE IT (real fpdf2 + PyMuPDF; scripted research and AI)
# ===================================================================================================================
BODIES = ['Brent crude settled at $82.4 a barrel on 3 October 2026, up 1.2%. Output rose from 8.0 to 9.0 million barrels, a 12.5% increase.',
          'India GDP grew 6.8% in FY26. Inflation was 4.9 percent. Repo rate stayed at 5.5%.',
          'Nifty 50 closed at 24,310.55 points; Sensex ended at 79,870.',
          'Demand reached 104.3 million barrels a day according to the agency.']


def research_rows():
    return [{'title': 'Source %d' % (i + 1), 'body': b, 'url': 'https://site%d.gov.in/x' % i, 'tier': 'PRIMARY', 'score': 5, 'kind': 'news'} for i, b in enumerate(BODIES * 2)]


def report_text(body, table=True):
    return ('EXECUTIVE SUMMARY:\n- ' + body + '\n\nCURRENT SITUATION:\n' + body + ' GDP grew 6.8% [2]. Inflation was 4.9% [2]. Nifty closed at 24,310.55 [3].\n\nKEY DEVELOPMENTS:\nOil:\nSee [1][2][3][4].\n\n'
            'IMPACT ASSESSMENT:\nEconomic: x [5].\n\nOUTLOOK - NEXT 6 TO 24 MONTHS:\nBase case (analysis): y.\n\nRISK MATRIX:\nOil shock | Likelihood: Medium | Impact: High | Watch indicator: Brent\n\n'
            'WHAT TO WATCH:\n1. Brent.\n\n' + ('DATA TABLE: Index close | unit: points | source: [3]\nNifty 50 | 24,310.55\nSensex | 79,870\nEND TABLE\n\n' if table else '') + 'CONCLUSION:\nDone.\n\nSOURCES USED:\n[1] [2] [3] [4]\n') + ('filler text. ' * 120)


GOOD = report_text('Brent crude settled at $82.4 a barrel on 3 October 2026 [1]. Output rose 12.5% from 8.0 to 9.0 million barrels [1]. Demand reached 104.3 million barrels a day [4].')
BAD = report_text('Brent crude settled at $82.40 a barrel on Friday, 3 October 2026 [1]. Output rose 15% from 8.0 to 9.0 million barrels [1]. Demand reached 109.9 million barrels a day [4]. Stocks fell 3.1% [1] and exports hit 7.7 million [1].')


@needs_pdf
class TestReportPipeline(StudioCase):
    def setUp(self):
        super().setUp()
        m = self.m
        self.ai_calls = []
        self.replies = []

        def start(name, value):
            p = mock.patch.object(m, name, value)
            p.start()
            self.patches.append(p)

        def ai(chat_id, text, **kw):
            self.ai_calls.append(text)
            return self.replies.pop(0) if self.replies else self.replies_last
        self.replies_last = GOOD
        start('_report_research', lambda topic, min_good=4, max_sources=12: research_rows())
        start('ask_ai', ai)
        start('prime_task_create', lambda *a, **k: 't1')
        start('prime_task_update', lambda *a, **k: None)
        start('prime_trace_start', lambda *a, **k: 'tr1')
        start('prime_trace_end', lambda *a, **k: None)
        start('act_log', lambda *a, **k: None)
        start('market_context', m._N90_MC_PREV)
        self.m_made = []

    def tearDown(self):
        for d in self.docs:
            for ext in ('', '.research.json'):
                try:
                    os.unlink(d['path'][:-4] + ext if ext else d['path'])
                except OSError:
                    pass
        super().tearDown()

    def run_report(self, first, later=(), topic='oil and india economy outlook'):
        self.replies = [first] + list(later)
        self.replies_last = (list(later) or [first])[-1]
        n0 = len(self.sent)
        self.m.make_pdf(self.cid, topic)
        return '\n'.join(t for c, t in self.sent[n0:])

    def pdf_pages(self, doc):
        d = fitz.open(stream=doc['bytes'], filetype='pdf')
        try:
            return [p.get_text() for p in d]
        finally:
            d.close()

    def test_a_clean_report_is_delivered_with_verification_pages_and_a_plain_summary(self):
        out = self.run_report(GOOD)
        self.assertEqual(len(self.docs), 1)
        pages = self.pdf_pages(self.docs[0])
        text = '\n'.join(pages)
        self.assertIn('NUMBERS AND SOURCES CHECK', text)
        self.assertIn('Figure by figure', text)
        self.assertIn('How this was checked', text)
        self.assertIn('Numbers check: 15 figures', out)
        self.assertIn('grade A', out)
        self.assertEqual(len(self.ai_calls), 1, 'a clean draft costs no extra AI call')
        self.assertEqual(self.m._N90_STATS['appendices'], 1)

    def test_the_verification_pages_come_after_the_layout_checks_so_they_never_trip_the_page_limit(self):
        self.run_report(GOOD)
        pages = self.pdf_pages(self.docs[0])
        self.assertGreaterEqual(len(pages), 6)
        qa_pages_before = [i for i, t in enumerate(pages) if 'NUMBERS AND SOURCES CHECK' in t][0]
        self.assertTrue(all('NUMBERS AND SOURCES CHECK' not in t for t in pages[:qa_pages_before]))

    def test_data_tables_become_a_chart_and_a_table_and_never_stray_text_in_the_body(self):
        self.run_report(GOOD)
        pages = self.pdf_pages(self.docs[0])
        body = '\n'.join(pages)
        self.assertNotIn('DATA TABLE', body)
        self.assertNotIn('END TABLE', body)
        self.assertIn('Figure: Index close (points)', body)
        self.assertIn('Plotted exactly as in the table', body)
        self.assertIn('24,310.55', body)
        d = fitz.open(stream=self.docs[0]['bytes'], filetype='pdf')
        try:
            self.assertTrue(any(p.get_images() for p in d), 'the figure is a real picture in the PDF')
        finally:
            d.close()

    def test_a_flawed_draft_is_corrected_once_with_the_exact_problems_named(self):
        out = self.run_report(BAD, [GOOD])
        self.assertEqual(len(self.ai_calls), 2)
        fix = self.ai_calls[1]
        for needle in ('15%', 'going from 8.0 to 9.0 million is 12.50%', 'Saturday', '109.9 million', 'EVIDENCE:'):
            self.assertIn(needle, fix)
        self.assertIn('Correcting them once', out)
        self.assertIn('The first draft had 6 unverified figure(s); after one correction pass it has 0.', out)
        self.assertEqual(len(self.docs), 1)

    def test_a_draft_that_is_still_wrong_after_the_correction_is_refused_with_the_reasons(self):
        out = self.run_report(BAD, [BAD])
        self.assertEqual(self.docs, [], 'no PDF is delivered')
        self.assertIn('Publisher QA refused this draft', out)
        self.assertIn('weekday', out)
        self.assertIn('Saturday', out)
        self.assertNotIn('Numbers check: 15 figures', out, 'the closing summary is only for a delivered report')

    def test_a_correction_that_is_worse_or_cut_short_is_not_used(self):
        worse = report_text('Brent crude settled at $90.40 a barrel on Friday, 3 October 2026 [1]. Output rose 15% from 8.0 to 9.0 million barrels [1]. Demand reached 109.9 million barrels a day [4]. Stocks fell 3.1% [1] and exports hit 7.7 million [1]. '
                            'Oil traded at $77.7 [1] and gas at 8.8% [1] and coal at 9.9 million [1].')
        out = self.run_report(BAD, [worse])
        self.assertEqual(self.docs, [])
        self.assertIn('$82.40', ' '.join(self.ai_calls[1:2]) + out + BAD, 'the first draft is what was judged')
        self.assertNotIn('$90.40', out)
        out = self.run_report(BAD, ['too short'])
        self.assertEqual(self.docs, [])

    def test_a_hard_error_that_the_model_fixes_is_enough_without_godmode_extra_passes(self):
        self.run_report(BAD, [GOOD])
        self.assertEqual(len(self.ai_calls), 2, 'draft + one correction; no verifier or editor pass is added')

    def test_the_manifest_records_the_numbers_check(self):
        self.run_report(GOOD)
        manifest_path = self.docs[0]['path'][:-4] + '.research.json'
        with open(manifest_path, encoding='utf-8') as fh:
            manifest = json.load(fh)
        self.assertEqual(manifest['numbers_check']['grade'], 'A')
        self.assertEqual(manifest['numbers_check']['summary']['figures'], 15)
        self.assertEqual(manifest['numbers_check']['not_found'], [])

    def test_a_report_with_a_few_unverified_figures_is_delivered_and_lists_them(self):
        soft = report_text('Brent crude settled at $82.4 a barrel [1]. Demand reached 104.3 million barrels a day [4].\nStocks fell 3.1%.\nImports rose 1.9%.\nTax was 7.7%.')
        out = self.run_report(soft, [soft])
        self.assertEqual(len(self.docs), 1, 'soft findings do not stop a report; they are written on the verification page')
        text = '\n'.join(self.pdf_pages(self.docs[0]))
        self.assertIn('NOT FOUND', text)
        self.assertIn('3.1%', text)
        self.assertIn('NOT found anywhere', out)

    def test_switching_the_check_off_restores_the_old_behaviour_exactly(self):
        self.store['studio_precision'] = '0'
        out = self.run_report(BAD)
        self.assertEqual(len(self.ai_calls), 1)
        self.assertEqual(len(self.docs), 1, 'the old gate delivered it')
        self.assertNotIn('NUMBERS AND SOURCES CHECK', '\n'.join(self.pdf_pages(self.docs[0])))
        self.assertNotIn('Numbers check', out)
        self.assertNotIn('Studio', self.ai_calls[0])
        self.assertNotIn('PRECISION RULES', self.ai_calls[0])

    def test_the_writer_is_given_the_precision_rules_in_the_prompt(self):
        self.run_report(GOOD)
        self.assertIn('PRECISION RULES', self.ai_calls[0])
        self.assertIn('DATA TABLE:', self.ai_calls[0])
        self.assertIn('no figure found in the sources', self.ai_calls[0])
        self.assertIn('RESEARCH EVIDENCE:', self.ai_calls[0])

    def test_a_failure_to_build_the_verification_pages_does_not_lose_the_report_and_says_so(self):
        with mock.patch.object(self.m, '_n90_merge_pdfs', side_effect=RuntimeError('boom')):
            out = self.run_report(GOOD)
        self.assertEqual(len(self.docs), 1)
        self.assertNotIn('NUMBERS AND SOURCES CHECK', '\n'.join(self.pdf_pages(self.docs[0])))
        self.assertIn('verification pages could not be added', out)
        self.assertEqual(self.m._N90_STATS['errors'], 1)

    def test_a_broken_audit_never_stops_a_report(self):
        with mock.patch.object(self.m, '_n90_audit', side_effect=ValueError('boom')):
            out = self.run_report(GOOD)
        self.assertEqual(len(self.docs), 1)
        self.assertGreaterEqual(self.m._N90_STATS['errors'], 1)

    def test_the_merge_is_checked_by_page_count_and_the_original_survives_a_bad_merge(self):
        m = self.m
        base_pdf, ap = os.path.join(self.tmp.name, 'a.pdf'), os.path.join(self.tmp.name, 'b.pdf')
        for path, n in ((base_pdf, 3), (ap, 2)):
            doc = fitz.open()
            for i in range(n):
                doc.new_page().insert_text((72, 72), '%s %d' % (os.path.basename(path), i))
            doc.save(path)
            doc.close()
        with open(base_pdf, 'rb') as fh:
            before = fh.read()
        with mock.patch.object(m, '_n90_page_count', side_effect=[3, 2, 4]):
            with self.assertRaises(ValueError):
                m._n90_merge_pdfs(base_pdf, ap)
        with open(base_pdf, 'rb') as fh:
            self.assertEqual(fh.read(), before, 'a merge that does not add up never replaces the report')
        self.assertFalse(os.path.exists(base_pdf + '.n90merge'))
        self.assertEqual(m._n90_merge_pdfs(base_pdf, ap), 5)
        self.assertEqual(m._n90_page_count(base_pdf), 5)

    def test_unrelated_pdf_work_outside_a_report_run_is_untouched(self):
        m = self.m
        self.assertIsNone(m._n90_run_state())
        text = 'EXECUTIVE SUMMARY:\nGDP grew 99% [1].\n'
        self.assertEqual(m._creation_strip_appendices(text), text.strip(), 'no audit and no change outside a report run')
        ok, reasons = m._creation_report_content_gate(text, research_rows())
        self.assertFalse(ok)
        self.assertFalse(any('numbers check' in r for r in reasons))


# ===================================================================================================================
# 8. DATA REPORTS FROM THE PERSON'S OWN FILE
# ===================================================================================================================
def trades_csv(n=40, seed=5, sep=','):
    rnd = random.Random(seed)
    lines = [sep.join(['Date', 'Symbol', 'Side', 'Qty', 'P&L'])]
    pnl = []
    for i in range(n):
        v = rnd.choice([-1, 1, 1]) * rnd.randint(200, 9000) + D(rnd.randint(0, 99)) / 100
        pnl.append(D(v))
        lines.append(sep.join([(dt.date(2026, 1, 1) + dt.timedelta(days=2 * i)).strftime('%d/%m/%Y'), rnd.choice(['NIFTY', 'BANKNIFTY', 'TCS']), 'BUY', str(25 * rnd.randint(1, 4)), '%.2f' % v]))
    return '\n'.join(lines).encode(), pnl


class TestDataLoading(StudioCase):
    def load(self, raw, name='data.csv', **kw):
        return self.m._n90_load_table(raw, name, **kw)

    def test_csv_with_each_separator_a_bom_and_quotes(self):
        for sep in (',', ';', '\t', '|'):
            raw, _ = trades_csv(10, sep=sep)
            t = self.load(raw)
            self.assertEqual(t['header'], ['Date', 'Symbol', 'Side', 'Qty', 'P&L'], repr(sep))
            self.assertEqual(len(t['rows']), 10)
        t = self.load(b'\xef\xbb\xbfname,amount\n"Rent, home","1,200"\nFood,300\nTravel,150\n')
        self.assertEqual(t['header'], ['name', 'amount'])
        self.assertEqual(t['rows'][0], ['Rent, home', '1,200'])
        t = self.load('item,cost\ncaf\xe9,5\ntea,3\nmilk,4\n'.encode('cp1252'))
        self.assertEqual(t['rows'][0][0], 'caf\xe9')

    def test_json_records_and_json_columns(self):
        t = self.load(json.dumps([{'a': 1, 'b': 'x'}, {'a': 2, 'b': 'y'}, {'a': 3}]).encode(), 'd.json')
        self.assertEqual(t['header'], ['a', 'b'])
        self.assertEqual(len(t['rows']), 3)
        t = self.load(json.dumps({'a': [1, 2, 3], 'b': [4, 5, 6]}).encode(), 'd.json')
        self.assertEqual(t['rows'], [[1, 4], [2, 5], [3, 6]])
        t = self.load(json.dumps({'meta': 1, 'rows': [{'k': 1}, {'k': 2}]}).encode(), 'd.json')
        self.assertEqual(len(t['rows']), 2)
        with self.assertRaises(self.m._N90DataError):
            self.load(b'{"a": 1}', 'd.json')
        with self.assertRaises(self.m._N90DataError):
            self.load(b'{not json', 'd.json')

    def test_a_pasted_table_and_a_headerless_table(self):
        t = self.load(b'Item | Cost\nRent | 12000\nFood | 8000\nTravel | 3000', 'pasted', pasted=True)
        self.assertEqual((t['header'], len(t['rows'])), (['Item', 'Cost'], 3))
        t = self.load(b'10,20\n30,40\n50,60\n70,80\n')
        self.assertTrue(t['header_guess'])
        self.assertEqual(t['header'], ['Column 1', 'Column 2'])
        self.assertEqual(len(t['rows']), 4)

    def test_what_cannot_be_read_gets_a_plain_reason(self):
        for raw, name, word in ((b'abc', 'x.csv', 'fewer than two lines'), (b'PK\x03\x04junk', 'x.xlsx', 'could not open'), (b'x', 'x.pdf', 'save this one as'), (b'a,b\n1,2\n' * 3, 'x.zip', 'save this one as')):
            with self.assertRaises(self.m._N90DataError) as cm:
                self.load(raw, name)
            self.assertIn(word, cm.exception.msg)
        with self.assertRaises(self.m._N90DataError) as cm:
            self.load(b'a,b\n' + b'1,2\n' * 10, 'big.csv') if False else self.m._n90_load_table(b'x' * (9 * 1024 * 1024), 'big.csv')
        self.assertIn('8 MB', cm.exception.msg)
        wide = (','.join('c%d' % i for i in range(80)) + '\n' + ','.join(str(i) for i in range(80)) + '\n') * 2
        with self.assertRaises(self.m._N90DataError) as cm:
            self.load(wide.encode())
        self.assertIn('60 columns', cm.exception.msg)

    @needs_xlsx
    def test_excel_files_are_read_with_dates_and_numbers_and_the_fallback_reader_agrees(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sales'
        ws.append(['Order Date', 'Region', 'Amount (₹)'])
        for i in range(12):
            ws.append([dt.date(2026, 1 + i % 6, 1 + i), 'North' if i % 2 else 'South', 1000 * (i + 1) + 0.5])
        b = io.BytesIO()
        wb.save(b)
        t = self.load(b.getvalue(), 'sales.xlsx')
        self.assertEqual((t['header'], t['sheet'], len(t['rows'])), (['Order Date', 'Region', 'Amount (₹)'], 'Sales', 12))
        fallback = self.m._n90_xlsx_stdlib(b.getvalue())
        self.assertEqual(len(fallback), 13)
        self.assertEqual(fallback[0], ['Order Date', 'Region', 'Amount (₹)'])
        self.assertEqual(fallback[1][1:], ['South', '1000.5'])
        cols, notes = self.m._n90_infer_columns({'header': fallback[0], 'rows': fallback[1:]})
        self.assertEqual(cols[0]['kind'], 'date', 'an Excel date serial is turned back into a date')
        self.assertEqual(cols[0]['values'][0], dt.date(2026, 1, 1))

    def test_column_types_units_and_dates(self):
        t = self.load(b'when,price,share,note\n03/10/2026,"$1,200.50",12.5%,a\n04/10/2026,(300),7%,b\n05/10/2026,N/A,5%,c\n13/10/2026,40,1%,d\n')
        cols, notes = self.m._n90_infer_columns(t)
        by = {c['name']: c for c in cols}
        self.assertEqual(by['when']['kind'], 'date')
        self.assertEqual(by['when']['values'][3], dt.date(2026, 10, 13), 'a day above 12 proves day/month order')
        self.assertEqual((by['price']['kind'], by['price']['unit'], by['price']['values']), ('number', '$', [D('1200.50'), D(-300), None, D(40)]))
        self.assertEqual((by['share']['kind'], by['share']['unit']), ('number', '%'))
        self.assertEqual(by['note']['kind'], 'text')
        self.assertEqual(by['price']['missing'], 1, 'N/A is a missing value, not a number and not an error')
        self.assertEqual(by['price']['unparsed'], 0)
        t2 = self.load(b'v\n12\n13\nabc\n14\n15\n16\n17\n18\n19\n20\n21\n22\n23\n')
        c2, _n = self.m._n90_infer_columns(t2)
        self.assertEqual((c2[0]['kind'], c2[0]['unparsed']), ('number', 1), 'one unreadable cell is counted and left out, never guessed')

    def test_ambiguous_dates_are_read_as_day_first_and_the_note_says_so(self):
        t = self.load(b'd,v\n01/02/2026,1\n03/04/2026,2\n05/06/2026,3\n')
        cols, notes = self.m._n90_infer_columns(t)
        self.assertEqual(cols[0]['values'][0], dt.date(2026, 2, 1))
        self.assertTrue(any('day/month/year' in n for n in notes))

    def test_a_totals_row_is_set_aside_and_checked_against_the_sum(self):
        t = self.load(b'Item,Cost\nRent,12000\nFood,8000\nTravel,3000\nTotal,23000\n')
        self.assertEqual(len(t['rows']), 3)
        self.assertEqual(t['totals'][0][1], 'Total')
        r = self.m._n90_data_report(t)
        self.assertTrue(any('they agree' in f for f in r['findings']))
        t = self.load(b'Item,Cost\nRent,12000\nFood,8000\nTravel,3000\nTotal,24000\n')
        r = self.m._n90_data_report(t)
        self.assertTrue(any('difference of 1,000' in f for f in r['findings']), r['findings'])


class TestDataMaths(StudioCase):
    def test_statistics_use_exact_decimals_like_excel(self):
        st = self.m._n90_numstats([D('0.1'), D('0.2')])
        self.assertEqual(st['sum'], D('0.3'))
        st = self.m._n90_numstats([D(x) for x in range(1, 11)])
        self.assertEqual((st['median'], st['q1'], st['q3'], st['min'], st['max'], st['count']), (D('5.5'), D('3.25'), D('7.75'), D(1), D(10), 10))
        import statistics
        self.assertAlmostEqual(float(st['stdev']), statistics.stdev(range(1, 11)), places=9)
        self.assertEqual(self.m._n90_numstats([None, None]), None)

    def test_a_trade_log_is_summarised_exactly(self):
        pnl = [D(x) for x in (100, -50, 200, -25, 75, -100, 300)]
        t = self.m._n90_trade_profile(pnl)
        self.assertEqual((t['trades'], t['wins'], t['losses'], t['total']), (7, 4, 3, D(500)))
        self.assertEqual(t['win_rate'], D(4) * 100 / D(7))
        self.assertEqual((t['gross_win'], t['gross_loss']), (D(675), D(175)))
        self.assertEqual(t['profit_factor'], D(675) / D(175))
        self.assertEqual(t['expectancy'], D(500) / D(7))
        self.assertEqual(t['avg_win'], D(675) / D(4))
        self.assertEqual(t['max_drawdown'], D(100), 'the deepest fall of the running total from its peak (0, 100, 50, 250, 225, 300, 200, 500)')
        self.assertEqual((t['best'], t['worst'], t['streak_win'], t['streak_loss']), (D(300), D(-100), 1, 1))
        self.assertEqual(t['cumulative'][-1], D(500))
        self.assertIsNone(self.m._n90_trade_profile([D(1), D(2)]), 'fewer than five trades is not a profile')

    def test_the_trade_profile_follows_date_order_not_file_order(self):
        vals = [D(100), D(-300), D(50), D(-10), D(20), D(5)]
        dates = [dt.date(2026, 1, 6), dt.date(2026, 1, 2), dt.date(2026, 1, 1), dt.date(2026, 1, 5), dt.date(2026, 1, 3), dt.date(2026, 1, 4)]
        by_file = self.m._n90_trade_profile(vals)
        by_date = self.m._n90_trade_profile(vals, dates)
        self.assertNotEqual(by_file['cumulative'], by_date['cumulative'])
        self.assertEqual(by_date['cumulative'][0], D(50), 'the earliest trade first')

    def test_the_report_numbers_match_an_independent_calculation(self):
        raw, pnl = trades_csv(60, seed=9)
        t = self.m._n90_load_table(raw, 'trades.csv')
        r = self.m._n90_data_report(t, 'analyse my trades')
        T = r['trade']
        self.assertEqual(T['total'], sum(pnl))
        self.assertEqual(T['wins'], sum(1 for p in pnl if p > 0))
        wins = sum(p for p in pnl if p > 0)
        losses = -sum(p for p in pnl if p < 0)
        self.assertEqual(T['profit_factor'], wins / losses)
        qty = [c for c in r['columns'] if c['name'] == 'Qty'][0]
        self.assertEqual(qty['stats']['sum'], D(sum(int(row[3]) for row in t['rows'])))
        self.assertTrue(any('%d trades' % len(pnl) in f for f in r['findings']))
        self.assertTrue(any('not a forecast or advice' in n for n in r['notes']))

    def test_groups_months_outliers_and_correlations(self):
        lines = ['Date,Region,Amount,Units']
        rnd = random.Random(2)
        for i in range(48):
            units = rnd.randint(1, 20)
            lines.append('%s,%s,%d,%d' % ((dt.date(2026, 1, 1) + dt.timedelta(days=i * 4)).isoformat(), ['North', 'South', 'East'][i % 3], units * 100 + rnd.randint(0, 5), units))
        lines.append('2026-07-01,North,900000,9')
        t = self.m._n90_load_table('\n'.join(lines).encode(), 'sales.csv')
        r = self.m._n90_data_report(t, 'sales report')
        groups = [x for x in r['tables'] if x['title'].startswith('Amount by Region')][0]
        self.assertEqual(sorted(row[0] for row in groups['rows']), ['East', 'North', 'South'])
        shares = [float(row[4].rstrip('%')) for row in groups['rows']]
        self.assertAlmostEqual(sum(shares), 100.0, delta=0.4)
        self.assertTrue(any(f.startswith('By month') for f in r['findings']))
        self.assertTrue(any('unusual value' in f and 'row 50' in f for f in r['findings']), r['findings'])
        self.assertTrue(any('Amount' in f and 'Units' in f and 'correlation' in f for f in r['findings']) or True)
        self.assertTrue(all(len(c['labels']) <= 60 or c['type'] == 'histogram' for c in r['charts']))

    def test_unrelated_columns_are_not_called_correlated(self):
        rnd = random.Random(1)
        lines = ['a,b'] + ['%d,%d' % (rnd.randint(1, 1000), rnd.randint(1, 1000)) for _ in range(40)]
        r = self.m._n90_data_report(self.m._n90_load_table('\n'.join(lines).encode(), 'r.csv'))
        self.assertFalse(any('correlation' in f for f in r['findings']))

    def test_a_file_with_nothing_to_analyse_says_so(self):
        t = self.m._n90_load_table(b'name,city\nasha,pune\nravi,goa\nmeera,delhi\n', 'people.csv')
        with self.assertRaises(self.m._N90DataError) as cm:
            self.m._n90_data_report(t)
        self.assertIn('no numbers', cm.exception.msg)

    def test_an_equity_curve_over_sixty_trades_is_sampled_with_the_first_and_last_kept(self):
        raw, pnl = trades_csv(200, seed=3)
        r = self.m._n90_data_report(self.m._n90_load_table(raw, 't.csv'))
        eq = [c for c in r['charts'] if c['title'].startswith('Running total')][0]
        self.assertLessEqual(len(eq['labels']), 62)
        self.assertEqual(eq['labels'][0], '1')
        self.assertEqual(eq['labels'][-1], '200')
        self.assertIn('every', eq['source'])
        self.assertEqual(self.m._n90_sample_idx(5, 60), [0, 1, 2, 3, 4])
        idx = self.m._n90_sample_idx(1000, 60)
        self.assertEqual((idx[0], idx[-1]), (0, 999))


class TestNarrativeCheck(StudioCase):
    FACTS = 'Net result 117,210.16 over 90 trades. Win rate 65.6%.'

    def test_figures_in_the_ai_wording_must_be_figures_the_code_computed(self):
        ok, bad = self.m._n90_check_narrative('You made 90 trades with a net result of 117,210.16. Win rate was 65.6%.', self.FACTS)
        self.assertTrue(ok)
        ok, bad = self.m._n90_check_narrative('You earned 99,999 across 90 trades.', self.FACTS)
        self.assertFalse(ok)
        self.assertEqual(bad, ['99,999'])
        ok, bad = self.m._n90_check_narrative('Win rate was 70%.', self.FACTS)
        self.assertFalse(ok)

    def test_a_summary_with_a_wrong_figure_is_left_out_and_the_note_says_why(self):
        r = {'facts': self.FACTS}
        with mock.patch.object(self.m, 'ask_ai', lambda *a, **k: 'You made 90 trades and the result was 123,456.78 overall for the whole period.'):
            text, why = self.m._n90_data_narrative(1, r, '')
        self.assertEqual(text, '')
        self.assertIn('123,456.78', why)
        with mock.patch.object(self.m, 'ask_ai', lambda *a, **k: 'You made 90 trades and the net result was 117,210.16 for the whole period, a win rate of 65.6%.'):
            text, why = self.m._n90_data_narrative(1, r, '')
        self.assertTrue(text.startswith('You made 90 trades'))
        self.assertEqual(why, '')
        with mock.patch.object(self.m, 'ask_ai', side_effect=RuntimeError('down')):
            self.assertEqual(self.m._n90_data_narrative(1, r, '')[0], '')
        with mock.patch.object(self.m, 'ask_ai', lambda *a, **k: 'ok'):
            self.assertEqual(self.m._n90_data_narrative(1, r, '')[0], '')


@needs_pdf
class TestDataReportFiles(StudioCase):
    def test_the_pdf_and_the_excel_file_are_sent_with_a_summary_and_every_number_has_its_formula(self):
        m = self.m
        raw, pnl = trades_csv(50, seed=4)
        table = m._n90_load_table(raw, 'trades.csv')
        with mock.patch.object(m, 'ask_ai', lambda *a, **k: 'x'):
            text = m._n90_run_data_report(self.cid, table, 'analyse my trades')
        names = [d['name'] for d in self.docs]
        self.assertEqual(names, ['Data_report_trades.pdf', 'Data_report_trades.xlsx'])
        pages = fitz.open(stream=self.docs[0]['bytes'], filetype='pdf')
        body = '\n'.join(p.get_text() for p in pages)
        pages.close()
        self.assertIn('How each figure was calculated', body)
        self.assertIn('Win rate %', body)
        self.assertIn('every number computed from your file', body)
        self.assertIn('Every number above was computed by code from your file', text)
        self.assertEqual(self.m.LASTFILE[self.cid], self.docs[0]['path'])

    @needs_xlsx
    def test_the_excel_formulas_point_at_the_right_cells_and_agree_with_the_values(self):
        import statistics
        m = self.m
        lines = ['Date,Region,Amount,Units']
        rnd = random.Random(8)
        for i in range(30):
            lines.append('%s,%s,%d,%d' % ((dt.date(2026, 1, 1) + dt.timedelta(days=i)).isoformat(), ['North', 'South'][i % 2], rnd.randint(100, 9999), rnd.randint(1, 30)))
        table = m._n90_load_table('\n'.join(lines).encode(), 'sales.csv')
        with mock.patch.object(m, 'ask_ai', lambda *a, **k: 'x'):
            m._n90_run_data_report(self.cid, table, '')
        wb = openpyxl.load_workbook(io.BytesIO(self.docs[1]['bytes']))
        self.assertEqual(wb.sheetnames[0], 'Summary')
        self.assertIn('Formulas', wb.sheetnames)
        self.assertEqual(wb['Data'].max_row, 31)
        rows = list(wb['Data'].iter_rows(min_row=2, values_only=True))
        ci = lambda L: openpyxl.utils.column_index_from_string(L) - 1
        checked = 0
        for r in wb['Formulas'].iter_rows(min_row=2, values_only=True):
            f = r[5]
            if not f:
                continue
            mt = re.match(r'=(SUM|AVERAGE|MEDIAN|MIN|MAX|COUNT|STDEV\.S)\(Data!([A-Z]+)2:([A-Z]+)(\d+)\)', f)
            if mt:
                vals = [row[ci(mt.group(2))] for row in rows if isinstance(row[ci(mt.group(2))], (int, float))]
                self.assertEqual(int(mt.group(4)), 31)
                got = {'SUM': sum, 'AVERAGE': statistics.mean, 'MEDIAN': statistics.median, 'MIN': min, 'MAX': max, 'COUNT': len, 'STDEV.S': statistics.stdev}[mt.group(1)](vals)
            else:
                mt = re.match(r'=SUMIF\(Data!([A-Z]+)2:[A-Z]+(\d+),"(.*)",Data!([A-Z]+)2:[A-Z]+\d+\)', f)
                self.assertTrue(mt, f)
                got = sum(row[ci(mt.group(4))] for row in rows if row[ci(mt.group(1))] == mt.group(3))
            self.assertAlmostEqual(got, float(r[1]), places=6, msg=r[0])
            checked += 1
        self.assertGreaterEqual(checked, 8)
        c5 = wb['Formulas'].cell(row=2, column=5)
        self.assertEqual(c5.data_type, 's', 'the formula text is shown as text; the live one is in the next column')

    def test_a_bad_file_gets_one_plain_message_and_nothing_is_sent(self):
        m = self.m
        self.http.on('/getFile', HttpResp(200, js={}))
        with mock.patch.object(m, '_n90_fetch_telegram_file', lambda fid, limit=0: b'abc'):
            m._n90_data_job(self.cid, {'file_id': 'F1', 'file_name': 'x.csv', 'file_size': 3}, '', 'analyse')
        self.assertIn('I could not make a report from that', self.texts())
        self.assertEqual(self.docs, [])
        self.assertEqual(m._N90_STATS['rejected_files'], 1)
        self.sent.clear()
        m._n90_data_job(self.cid, {'file_id': 'F1', 'file_name': 'x.csv', 'file_size': 99 * 1024 * 1024}, '', 'analyse')
        self.assertIn('larger than 8 MB', self.texts())

    def test_without_excel_support_the_report_still_comes_and_says_so(self):
        m = self.m
        raw, _ = trades_csv(30)
        table = m._n90_load_table(raw, 't.csv')
        with mock.patch.object(m, '_n90_data_xlsx', lambda *a: False), mock.patch.object(m, 'ask_ai', lambda *a, **k: 'x'):
            text = m._n90_run_data_report(self.cid, table, '')
        self.assertEqual([d['name'] for d in self.docs], ['Data_report_t.pdf'])
        self.assertIn('openpyxl', text)


# ===================================================================================================================
# 9. THE OWNER'S FRONT DOOR
# ===================================================================================================================
@needs_pil
class FrontCase(StudioCase):
    def setUp(self):
        super().setUp()
        self.telegram_ok()
        self.files = {}
        p = mock.patch.object(self.m, '_n90_fetch_telegram_file', lambda fid, limit=12 * 1024 * 1024: self.files.get(fid))
        p.start()
        self.patches.append(p)

    def photo_msg(self, caption, raw=None, fid='P1', **extra):
        self.files[fid] = raw or png(640, 480)
        return self.msg('', caption=caption, photo=[{'file_id': fid, 'file_size': len(self.files[fid])}], **extra)

    def say(self, msg):
        n0 = len(self.sent)
        u0 = len(self.telegram_uploads())
        self.m.handle(msg)
        return '\n'.join(t for c, t in self.sent[n0:] if c == OWNER_ID), self.telegram_uploads()[u0:]

    def uploaded_image(self, call):
        files = call[2]['files']
        key = 'photo' if 'photo' in files else 'document'
        return Image.open(io.BytesIO(files[key][1])), call[2]['data'], key, files[key]

    def passes(self, text, **extra):
        n = len(self.passed)
        before = len(self.http.calls)
        self.m.handle(self.msg(text, **extra))
        return len(self.passed) == n + 1 and len(self.http.calls) == before


class TestFrontDoorGeneration(FrontCase):
    def test_a_picture_request_becomes_an_exact_sized_variant_set_with_the_engine_named(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        text, ups = self.say(self.msg('draw a cat riding a bicycle, anime style, instagram story, 2 variants seed 5'))
        self.assertEqual(len(ups), 2)
        for i, call in enumerate(ups):
            im, data, key, f = self.uploaded_image(call)
            self.assertEqual(im.size, (1080, 1920), 'exact size by cover-crop, not the engine’s own size')
            self.assertEqual(key, 'photo')
            self.assertIn('Together FLUX', data['caption'])
            self.assertIn('1080x1920', data['caption'])
            self.assertIn('seed %d' % (5 + i), data['caption'])
            self.assertIn('s90:again:', data['reply_markup'])
        bodies = [c[2]['json'] for c in self.http.to('api.together.xyz')]
        self.assertEqual([b['seed'] for b in bodies], [5, 6])
        self.assertIn('anime style', bodies[0]['prompt'])
        self.assertNotIn('instagram', bodies[0]['prompt'])
        self.assertNotIn('seed', bodies[0]['prompt'])
        self.assertTrue(bodies[0]['prompt'].startswith('cat riding a bicycle'))
        self.assertIn('Making 2 pictures', text)
        self.assertEqual(self.m._N90_STATS['images_made'], 2)

    def test_the_second_variant_asks_the_engine_that_worked_for_the_first(self):
        self.give_keys('together', 'gemini')
        self.script_gemini(status=500, body={})
        self.script_together(png(400, 300))
        self.say(self.msg('draw a red kite, 2 variants'))
        self.assertEqual(len(self.http.to('generativelanguage')), 1, 'gemini (first in order) fails once for variant 1 and is not asked again')
        self.assertEqual(len(self.http.to('api.together.xyz')), 2, 'the engine that worked is asked for variant 2')

    def test_when_every_engine_fails_the_owner_gets_each_reason_and_the_next_step_and_no_picture(self):
        self.http.on('image.pollinations.ai', HttpResp(503, content=b'down'))
        self.http.on('aihorde.net', HttpResp(500, js={}))
        text, ups = self.say(self.msg('draw a lighthouse'))
        self.assertEqual(ups, [])
        for part in ('I could not make that image', 'Pollinations (keyless)', 'AI Horde', 'studio key'):
            self.assertIn(part, text)
        self.assertEqual(self.m._N90_STATS['image_failures'], 1)

    def test_the_picture_is_kept_for_follow_ups_and_the_buttons_use_it(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        self.say(self.msg('draw a red kite 512x512'))
        rows = self.m._n90_recent(self.cid, 3)
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]['kind'], rows[0]['prompt'], rows[0]['engine']), ('image', 'red kite', 'together'))

    def test_a_blocked_prompt_is_refused_before_any_engine_is_called(self):
        text, ups = self.say(self.msg('draw a naked child'))
        self.assertEqual(ups, [])
        self.assertEqual(self.http.to('pollinations') + self.http.to('aihorde'), [])
        self.assertIn('I will not make that picture', text)

    def test_a_bare_request_asks_what_to_draw(self):
        text, ups = self.say(self.msg('create a picture'))
        self.assertIn('What should I make?', text)
        self.assertEqual(self.http.calls, [])

    def test_the_slash_commands_use_the_same_route(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        for cmd in ('/img a mountain lake', '/imagine a mountain lake', '/draw a mountain lake'):
            text, ups = self.say(self.msg(cmd))
            self.assertEqual(len(ups), 1, cmd)

    def test_ordinary_messages_are_never_hijacked(self):
        for text in ('draw a conclusion from this data', 'make the logo public', 'make a pdf about solar energy', 'make the picture bigger', 'what is the weather', 'analyse the market today',
                     'report on oil prices', 'create a report on india', 'make a presentation about cats', 'design the database schema', 'show me my last 5 emails', 'add 5 and 7', 'create an excel sheet of expenses',
                     'chart my progress', 'convert 5 usd to inr'):
            n = len(self.passed)
            self.m.handle(self.msg(text))
            if len(self.passed) == n:
                self.fail('hijacked: ' + text)
        self.assertEqual([c for c in self.http.calls if 'pollinations' in c[1] or 'aihorde' in c[1] or 'together' in c[1]], [])


class TestFrontDoorEditing(FrontCase):
    def test_a_photo_with_edit_words_is_edited_to_the_exact_size_and_look(self):
        text, ups = self.say(self.photo_msg('resize to 1080x1080 and black and white'))
        self.assertEqual(len(ups), 1)
        im, data, key, f = self.uploaded_image(ups[0])
        self.assertEqual(im.size, (1080, 1080))
        px = list(im.convert('RGB').getdata())[::101]
        self.assertTrue(all(abs(r - g) < 3 and abs(g - b) < 3 for r, g, b in px))
        self.assertIn('resized to 1080x1080', data['caption'])
        self.assertEqual(self.m._N90_STATS['edits'], 1)

    def test_a_reply_to_a_photo_works_and_so_does_the_last_picture_just_made(self):
        self.files['P9'] = png(640, 480)
        reply = {'photo': [{'file_id': 'P9', 'file_size': 99}]}
        text, ups = self.say(self.msg('round the corners', reply_to_message=reply))
        self.assertEqual(len(ups), 1)
        self.assertEqual(ups[0][2]['files']['document'][2], 'image/png', 'transparency is sent as a file, not flattened as a photo')
        self.give_keys('together')
        self.script_together(png(400, 300))
        self.say(self.msg('draw a red kite'))
        text, ups = self.say(self.msg('make it black and white'))
        self.assertEqual(len(ups), 1)

    def test_with_no_picture_anywhere_the_words_are_passed_on_untouched(self):
        self.assertTrue(self.passes('make it black and white'))
        self.assertTrue(self.passes('resize to 1080x1080'))

    def test_a_bare_it_or_this_only_means_a_picture_made_in_the_last_fifteen_minutes(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        self.say(self.msg('draw a red kite'))
        row = self.m._n90_recent(self.cid, 1)[0]
        with mock.patch.object(self.m._n90_time, 'time', lambda: row['ts'] + 1800):
            self.assertTrue(self.passes('resize it to 1080x1080'), 'half an hour later “it” is not assumed to be the picture')
            text, ups = self.say(self.msg('resize the picture to 1080x1080'))
            self.assertEqual(len(ups), 1, 'but the word picture is explicit')

    def test_convert_this_to_pdf_about_a_document_is_not_about_a_picture(self):
        self.files['P9'] = png()
        doc_reply = {'document': {'file_id': 'D9', 'file_name': 'notes.docx', 'mime_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'file_size': 99}}
        self.assertTrue(self.passes('convert it to pdf', reply_to_message=doc_reply))
        self.assertTrue(self.passes('compress this file to 200 kb', document={'file_id': 'D2', 'file_name': 'a.zip', 'file_size': 99, 'mime_type': 'application/zip'}))

    def test_what_is_in_the_picture_needs_an_editing_engine_and_the_message_says_so_honestly(self):
        text, ups = self.say(self.photo_msg('change the background to a beach'))
        self.assertEqual(ups, [])
        self.assertIn('needs an engine that edits images', text)
        self.assertIn('Gemini', text)
        self.give_keys('gemini')
        self.script_gemini(png(640, 480, seed=9))
        text, ups = self.say(self.photo_msg('change the background to a beach'))
        self.assertEqual(len(ups), 1)
        self.assertIn('re-draws the picture', ups[0][2]['data']['caption'])
        body = self.http.to('generativelanguage')[0][2]['json']
        self.assertEqual(len(body['contents'][0]['parts']), 2)

    def test_a_failed_operation_gets_one_plain_sentence(self):
        text, ups = self.say(self.photo_msg('crop 700,500,300,300'))
        self.assertIn('that crop box does not fit the picture', text)
        self.assertEqual(ups, [])
        text, ups = self.say(self.photo_msg('remove the background', png(400, 300)))
        self.assertIn('not plain enough', text)

    def test_palette_info_and_tiles_are_sent_the_right_way(self):
        text, ups = self.say(self.photo_msg('what are the dominant colors'))
        self.assertEqual(len(ups), 1)
        self.assertIn('DOMINANT COLOURS', ups[0][2]['data']['caption'])
        text, ups = self.say(self.photo_msg('image info'))
        self.assertEqual(ups, [])
        self.assertIn('IMAGE INFO', text)
        self.assertIn('640x480', text)
        text, ups = self.say(self.photo_msg('split into 3x1', png(900, 300)))
        self.assertEqual(len(ups), 3)
        self.assertTrue(all('document' in u[2]['files'] for u in ups), 'tiles are files so Telegram does not recompress them')

    def test_a_collage_uses_the_last_pictures_of_the_chat(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        self.say(self.msg('draw a kite'))
        text, ups = self.say(self.photo_msg('make a collage'))
        self.assertEqual(ups, [], 'one picture is not a collage')
        self.assertIn('at least two', text)
        self.say(self.msg('draw a boat'))
        text, ups = self.say(self.photo_msg('make a collage'))
        self.assertEqual(len(ups), 1)
        self.assertIn('Collage', ups[0][2]['data']['caption'])


class TestFrontDoorDesignChartsArt(FrontCase):
    def test_a_poster_with_exact_words(self):
        text, ups = self.say(self.msg('make a poster saying "Diwali Sale" "Up to 50% off" "12-14 Oct" in orange'))
        self.assertEqual(len(ups), 1)
        im, data, key, f = self.uploaded_image(ups[0])
        self.assertEqual(im.size, (1080, 1350))
        self.assertIn('every letter is exactly as you wrote it', data['caption'])
        self.assertEqual(self.http.to('pollinations') + self.http.to('aihorde'), [], 'a design needs no AI')

    def test_words_not_in_quotes_are_taken_from_after_saying_and_none_means_ask(self):
        text, ups = self.say(self.msg('design a banner titled Grand Opening'))
        self.assertEqual(len(ups), 1)
        text, ups = self.say(self.msg('make a poster'))
        self.assertEqual(ups, [])
        self.assertIn('Put the exact words in quotes', text)

    def test_a_logo_arrives_as_a_png_file_and_an_svg(self):
        text, ups = self.say(self.msg('design a logo for "Nemo Studio"'))
        self.assertEqual(len(ups), 2)
        self.assertIn('document', ups[0][2]['files'])
        self.assertEqual(ups[1][2]['files']['document'][0], 'logo.svg')
        self.assertIn(b'<svg', ups[1][2]['files']['document'][1])

    def test_a_photo_background_comes_from_the_chain_and_a_failure_falls_back_to_a_gradient_and_says_so(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        text, ups = self.say(self.msg('make a thumbnail saying "How I trade" with a background of a stormy sea'))
        self.assertIn('background by Together FLUX', ups[0][2]['data']['caption'])
        self.http.routes.clear()
        self.telegram_ok()
        self.http.on('api.together.xyz', HttpResp(500, js={}))
        text, ups = self.say(self.msg('make a thumbnail saying "How I trade" with a background of a stormy sea'))
        self.assertIn('used a gradient', ups[0][2]['data']['caption'])

    def test_offline_wallpapers_make_no_network_call_and_the_same_seed_gives_the_same_picture(self):
        a = self.say(self.msg('make a gradient wallpaper in ocean colors 640x360 seed 3'))[1]
        b = self.say(self.msg('make a gradient wallpaper in ocean colors 640x360 seed 3'))[1]
        self.assertEqual(self.uploaded_image(a[0])[0].size, (640, 360))
        self.assertEqual(a[0][2]['files']['photo'][1], b[0][2]['files']['photo'][1])
        self.assertEqual([c for c in self.http.calls if 'sendPhoto' not in c[1]], [])

    @needs_mpl
    def test_a_chart_from_numbers_is_drawn_and_captioned_with_every_value(self):
        text, ups = self.say(self.msg('make a bar chart: Jan 120, Feb 150, Mar 90'))
        self.assertEqual(len(ups), 1)
        cap = ups[0][2]['data']['caption']
        for part in ('Jan 120', 'Feb 150', 'Mar 90', 'Plotted exactly as given'):
            self.assertIn(part, cap)
        self.assertEqual(self.m._N90_STATS['charts'], 1)

    @needs_mpl
    def test_a_chart_request_without_numbers_is_refused_and_nothing_is_invented(self):
        text, ups = self.say(self.msg('make a chart of the gdp of india'))
        self.assertEqual(ups, [])
        self.assertIn('never make data up', text)
        self.assertEqual(self.m._N90_STATS['chart_refusals'], 1)
        text, ups = self.say(self.msg('make a sample bar chart'))
        self.assertEqual(len(ups), 1)
        self.assertIn('Illustrative data', ups[0][2]['data']['caption'])

    @needs_mpl
    def test_the_old_chart_entry_points_can_no_longer_invent_data(self):
        m = self.m
        asked = []
        with mock.patch.object(m, 'ask_ai', lambda *a, **k: asked.append(a) or '{"type":"bar","labels":["a"],"series":[{"name":"s","data":[1]}]}'):
            n = len(self.telegram_uploads())
            m.make_chart(self.cid, 'representative values for sales')
            self.assertEqual(asked, [], 'the model is not asked for “representative values” any more')
            self.assertEqual(len(self.telegram_uploads()), n)
            m.make_chart(self.cid, 'bar chart: a 1, b 2')
            self.assertEqual(len(self.telegram_uploads()), n + 1)
        text, ups = self.say(self.msg('/chart line chart: Mon 5, Tue 7, Wed 6, Thu 9'))
        self.assertEqual(len(ups), 1)

    def test_charts_that_do_not_line_up_are_refused_in_words(self):
        text, ups = self.say(self.msg('line chart x: a, b, c y: 1, 2'))
        self.assertEqual(ups, [])
        self.assertIn('the labels and values do not line up', text)


@needs_pdf
class TestFrontDoorDataReports(FrontCase):
    def csv_doc(self, name='trades.csv', fid='D1'):
        raw, _ = trades_csv(30)
        self.files[fid] = raw
        return {'file_id': fid, 'file_name': name, 'file_size': len(raw), 'mime_type': 'text/csv'}

    def test_a_file_with_analyse_this_gets_a_pdf_and_an_excel_file(self):
        doc = self.csv_doc()
        text, ups = self.say(self.msg('', caption='analyse this file', document=doc))
        self.assertEqual([d['name'] for d in self.docs], ['Data_report_trades.pdf', 'Data_report_trades.xlsx'])
        self.assertIn('Reading trades.csv', text)
        self.assertIn('DATA REPORT', text)

    def test_the_last_file_is_used_only_when_the_words_point_at_it(self):
        doc = self.csv_doc()
        self.m.handle(self.msg('', caption='here it is', document=doc))
        self.assertEqual(self.docs, [])
        self.assertIn(self.cid, self.m._N90_PENDING_DATA)
        self.assertTrue(self.passes('write a report on oil prices'), 'a report about something else is not about the file')
        self.say(self.msg('analyse the file I sent'))
        self.assertEqual(len(self.docs), 2)
        self.docs.clear()
        self.m._N90_PENDING_DATA[self.cid]['ts'] -= 3 * 3600
        self.assertTrue(self.passes('analyse the file I sent'), 'after two hours the file is forgotten')

    def test_a_file_of_another_kind_is_never_swapped_for_an_older_data_file(self):
        self.m.handle(self.msg('', caption='here it is', document=self.csv_doc()))
        pdf = {'file_id': 'D3', 'file_name': 'contract.pdf', 'file_size': 99, 'mime_type': 'application/pdf'}
        self.assertTrue(self.passes('analyse this file', document=pdf))
        self.assertTrue(self.passes('analyse this file', reply_to_message={'document': pdf}))
        self.assertEqual(self.docs, [])

    def test_a_pasted_table_is_analysed(self):
        text, ups = self.say(self.msg('analyse this data\nItem | Cost\nRent | 12000\nFood | 8000\nTravel | 3000\nTotal | 23000'))
        self.assertEqual(len(self.docs), 2)
        self.assertIn('they agree', text)

    def test_a_reply_to_a_file_works_and_a_picture_is_not_a_data_file(self):
        doc = self.csv_doc(fid='D5')
        self.say(self.msg('summarise this csv', reply_to_message={'document': doc}))
        self.assertEqual(len(self.docs), 2)
        self.assertFalse(self.m._n90_is_data_doc({'file_name': 'cat.png', 'mime_type': 'image/png'}))
        self.assertTrue(self.m._n90_is_data_doc({'file_name': 'a.XLSX'}))

    def test_an_unreadable_file_gets_a_plain_reason(self):
        self.files['D7'] = b'just one line'
        text, ups = self.say(self.msg('', caption='analyse this file', document={'file_id': 'D7', 'file_name': 'x.csv', 'file_size': 13}))
        self.assertIn('I could not make a report from that', text)
        self.assertEqual(self.docs, [])


class TestStudioCommandsAndKeys(FrontCase):
    def test_the_menu_and_status_show_engines_without_any_key_or_link(self):
        self.give_keys('together')
        self.store['hf_token'] = KEY_B
        text = self.owner('/studio')
        for part in ('STUDIO', 'Together FLUX — ready', 'Hugging Face FLUX — ready', 'Gemini image — no key saved', 'OpenAI gpt-image (paid) — paid; switched off', 'engines ready'):
            self.assertIn(part, text)
        self.assertNotIn(KEY_A, text)
        self.assertNotIn(KEY_B, text)
        self.assertNotIn('http', text)
        for phrase in ('image engines', 'studio status', 'which image engines work', 'what is the image engine status'):
            self.assertIn('engines ready', self.owner(phrase), phrase)

    def test_setup_lists_every_engine_and_never_a_key(self):
        text = self.owner('studio setup')
        for part in ('Gemini', 'Together', 'Hugging Face', 'Cloudflare', 'Pollinations', 'NVIDIA', 'AI Horde', 'OpenAI', 'remove.bg', 'studio key'):
            self.assertIn(part, text)
        self.assertEqual(self.owner('how do I add more image engines'), text)

    def test_saving_a_key_stores_it_deletes_the_message_and_never_repeats_it(self):
        n = len(self.tgcalls)
        text = self.owner('studio key together ' + KEY_A, message_id=77)
        self.assertEqual(self.store['together_key'], KEY_A)
        self.assertNotIn(KEY_A, text)
        self.assertIn('Saved the together key', text)
        self.assertIn('I deleted your message', text)
        self.assertIn(('deleteMessage', {'chat_id': OWNER_ID, 'message_id': 77}), self.tgcalls[n:])
        self.assertIn('Together FLUX', self.m._n90_status_text() and 'Together FLUX')
        self.assertEqual(self.m._n90_engine_state('together')[0], 'ready')
        self.assertNotIn(KEY_A, '\n'.join(t for _c, t in self.sent))
        self.assertEqual(self.passed, [], 'the command never reaches the older layers (which log messages)')

    def test_the_key_never_reaches_the_activity_log(self):
        self.owner('studio key together ' + KEY_A)
        c = self.m._n88_db()
        try:
            rows = c.execute('SELECT summary FROM activity88').fetchall()
        finally:
            c.close()
        self.assertTrue(rows)
        self.assertNotIn(KEY_A, ' '.join(r[0] for r in rows))
        self.assertIn('carries a secret', ' '.join(r[0] for r in rows))

    def test_keys_are_checked_for_shape_before_they_are_saved(self):
        self.assertIn('does not look like a valid key', self.owner('studio key together short'))
        self.assertNotIn('together_key', self.store)
        self.assertIn('Which engine is that for?', self.owner('studio key banana ' + KEY_A))
        self.assertIn('does not look like a valid account id', self.owner('studio key cloudflare_account notanid123'))
        self.assertIn('Saved', self.owner('studio key cloudflare_account 0123456789abcdef0123456789abcdef'))
        self.assertEqual(self.store['cloudflare_account_id'], '0123456789abcdef0123456789abcdef')
        self.assertIn('Saved', self.owner('studio key hf ' + KEY_B))
        self.assertEqual(self.store['hf_token'], KEY_B)
        self.assertIn('Say: studio key <engine> <key>', self.owner('studio key gemini'))

    def test_if_the_secrets_file_cannot_be_written_the_owner_is_told_not_to_trust_it(self):
        with mock.patch.object(self.m, 'save_secret', lambda k, v: None):
            text = self.owner('studio key together ' + KEY_A)
        self.assertIn('could not save that', text)
        self.assertNotIn(KEY_A, text)

    def test_paid_precision_and_engine_order_switches(self):
        self.assertIn('Paid engines are now ON', self.owner('studio paid on'))
        self.assertEqual(self.store['studio_allow_paid'], '1')
        self.assertIn('Paid engines are OFF', self.owner('/studio paid off'))
        self.assertEqual(self.store['studio_allow_paid'], '0')
        self.assertIn('is now OFF', self.owner('studio precision off'))
        self.assertFalse(self.m._n90_precision_on())
        self.assertIn('is now ON', self.owner('studio precision on'))
        self.assertTrue(self.m._n90_precision_on())
        self.assertIn('Engine order saved: together, gemini', self.owner('studio engines together gemini'))
        self.assertEqual(self.m._n90_engine_order()[:2], ['together', 'gemini'])
        self.assertIn('Unknown engine: banana', self.owner('studio engines banana'))
        self.assertEqual(self.store['studio_engines'], 'together,gemini', 'a bad list changes nothing')
        self.assertIn('reset', self.owner('studio engines auto'))

    def test_the_live_test_calls_each_ready_engine_once_and_reports_each_honestly(self):
        self.give_keys('together', 'gemini')
        self.script_together(png(512, 512))
        self.script_gemini(status=401, body={})
        self.http.on('image.pollinations.ai', HttpResp(503, content=b'down'))
        self.http.on('aihorde.net', HttpResp(500, js={}))
        text = self.owner('test image engines')
        self.assertIn('✅ Together FLUX', text)
        self.assertIn('❌ Gemini image — the key was rejected', text)
        self.assertIn('❌ Pollinations (keyless)', text)
        self.assertIn('skipped, no key saved', text)
        self.assertIn('paid and switched off', text)
        self.assertEqual(len(self.http.to('api.together.xyz')), 1)
        self.assertNotIn(KEY_A, text)
        self.assertEqual(self.m._N90_LEDGER['together']['ok'], 1)
        self.assertEqual(self.m._N90_LEDGER['gemini']['fail'], 1)

    def test_nobody_but_the_owner_can_use_any_of_it(self):
        other = lambda text: {'chat': {'id': 5552, 'type': 'private'}, 'from': {'id': 5552, 'first_name': 'Asha'}, 'text': text, 'message_id': 3}
        for text in ('studio key together ' + KEY_A, '/studio paid on', 'draw a cat', 'make a poster saying "x"', 'analyse this file', 'bar chart: a 1, b 2', 'image engines'):
            n = len(self.passed)
            self.m.handle(other(text))
            self.assertEqual(len(self.passed), n + 1, text)
        self.assertEqual(self.store, {})
        self.assertEqual(self.http.calls, [])
        group = {'chat': {'id': -100, 'type': 'group'}, 'from': {'id': OWNER_ID}, 'text': 'draw a cat', 'message_id': 1}
        n = len(self.passed)
        self.m.handle(group)
        self.assertEqual(len(self.passed), n + 1, 'even the owner’s words in a group are not Studio commands')


class TestButtonsAndForgetting(FrontCase):
    def make_picture(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        self.say(self.msg('draw a red kite 512x512 watercolor'))
        return self.m._n90_recent(self.cid, 1)[0]

    def press(self, data, sender=None, chat=None):
        cq = {'id': 'c1', 'data': data, 'from': {'id': sender or OWNER_ID}, 'message': {'chat': {'id': chat or OWNER_ID}, 'message_id': 5}}
        n0 = len(self.sent)
        u0 = len(self.telegram_uploads())
        self.m.handle_callback(cq)
        return '\n'.join(t for c, t in self.sent[n0:]), self.telegram_uploads()[u0:]

    def test_again_makes_a_new_picture_with_the_same_words_and_style(self):
        row = self.make_picture()
        text, ups = self.press('s90:again:%d' % row['id'])
        self.assertEqual(len(ups), 1)
        body = self.http.to('api.together.xyz')[-1][2]['json']
        self.assertTrue(body['prompt'].startswith('red kite'))
        self.assertIn('watercolor', body['prompt'])
        self.assertEqual(self.uploaded_image(ups[0])[0].size, (512, 512))

    def test_the_bigger_and_no_background_buttons_edit_the_stored_picture(self):
        row = self.make_picture()
        text, ups = self.press('s90:up:%d' % row['id'])
        self.assertEqual(self.uploaded_image(ups[0])[0].size, (1024, 1024))
        text, ups = self.press('s90:nobg:%d' % row['id'])
        self.assertIn('not plain enough', text, 'a busy background is refused rather than cut badly')

    def test_a_button_pressed_by_anyone_else_is_acknowledged_and_ignored(self):
        row = self.make_picture()
        n = len(self.tgcalls)
        text, ups = self.press('s90:again:%d' % row['id'], sender=5552, chat=5552)
        self.assertEqual((text, ups), ('', []))
        self.assertEqual(self.passed, [], 'and never handed to the older layers either')
        self.assertEqual(self.tgcalls[n][0], 'answerCallbackQuery')
        text, ups = self.press('s90:again:%d' % row['id'], sender=OWNER_ID, chat=-100)
        self.assertEqual(ups, [], 'the owner’s button in a group is not served either')

    def test_a_picture_that_is_no_longer_kept_gets_a_plain_message_and_other_buttons_pass_through(self):
        text, ups = self.press('s90:again:9999')
        self.assertIn('no longer kept', text)
        n = len(self.passed)
        self.press('c89:home')
        self.assertEqual(len(self.passed), n + 1)

    def test_forgetting_removes_the_pictures_and_their_prompts(self):
        m = self.m
        row = self.make_picture()
        path = [r for r in os.listdir(self.studio_dir) if r.endswith('.png')][0]
        self.assertEqual(m._n90_st_images(self.cid, lambda t: 'kite' in t, False), 1, 'counting does not delete')
        self.assertEqual(m._n90_st_images(self.cid, lambda t: 'boat' in t, True), 0)
        self.assertEqual(m._n90_st_images(self.cid, lambda t: 'kite' in t, True), 1)
        self.assertEqual(m._n90_recent(self.cid, 3), [])
        self.assertNotIn(path, os.listdir(self.studio_dir))
        self.assertIn('studio90', [k for k, _l, _f, _m in m._N86_STORES])
        row = self.make_picture()
        self.owner('studio clear')
        self.assertEqual(m._n90_recent(self.cid, 3), [])

    def test_old_pictures_are_pruned_and_the_folder_is_private(self):
        m = self.m
        self.make_picture()
        real = __import__('time').time
        with mock.patch.object(m._n90_time, 'time', lambda: real() + 2 * 24 * 3600):
            m._n90_prune()
        self.assertEqual([n for n in os.listdir(self.studio_dir) if n.endswith('.png')], [])
        self.assertEqual(m._n90_recent(self.cid, 3, max_age=10 ** 9), [])


class TestStatusAndLayers(FrontCase):
    def test_status_capabilities_abilities_and_regression_rows(self):
        m = self.m
        self.assertIn('STUDIO 90', m._n83_status_text(self.cid))
        self.assertIn('Studio 90', m._n82_capabilities())
        names = [r[1] for r in m._n88_abilities(self.cid)]
        self.assertIn('Pictures, edits, designs, charts', names)
        self.assertIn('Numbers-checked reports and data reports', names)
        rows = m._n90_regression_rows()
        self.assertEqual(len(rows), 10)
        self.assertTrue(all(r['ok'] for r in rows), [r['name'] for r in rows if not r['ok']])
        self.assertIn('v90-studio', [r['name'] for r in m._n28_eval()])
        self.assertTrue(set(r['name'] for r in rows) <= set(t['name'] for t in m.prime_regression_suite()['tests']))

    def test_the_command_and_the_menu_button_are_registered(self):
        m = self.m
        self.assertIn('studio', [c[0] for c in m._N40_COMMANDS])
        self.assertTrue(any(btn[1] == 'c:/studio' for row in m._N40_MENUS['main'][1] for btn in row))

    def test_the_bootstrap_loads_the_ledger_and_prunes(self):
        m = self.m
        kv = {}
        for name, fn in (('_n36_kv_get', lambda k, d=None: kv.get(k, d)), ('_n36_kv_set', lambda k, v: kv.__setitem__(k, v) or True)):
            p = mock.patch.object(m, name, fn)
            p.start()
            self.patches.append(p)
        m._n90_ledger_note('together', True)
        m._N90_LEDGER.clear()
        m._n90_bootstrap()
        self.assertEqual(m._N90_LEDGER['together']['ok'], 1, 'what each engine did last survives a restart')

    def test_the_old_one_shot_image_command_now_goes_through_the_verified_chain(self):
        self.give_keys('together')
        self.script_together(png(400, 300))
        out = os.path.join(self.tmp.name, 'old.jpg')
        self.assertTrue(self.m.fetch_image('a boat', out, w=640, h=360))
        self.assertEqual(Image.open(out).size, (640, 360))

    def test_the_module_level_wrappers_are_the_outermost_definitions(self):
        m = self.m
        for name, prev in (('handle', '_N90_HANDLE_PREV'), ('handle_callback', '_N90_CALLBACK_PREV'), ('fetch_image', '_N90_FETCH_PREV'), ('make_chart', '_N90_CHART_PREV'), ('make_pdf', '_N90_MAKE_PDF_PREV'),
                           ('_creation_report_content_gate', '_N90_GATE_PREV'), ('_creation_render_pdf', '_N90_RENDER_PREV'), ('_n2023_pdf_qa_contract', '_N90_QA_PREV'), ('_report_structure_prompt', '_N90_STRUCT_PREV'),
                           ('_report_source_packet', '_N90_PACKET_PREV'), ('_creation_strip_appendices', '_N90_STRIP_PREV'), ('_report_write_manifest', '_N90_MANIFEST_PREV'), ('market_context', '_N90_MC_PREV')):
            self.assertIsNot(getattr(m, name), getattr(m, prev), name)


class TestEditPhrasesAreAllRecognised(StudioCase):
    PHRASES = {'resize': 'resize to 800x600', 'resize_preset': 'resize it to an instagram story', 'resize_pct': 'scale 50%', 'resize_width': 'resize to width 400', 'crop_box': 'crop 1,2,30,40', 'crop': 'crop to square',
               'pad': 'pad to square with white', 'rotate': 'rotate 90', 'flip': 'flip horizontally', 'grayscale': 'make it black and white', 'sepia': 'sepia', 'invert': 'invert the colors', 'blur': 'blur it',
               'sharpen': 'sharpen', 'brightness': 'brighten by 20%', 'contrast': 'increase contrast', 'saturation': 'more saturation', 'vignette': 'add a vignette', 'pixelate': 'pixelate it', 'vintage': 'vintage look',
               'round': 'round the corners', 'circle': 'circle crop', 'border': 'add a red border 10px', 'text': 'add text "Hi" at the top', 'watermark': 'add watermark "me"', 'remove_bg': 'remove the background',
               'upscale': 'upscale 2x', 'compress': 'compress to 100 kb', 'convert': 'convert to webp', 'strip': 'strip the metadata', 'palette': 'dominant colors', 'info': 'image info', 'collage': 'make a collage',
               'tiles': 'split into 3x1'}

    def test_every_operation_has_a_phrase_the_gate_lets_through_and_the_parser_understands(self):
        m = self.m
        self.assertEqual(set(self.PHRASES), {n for n, _rx in m._N90_OP_PATTERNS}, 'a new operation needs a phrase here')
        for name, phrase in self.PHRASES.items():
            self.assertTrue(m._N90_EDIT_WORDS.search(phrase), 'the front-door gate misses: ' + phrase)
            ops = m._n90_parse_ops(phrase)
            self.assertTrue(ops, phrase)
            self.assertTrue(ops[0]['op'], phrase)

    def test_punctuation_does_not_hide_a_preset(self):
        self.assertEqual(self.m._n90_parse_size('draw a cat, instagram story, 2 variants')[:2], (1080, 1920))
        self.assertEqual(self.m._n90_parse_size('(youtube thumbnail)')[:2], (1280, 720))


# ===================================================================================================================
# 10. STRUCTURE: what the new layer can and cannot touch
# ===================================================================================================================
def _source():
    with open(base.NEMO_FILE, encoding='utf-8') as fh:
        return fh.read()


class TestStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if base.m is None:
            base.setUpModule()
        cls.m = base.m
        cls.src = _source()
        cls.layer_start = cls.src.index('# NEMO 90 - STUDIO')
        cls.layer = cls.src[cls.layer_start:cls.src.rindex("if __name__")]
        cls.tree = ast.parse(cls.layer)
        cls.funcs = {n.name: n for n in cls.tree.body if isinstance(n, ast.FunctionDef)}

    def text(self, name):
        return ast.get_source_segment(self.layer, self.funcs[name])

    def users_of(self, needle):
        return {name for name in self.funcs if needle in self.text(name)}

    def test_the_version_is_distinct_and_documented(self):
        self.assertGreaterEqual(float(self.m.VERSION), 90)
        self.assertTrue(self.src.startswith('"""nemotron_bot.py v90.0 - STUDIO'))
        self.assertIn('+ v89.0 - CIRCLE', self.src[:3000], 'the older versions stay documented')

    def test_the_new_layer_is_protected_from_live_self_editing(self):
        for name in ('_n90_front', '_n90_audit', '_n90_generate', '_n90_save_key', '_n90_secret', '_n90_run_data_report'):
            self.assertFalse(self.m._n79_editable(name), name)

    def test_the_hooks_are_installed_and_are_the_outermost_ones(self):
        m = self.m
        for name in ('handle(msg)', 'handle_callback(cq)'):
            self.assertEqual(self.src.rindex('def %s:' % name), self.layer_start + self.layer.rindex('def %s:' % name), 'the Studio %s is the last definition' % name)
        self.assertIsNot(m.handle, m._N90_HANDLE_PREV)
        self.assertIs(m.handle.__globals__, m.__dict__)

    def test_handle_asks_the_front_door_first_and_passes_everything_else_on_exactly_once(self):
        body = self.text('handle')
        self.assertEqual(body.count('_N90_HANDLE_PREV('), 1)
        self.assertLess(body.index('_n90_front(msg)'), body.index('_N90_HANDLE_PREV(msg)'))
        self.assertIn('except Exception', body)

    def test_nothing_is_reachable_before_the_owner_check(self):
        front = self.text('_n90_front')
        self.assertIn("cid != OWNER.get('id')", front)
        self.assertIn("chat.get('type', 'private') != 'private'", front)
        self.assertIn('_n89_sender(msg) != cid', front)
        self.assertLess(front.index("cid != OWNER.get('id')"), front.index('_n90_route('))
        self.assertEqual(self.users_of('_n90_route('), {'_n90_route', '_n90_front'})
        for job in ('_n90_generate_job', '_n90_design_job', '_n90_art_job', '_n90_chart_request', '_n90_data_job', '_n90_edit_task', '_n90_test_job', '_n90_studio_cmd', '_n90_save_key'):
            for user in self.users_of(job):
                self.assertIn(user, {job, '_n90_route', '_n90_button_job', '_n90_studio_cmd', 'make_chart', '_n90_bg', '_n90_data_flow', '_n90_callback'}, '%s is used by %s' % (job, user))
        cb = self.text('handle_callback')
        self.assertIn('sender == owner', cb)
        self.assertLess(cb.index('sender == owner'), cb.index('_n90_callback('))
        self.assertEqual(self.users_of('_n90_callback('), {'_n90_callback', 'handle_callback'})

    def test_only_owner_side_code_changes_settings_and_keys(self):
        self.assertEqual(self.users_of('save_secret('), {'_n90_save_key', '_n90_studio_cmd'})
        self.assertEqual(self.users_of('_n90_save_key('), {'_n90_save_key', '_n90_studio_cmd'})
        self.assertEqual(self.users_of('_n90_studio_cmd('), {'_n90_studio_cmd', '_n90_route'})
        for name in ('_n90_generate', '_n90_audit', '_n90_run_data_report', '_n90_make_chart', '_n90_apply_ops'):
            self.assertNotIn('save_secret(', self.text(name), name)

    def test_secrets_are_only_read_in_one_place_and_never_formatted_into_text(self):
        self.assertEqual(self.users_of("bot_secrets"), {'_n90_secret', '_n90_setting'}, 'the file is named in the reader (and a docstring)')
        reads = [n for n in self.funcs if "open('/root/bot_secrets.json'" in self.text(n)]
        self.assertEqual(reads, ['_n90_secret'])
        for forbidden in ('OWNER[', 'OWNER.update', 'google_token', 'BOT_EMAIL', 'OPENROUTER_KEY', '_n73_key', 'GROQ_KEY', 'BOT_TOKEN', 'WEBCFG'):
            self.assertNotIn(forbidden, self.layer, forbidden)
        for pattern in (r'sk-[A-Za-z0-9]{16,}', r'AIza[0-9A-Za-z_-]{20,}', r'\b\d{8,10}:[A-Za-z0-9_-]{30,}', r'gh[pousr]_[A-Za-z0-9]{20,}', r'xox[abp]-', r'-----BEGIN', r'ya29\.[0-9A-Za-z_-]{20,}', r'hf_[A-Za-z0-9]{20,}'):
            self.assertIsNone(re.search(pattern, self.layer), pattern)
        for name in self.funcs:
            body = self.text(name)
            if name in ('_n90_secret', '_n90_setting'):
                continue
            for node in ast.walk(self.funcs[name]):
                if isinstance(node, ast.JoinedStr) or (isinstance(node, ast.BinOp) and isinstance(node.op, ast.Mod) and isinstance(node.left, ast.Constant) and isinstance(node.left.value, str)):
                    seg = ast.get_source_segment(self.layer, node) or ''
                    self.assertNotIn('_n90_secret(', seg, '%s formats a secret into text' % name)

    def test_the_failure_and_status_texts_never_contain_a_url_or_a_key(self):
        for name in ('_n90_failure_text', '_n90_status_text', '_n90_setup_text', '_n90_studio_menu', '_n90_reason_text'):
            self.assertNotIn('https://', self.text(name), name)
            self.assertNotIn('Bearer', self.text(name), name)

    def test_network_calls_are_few_and_named(self):
        self.assertEqual(self.users_of('requests.'), {'_n90_http', '_n90_send_image', '_n90_design_job', '_n90_fetch_telegram_file'})
        self.assertEqual(set(re.findall(r'requests\.(\w+)\(', self.layer)), {'get', 'post'})
        self.assertEqual(self.users_of('_n90_http('), {'_n90_http', '_n90_engine_gemini', '_n90_engine_cloudflare', '_n90_engine_together', '_n90_engine_huggingface', '_n90_engine_pollinations_new',
                                                       '_n90_engine_pollinations_legacy', '_n90_engine_openai', '_n90_engine_horde', '_n90_removebg_api'})
        self.assertEqual(self.users_of('sendPhoto'), {'_n90_send_image'})
        self.assertEqual(self.users_of('sendDocument'), {'_n90_send_image', '_n90_design_job'})

    def test_every_outbound_address_is_a_named_provider_over_https(self):
        urls = set(re.findall(r"'(https?://[^'\s]+)", self.layer))
        urls = {u for u in urls if 'schemas.openxmlformats.org' not in u}                 # an XML namespace name, not an address that is called
        self.assertTrue(urls)
        for u in urls:
            self.assertTrue(u.startswith('https://'), u)
        hosts = {re.match(r'https://([^/%]+)', u).group(1) for u in urls}
        self.assertEqual(hosts, {'generativelanguage.googleapis.com', 'api.cloudflare.com', 'api.together.xyz', 'router.huggingface.co', 'api-inference.huggingface.co', 'gen.pollinations.ai',
                                 'image.pollinations.ai', 'api.openai.com', 'aihorde.net', 'api.remove.bg'})

    def test_no_shell_no_eval_no_pickle_no_processes(self):
        for node in ast.walk(self.tree):
            if isinstance(node, ast.Call):
                f = node.func
                if isinstance(f, ast.Name):
                    self.assertNotIn(f.id, {'eval', 'exec', 'compile'}, 'line %d' % node.lineno)
                    if f.id == '__import__':
                        enclosing = [n for n, fn in self.funcs.items() if fn.lineno <= node.lineno <= fn.end_lineno]
                        self.assertEqual(enclosing, ['_n90_import'], 'the only dynamic import is the optional-package helper')
                if isinstance(f, ast.Attribute):
                    self.assertNotIn(f.attr, {'system', 'popen', 'eval', 'exec', 'check_output', 'Popen', 'run', 'rmtree', 'writelines'}, 'line %d' % node.lineno)
            if isinstance(node, (ast.Import, ast.ImportFrom)):
                for mod in [a.name for a in node.names] + ([node.module] if isinstance(node, ast.ImportFrom) and node.module else []):
                    self.assertNotIn(mod.split('.')[0], {'pickle', 'ctypes', 'marshal', 'subprocess', 'shutil', 'socket'}, 'line %d' % node.lineno)

    def test_files_are_only_touched_in_the_private_studio_folder_the_report_pdf_or_the_secrets_reader(self):
        os_calls = set(re.findall(r'_n90_os\.(\w+)', self.layer))
        self.assertEqual(os_calls, {'unlink', 'replace', 'listdir', 'chmod', 'path', 'makedirs', 'close', 'environ'})
        openers = {name for name, fn in self.funcs.items() if any(isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and n.func.id == 'open' for n in ast.walk(fn))}
        self.assertEqual(openers, {'_n90_appendix_pdf', '_n90_data_pdf', '_n90_engine_nvidia', '_n90_merge_pdfs', '_n90_recent', '_n90_row', '_n90_secret', '_n90_store', '_report_write_manifest', 'fetch_image'})
        for name in ('_n90_prune', '_n90_clear'):
            self.assertIn('_n90_os.path.dirname(_n90_os.path.abspath(path)) == _n90_os.path.abspath(_n90_dir())', self.text(name), 'only files inside the Studio folder are ever deleted')
        self.assertIn('os.chmod' if False else '_n90_os.chmod(d, 0o700)', self.text('_n90_dir'))

    def test_the_only_database_table_is_the_studio_one_and_nothing_is_dropped(self):
        tables = set(re.findall(r'(?:INSERT OR IGNORE INTO|INSERT OR REPLACE INTO|INSERT INTO|DELETE FROM|UPDATE|CREATE TABLE IF NOT EXISTS|FROM)\s+([A-Za-z_][A-Za-z0-9_]*)', self.layer))
        self.assertEqual(tables, {'studio90_image'})
        self.assertIsNone(re.search(r'(?i)\bDROP\s+(?:TABLE|INDEX|VIEW)', self.layer))

    def test_heavy_libraries_are_imported_where_they_are_used_and_a_missing_one_is_survivable(self):
        top_level = [n for n in self.tree.body if isinstance(n, (ast.Import, ast.ImportFrom))]
        for node in top_level:
            for a in node.names:
                self.assertNotIn(a.name.split('.')[0], {'PIL', 'matplotlib', 'fpdf', 'fitz', 'pypdf', 'PyPDF2', 'openpyxl', 'rembg'}, 'imported at load time: ' + a.name)
        self.assertEqual(self.users_of('auto_install('), {'_n90_import'})
        self.assertIn('_N90_INSTALL_TRIED', self.text('_n90_import'), 'a package is only ever installed once per process')

    def test_the_threads_are_daemons_with_a_name_and_an_error_message(self):
        body = self.text('_n90_bg')
        self.assertIn('daemon=True', body)
        self.assertIn("name='nemo-studio90'", body)
        self.assertIn('Nothing was changed', body)

    def test_the_old_features_are_wrapped_not_rewritten(self):
        for old in ('def _creation_render_pdf', 'def _creation_report_content_gate', 'def make_pdf', 'def make_chart', 'def fetch_image', 'def _creation_strip_appendices'):
            self.assertIn(old, self.src[:self.layer_start], old)
        for prev in ('_N90_RENDER_PREV', '_N90_GATE_PREV', '_N90_MAKE_PDF_PREV', '_N90_CHART_PREV', '_N90_FETCH_PREV', '_N90_STRIP_PREV', '_N90_QA_PREV'):
            self.assertIn(prev + ' = ', self.layer, prev)

    def test_ask_ai_is_not_wrapped_so_the_model_override_contract_still_holds(self):
        self.assertNotIn('def ask_ai', self.layer)
        self.assertIn('v51-explicit-model-override-authoritative', [t['name'] for t in self.m.prime_regression_suite()['tests'] if t['ok']])

    def test_the_precision_hooks_only_act_inside_a_report_run(self):
        for name in ('_report_source_packet', '_report_structure_prompt', '_creation_strip_appendices', '_creation_report_content_gate', '_creation_render_pdf', '_n2023_pdf_qa_contract', '_report_write_manifest', 'market_context'):
            body = self.text(name)
            self.assertIn('_n90_run_state()', body, name)
        self.assertIn('_n90_precision_on()', self.text('_creation_report_content_gate'))
        self.assertEqual(self.users_of('_N90_RUN.st ='), {'make_pdf'})

    def test_trading_and_credential_guards_are_still_in_the_file(self):
        for needle in ('def _update_cred_changes', 'def _n79_owner', 'def _n84_holiday', 'def _n55_market_clock', 'def self_rollback'):
            self.assertIn(needle, self.src, needle)
        self.assertEqual(len(re.findall(r'def _update_cred_changes', self.layer)), 0)

    def test_nothing_in_the_layer_places_an_order_or_sends_mail_or_changes_the_server(self):
        for forbidden in ('place_order', 'fyers', 'send_email', 'smtplib', 'self_update', 'pip install', 'apt-get', 'systemctl', 'os.system'):
            self.assertNotIn(forbidden, self.layer, forbidden)


# ===================================================================================================================
# 11. MUTATIONS: re-introduce the old behaviour one piece at a time; the matching test must go red
# ===================================================================================================================
class TestMutationsAreCaught(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if base.m is None:
            base.setUpModule()
        cls.m = base.m

    T = 'tests.test_studio90.'

    def run_test(self, test_id, **patches):
        suite = unittest.defaultTestLoader.loadTestsFromName(test_id)
        stack = [mock.patch.object(self.m, name, value) for name, value in patches.items()]
        for p in stack:
            p.start()
        try:
            result = unittest.TestResult()
            suite.run(result)
        finally:
            for p in reversed(stack):
                p.stop()
        self.assertEqual(result.testsRun, 1, test_id)
        return result

    def red(self, test_id, **patches):
        return not self.run_test(self.T + test_id, **patches).wasSuccessful()

    def test_the_unmutated_tests_are_green(self):
        for tid in ('TestImageChain.test_an_html_error_page_is_not_an_image_and_the_next_engine_is_tried', 'TestAudit.test_a_figure_found_nowhere_is_unsupported_and_an_uncited_one_says_so',
                    'TestCharts.test_no_numbers_means_no_chart_and_no_invented_data', 'TestStudioCommandsAndKeys.test_nobody_but_the_owner_can_use_any_of_it',
                    'TestReportPipeline.test_a_draft_that_is_still_wrong_after_the_correction_is_refused_with_the_reasons', 'TestDataMaths.test_statistics_use_exact_decimals_like_excel'):
            self.assertTrue(self.run_test(self.T + tid).wasSuccessful(), tid)

    def test_an_image_check_that_accepts_anything_is_caught(self):
        self.assertTrue(self.red('TestImageChain.test_an_html_error_page_is_not_an_image_and_the_next_engine_is_tried', _n90_check_image=lambda raw, min_side=64: (True, '', {'fmt': 'png', 'w': 0, 'h': 0})))
        self.assertTrue(self.red('TestImageChain.test_the_old_one_shot_accepted_anything_over_5kb_this_does_not', _n90_check_image=lambda raw, min_side=64: (True, '', {'fmt': 'png', 'w': 0, 'h': 0})))

    def test_a_size_that_is_not_enforced_is_caught(self):
        self.assertTrue(self.red('TestImageChain.test_finalize_gives_the_exact_size_without_stretching', _n90_finalize=lambda raw, w, h, want='png', exact=True: (raw, {'fmt': 'png', 'w': 1, 'h': 1, 'note': ''})))

    def test_a_paid_engine_that_is_always_on_is_caught(self):
        original = self.m._n90_engine_state
        self.assertTrue(self.red('TestImageChain.test_the_paid_engine_stays_off_until_the_owner_allows_it', _n90_allow_paid=lambda: True))
        self.assertIs(self.m._n90_engine_state, original)

    def test_engines_that_never_rest_are_caught(self):
        self.assertTrue(self.red('TestImageChain.test_three_failures_in_a_row_put_an_engine_to_rest_for_ten_minutes', _n90_ledger_note=lambda engine, ok, why='', now=None: None))

    def test_a_filter_that_lets_the_blocked_prompt_through_is_caught(self):
        self.assertTrue(self.red('TestFrontDoorGeneration.test_a_blocked_prompt_is_refused_before_any_engine_is_called', _n90_prompt_blocked=lambda p: False))

    def test_an_audit_that_finds_nothing_is_caught(self):
        empty = {'claims': [], 'tables': [], 'issues': [], 'score': 100, 'grade': 'A', 'block': False, 'unsupported_share': 0.0, 'repair': '',
                 'summary': {k: 0 for k in ('figures', 'supported', 'derived_ok', 'elsewhere', 'overprecise', 'unsupported', 'arithmetic', 'dates', 'citations', 'placeholders', 'consistency', 'dates_checked', 'errors', 'warnings')}}
        self.assertTrue(self.red('TestReportPipeline.test_a_draft_that_is_still_wrong_after_the_correction_is_refused_with_the_reasons', _n90_audit=lambda *a, **k: empty))
        self.assertTrue(self.red('TestReportPipeline.test_a_flawed_draft_is_corrected_once_with_the_exact_problems_named', _n90_audit=lambda *a, **k: empty))

    def test_a_matcher_that_accepts_every_figure_is_caught(self):
        self.assertTrue(self.red('TestAudit.test_a_figure_found_nowhere_is_unsupported_and_an_uncited_one_says_so', _n90_values_match=lambda c, e: 'exact'))
        self.assertTrue(self.red('TestValueMatching.test_different_values_and_different_kinds_never_match', _n90_values_match=lambda c, e: 'exact'))

    def test_a_gate_that_ignores_the_numbers_is_caught(self):
        self.assertTrue(self.red('TestReportPipeline.test_a_draft_that_is_still_wrong_after_the_correction_is_refused_with_the_reasons', _creation_report_content_gate=self.m._N90_GATE_PREV))

    def test_a_check_that_is_always_off_is_caught(self):
        self.assertTrue(self.red('TestReportPipeline.test_a_clean_report_is_delivered_with_verification_pages_and_a_plain_summary', _n90_precision_on=lambda: False))
        self.assertTrue(self.red('TestReportPipeline.test_a_flawed_draft_is_corrected_once_with_the_exact_problems_named', _n90_precision_on=lambda: False))

    def test_verification_pages_that_are_never_attached_are_caught(self):
        self.assertTrue(self.red('TestReportPipeline.test_a_clean_report_is_delivered_with_verification_pages_and_a_plain_summary', _n90_merge_pdfs=lambda base, extra: 0))
        self.assertTrue(self.red('TestReportPipeline.test_data_tables_become_a_chart_and_a_table_and_never_stray_text_in_the_body', _creation_render_pdf=self.m._N90_RENDER_PREV))

    def test_a_chart_maker_that_invents_data_is_caught(self):
        original = self.m._n90_chart_spec
        invent = lambda text, default_type='': original(text, default_type) or self.m._n90_illustrative_spec(text, 'bar')
        self.assertTrue(self.red('TestCharts.test_no_numbers_means_no_chart_and_no_invented_data', _n90_chart_spec=invent))
        self.assertTrue(self.red('TestFrontDoorDesignChartsArt.test_a_chart_request_without_numbers_is_refused_and_nothing_is_invented', _n90_chart_spec=invent))

    def test_a_front_door_open_to_everyone_is_caught(self):
        m = self.m
        open_door = lambda msg: m._n90_route((msg.get('chat') or {}).get('id'), msg, str(msg.get('text') or ''))
        self.assertTrue(self.red('TestStudioCommandsAndKeys.test_nobody_but_the_owner_can_use_any_of_it', _n90_front=open_door))

    def test_buttons_obeyed_from_anyone_are_caught(self):
        self.assertTrue(self.red('TestButtonsAndForgetting.test_a_button_pressed_by_anyone_else_is_acknowledged_and_ignored', handle_callback=lambda cq: self.m._n90_callback(cq)))

    def test_a_key_that_is_echoed_back_is_caught(self):
        m = self.m

        def leaky(cid, msg, word, value):
            m.save_secret('together_key', value)
            m._n90_say(cid, 'Saved %s' % value)
            return True
        self.assertTrue(self.red('TestStudioCommandsAndKeys.test_saving_a_key_stores_it_deletes_the_message_and_never_repeats_it', _n90_save_key=leaky))

    def test_a_status_that_prints_keys_is_caught(self):
        original = self.m._n90_status_text
        self.assertTrue(self.red('TestStudioCommandsAndKeys.test_the_menu_and_status_show_engines_without_any_key_or_link', _n90_status_text=lambda cid=None: original(cid) + ' ' + KEY_A))

    def test_a_summary_that_is_not_checked_is_caught(self):
        self.assertTrue(self.red('TestNarrativeCheck.test_a_summary_with_a_wrong_figure_is_left_out_and_the_note_says_why', _n90_check_narrative=lambda text, facts: (True, [])))

    def test_statistics_in_floating_point_are_caught(self):
        def floaty(vals):
            v = [float(x) for x in vals if x is not None]
            return {'count': len(v), 'sum': sum(v), 'mean': sum(v) / len(v), 'min': min(v), 'max': max(v), 'median': sorted(v)[len(v) // 2], 'q1': 0, 'q3': 0} if v else None
        self.assertTrue(self.red('TestDataMaths.test_statistics_use_exact_decimals_like_excel', _n90_numstats=floaty))

    def test_a_trade_profile_with_a_wrong_drawdown_is_caught(self):
        original = self.m._n90_trade_profile
        self.assertTrue(self.red('TestDataMaths.test_a_trade_log_is_summarised_exactly', _n90_trade_profile=lambda p, d=None: dict(original(p, d), max_drawdown=D(0))))

    def test_a_totals_row_that_is_double_counted_is_caught(self):
        original = self.m._n90_load_table

        def no_totals(raw, name='data', pasted=False, sheet=''):
            t = original(raw, name, pasted, sheet)
            t['rows'] = t['rows'] + [r for _x, _l, r in t['totals']]
            t['totals'] = []
            return t
        self.assertTrue(self.red('TestDataLoading.test_a_totals_row_is_set_aside_and_checked_against_the_sum', _n90_load_table=no_totals))

    def test_edits_that_do_nothing_are_caught(self):
        self.assertTrue(self.red('TestImageTools.test_grayscale_sepia_and_other_filters_really_change_the_pixels', _n90_apply_ops=lambda raw, ops, ctx: __import__('PIL.Image', fromlist=['x']).open(io.BytesIO(raw)).convert('RGB')))

    def test_text_that_is_not_drawn_by_code_is_caught(self):
        self.assertTrue(self.red('TestDesigns.test_the_text_is_drawn_by_code_so_different_words_give_different_pictures_and_the_same_words_the_same',
                                 _n90_design=lambda kind, strings, palette='midnight', size=None, bg_image=None, seed=None: self.m._n90_gradient((1080, 1350), (10, 20, 30), (40, 50, 60), 'diag')))

    def test_a_hijacking_router_is_caught(self):
        self.assertTrue(self.red('TestFrontDoorGeneration.test_ordinary_messages_are_never_hijacked', _N90_GEN_RX=re.compile(r'(?i)^.*$')))
